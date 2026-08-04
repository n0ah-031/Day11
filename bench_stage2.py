#!/usr/bin/env python3
"""2단계 AI 재검증 병렬화 실측 — 순차(동시성 1) 대비 병렬(기본 6) 배수.

HANDOFF §4의 '미측정: 실제 키로 돌린 순차 대비 병렬 배수'를 재는 스크립트다.
같은 입력을 두 번 돌리므로 호출 수는 파일 수 × 2다. `--dry`는 호출을 스텁으로 바꿔
호출 수·항목 수만 센다(비용 없음).

실행: python3 bench_stage2.py [--files 30] [--workers 6] [--dry]
"""
from __future__ import annotations

import argparse
import copy
import os
import sys
import tempfile
import time
from pathlib import Path

import openpyxl

import aggregate as ag

ROWS = [
    ("기획팀", "정보시스템 유지관리 용역", 12_500_000, "2026-03-02"),
    ("기획팀", "업무용 노트북 교체", 8_400_000, "2026-03-11"),
    ("인사팀", "신입사원 입문 교육 위탁", 5_200_000, "2026-03-05"),
    ("인사팀", "직원 건강검진 지원", 3_100_000, "2026-03-19"),
    ("총무팀", "청사 소방설비 정기점검", 2_750_000, "2026-03-07"),
    ("총무팀", "사무용품 일괄 구매", 1_180_000, "2026-03-22"),
    ("안전팀", "지사 안전점검 외부 진단", 9_900_000, "2026-03-14"),
    ("안전팀", "보호구 정기 교체", 4_300_000, "2026-03-26"),
    ("회계팀", "세무 자문 수수료", 2_200_000, "2026-03-09"),
    ("회계팀", "결산 시스템 라이선스 갱신", 6_700_000, "2026-03-28"),
]


def make_file(path: Path, dept: str) -> None:
    """1단계를 전부 통과하는(= 2단계에 진입하는) 실무 형태 파일."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "예산집행"
    ws["A1"] = f"{dept} 2026년 1분기 예산 집행 현황"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.append([None, None, None, None])
    ws.append(["부서", "사업내용", "예산액", "집행일자"])
    for row in ROWS:
        ws.append(list(row))
    wb.save(path)


def parse(paths: list[Path]) -> list[ag.UploadedFile]:
    files = []
    for p in paths:
        uf = ag.read_file(p)
        ag.review_stage1(uf, {})
        files.append(uf)
    return files


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--files", type=int, default=30)
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--dry", action="store_true")
    a = ap.parse_args()

    ag.load_env(Path(__file__).parent / ".env")
    model = os.environ.get("OPENAI_MODEL", "gpt-5-mini")

    tmp = Path(tempfile.mkdtemp(prefix="bench-"))
    paths = []
    for i in range(a.files):
        p = tmp / f"{i + 1:02d}_지사.xlsx"
        make_file(p, f"{i + 1:02d}지사")
        paths.append(p)

    files = parse(paths)
    clean = [uf for uf in files if uf.readable and not uf.issues]
    items = sum(len(ag._ai_items(uf)) for uf in clean)
    batches = sum(-(-len(ag._ai_items(uf)) // 100) for uf in clean)
    print(f"파일 {len(files)}개 · 2단계 진입 {len(clean)}개 · 항목 {items}건 · 배치(=호출) {batches}건")
    print(f"→ 두 번 돌리므로 실제 호출은 {batches * 2}건")
    if len(clean) != len(files):
        print("경고: 1단계에서 걸린 파일이 있어 벤치마크 대상이 줄었다", file=sys.stderr)

    if a.dry:
        calls = {"n": 0}
        real = ag._ai_call

        def stub(client, model, batch):
            calls["n"] += 1
            time.sleep(0.05)
            return {it["id"]: {"id": it["id"], "verdict": "적합", "reason": ""} for it in batch}

        ag._ai_call = stub
        ag._ai_client = lambda: object()
        for workers in (1, a.workers):
            fresh = copy.deepcopy(files)
            t0 = time.perf_counter()
            ag.review_stage2_many(fresh, model, workers=workers)
            print(f"[dry] 동시성 {workers}: {time.perf_counter() - t0:.2f}초 (호출 {calls['n']}건 누적)")
        ag._ai_call = real
        return 0

    results = {}
    for workers in (1, a.workers):
        fresh = copy.deepcopy(files)
        t0 = time.perf_counter()
        ag.review_stage2_many(fresh, model, workers=workers)
        elapsed = time.perf_counter() - t0
        verdicts = sum(1 for uf in fresh if not uf.ai_unverified)
        unverified = [uf.name for uf in fresh if uf.ai_unverified]
        flagged = sum(len(uf.issues) for uf in fresh)
        results[workers] = elapsed
        print(f"동시성 {workers:>2}: {elapsed:6.1f}초 · 판정 완료 {verdicts}/{len(fresh)}개 · "
              f"AI 미검증 {len(unverified)}개 · 2단계 지적 {flagged}건", flush=True)
        if unverified:
            print(f"  미검증 파일: {unverified[:5]}", flush=True)

    seq, par = results[1], results[a.workers]
    print(f"\n순차 {seq:.1f}초 → 병렬(동시성 {a.workers}) {par:.1f}초 = **{seq / par:.2f}배**")
    print(f"파일당 순차 {seq / max(1, len(clean)):.1f}초")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
