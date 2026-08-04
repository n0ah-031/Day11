"""F1 AI 양식 생성 — 문답으로 사양을 모아 xlsx 양식과 작성기준을 만든다.

두 번의 LLM 역할이 있다(생성기능 기술명세 D1·D3).

1. **문답**(`intake_turn`) — 질문 순서·완료 판정을 LLM이 주도한다. 대신 백엔드가 두 가지
   가드를 건다: 완결성 루브릭(§4.3)과 최대 턴 수. LLM이 "다 됐다"고 해도 사양에 빈 곳이
   있으면 생성으로 넘기지 않는다.
2. **저작**(`author_workbook`) — LLM이 셀 단위로 파일 내용을 짓고(`workbook_json`),
   `materialize`는 그걸 해석 없이 xlsx로 옮기기만 한다. 레이아웃 규칙을 코드에 넣지 않는
   이유는, 양식마다 제목행·안내문·시트 분리 요구가 달라서 규칙을 넣는 만큼 갇히기 때문이다.

**LLM 산출물은 신뢰하지 않는다.** `validate_workbook`이 통과시킨 것만 파일로 만든다
(수식 함수 화이트리스트·외부 참조 금지·셀 수 상한·스타일 속성 화이트리스트).

마스킹은 `aggregate.mask`를 그대로 쓴다 — 외부 전송 직전 공통 통과 지점이 이미 거기 있고
(명세 §13.2), 두 번째 마스커를 만들면 규칙이 갈라진다. 저장은 원문이고 마스킹은 전송할
때만 한다(§9.2).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import aggregate as ag

TURN_LIMIT = 30           # §4.3 최대 턴 가드
AUTHOR_RETRY = 2          # §6.3 검증 실패 시 재저작 횟수
MAX_SHEETS = 10
MAX_CELLS = 10_000
ATTACH_CHARS = 8_000      # §6.5 첨부 추출 텍스트 절단

FIELD_TYPES = ("text", "number", "amount", "date", "name", "dept_code")
STYLE_KEYS = ("bold", "italic", "size", "bg", "color", "border", "align", "number_format")
ALLOWED_FUNCS = ("SUM", "SUMIF", "SUMIFS", "COUNT", "IF", "TODAY")
VALIDATION_TYPES = ("date", "list", "whole", "decimal")

REF_RE = re.compile(r"^\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}$")
RANGE_RE = re.compile(r"^\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}(:\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})?$")
FUNC_RE = re.compile(r"([A-Za-z][A-Za-z0-9_.]*)\s*\(")
BAD_SHEET_CHARS = set(r"\/?*[]:")

INTAKE_SYSTEM = """당신은 공공기관 업무 양식을 설계하는 인터뷰어입니다.
사용자가 배포할 엑셀 양식의 사양을 문답으로 확정하는 것이 목표입니다.

반드시 확인할 것:
1) 양식 이름
2) 배포 대상(전체 부서인지, 특정 부서 목록인지)
3) 포함할 항목(열)
4) 항목별 입력 형식 — 날짜 표기(예: YYYY-MM-DD), 금액 단위·구분기호 유무, 코드표 사용 여부
5) 항목별 필수/선택 구분

규칙:
- 한 번에 한 주제만 묻습니다. 여러 질문을 한꺼번에 쏟지 마세요.
- 사용자가 "알아서 해달라"고 하면 실무 관례에 맞는 기본값을 제안하고 동의를 구합니다.
- 이미 답이 나온 것을 다시 묻지 마세요.
- **사용자가 정정하면 그 지시가 이깁니다.** 예를 들어 "비고만 선택이고 나머지는 필수"라고
  하면 비고를 제외한 모든 항목의 required를 true로 바꿉니다. 앞서 당신이 제안했던 값을
  그대로 남겨두지 마세요. 필수/선택이 틀리면 회신본 검토에서 그대로 오판으로 이어집니다.
- 확인된 내용은 매 턴 spec_json 전체 갱신본으로 누적합니다.
- 남은 확인 사항이 없을 때만 spec_complete를 true로 둡니다.

반드시 아래 JSON만 출력합니다.
{
  "reply": "사용자에게 보여줄 한국어 메시지(질문 또는 확인 요청)",
  "spec_json": {
    "form_title": "양식 이름",
    "target_depts": ["전체 부서"],
    "fields": [
      {"name": "항목명", "type": "text|number|amount|date|name|dept_code",
       "required": true, "format": "형식 설명 또는 null", "notes": "비고 또는 null"}
    ],
    "layout_hints": "레이아웃 요구(자유 텍스트) 또는 null",
    "locale": "ko"
  },
  "spec_complete": false,
  "coverage": {"confirmed_topics": [], "remaining_topics": []}
}
type이 date이거나 amount인 항목은 format을 반드시 채웁니다."""

AUTHOR_SYSTEM = f"""당신은 확정된 사양(spec_json)을 받아 배포용 엑셀 양식 파일의 내용을
셀 단위로 저작합니다. 회신자가 값을 채워 되돌려줄 빈 양식입니다.

실무 관례를 따르세요:
- 최상단에 양식 제목 행(가로 병합, 굵게, 가운데)
- 필요하면 제목 아래 작성 안내 문구 한두 줄
- 헤더 행은 굵게 + 배경색 + 테두리, 필수 항목은 이름 끝에 * 표시
- 헤더 아래로 회신자가 채울 빈 행을 넉넉히(100행 안팎) 두고, 그 범위에 입력 형식에 맞는
  유효성 검사를 겁니다
- 유효성 검사 type은 {", ".join(VALIDATION_TYPES)} 만 씁니다(정수는 whole). 드롭다운은
  type을 list로 두고 선택지를 source 배열에 넣습니다
- **선택지가 사양에 실제 목록으로 주어진 경우에만** 드롭다운을 만듭니다. "코드표 사용",
  "전체 부서"처럼 목록이 아닌 설명만 있으면 드롭다운을 만들지 말고 자유 입력으로 두세요 —
  없는 선택지를 지어내면 회신자가 값을 넣을 수 없습니다
- 열 너비는 항목 성격에 맞게 지정
- 헤더 행에 freeze_panes를 걸어 스크롤해도 보이게 합니다

제약(위반 시 파일로 만들지 않습니다):
- 시트 {MAX_SHEETS}개 이하, 셀 합계 {MAX_CELLS}개 이하
- 시트명 31자 이하, {"".join(sorted(BAD_SHEET_CHARS))} 사용 금지, 중복 금지
- 수식은 {", ".join(ALLOWED_FUNCS)} 만 사용, 다른 파일 참조 금지
- style 속성은 {", ".join(STYLE_KEYS)} 만 사용

반드시 아래 JSON만 출력합니다.
{{
  "sheets": [
    {{
      "name": "시트명",
      "columns": [{{"index": 1, "width": 14}}],
      "merges": ["A1:F1"],
      "freeze_panes": "A4",
      "styles": {{"title": {{"bold": true, "size": 14, "align": "center"}}}},
      "cells": [{{"ref": "A1", "value": "제목", "style": "title"}}],
      "validations": [{{"range": "B4:B103", "type": "date",
                        "prompt": "YYYY-MM-DD 형식으로 입력", "error": "날짜 형식이 올바르지 않습니다"}}]
    }}
  ]
}}"""


class FormGenError(Exception):
    """생성을 진행할 수 없는 상태. 사용자에게 그대로 보여줄 문장을 담는다."""


# ── LLM 호출 ─────────────────────────────────────────────────────────────────

def _client():
    """OpenAI 클라이언트.

    F2의 2단계 재검증은 AI가 없으면 건너뛰고 취합을 계속하지만(§13.3), 양식 생성은
    AI가 기능 자체다. 조용히 빈 양식을 만드는 대신 사유를 들고 실패한다(§9.4).
    """
    try:
        from openai import OpenAI
        return OpenAI()
    except Exception as exc:
        raise FormGenError(
            "AI 서비스에 연결할 수 없어 양식을 생성할 수 없습니다. "
            f"잠시 후 다시 시도해주세요. (사유: {exc})") from exc


def _ask(system: str, payload: str, model: str, client=None) -> dict:
    """JSON 응답 1회. payload는 마스킹을 이미 통과한 문자열이어야 한다."""
    client = client or _client()
    try:
        res = client.chat.completions.create(
            model=model,
            response_format={"type": "json_object"},
            messages=[{"role": "system", "content": system},
                      {"role": "user", "content": payload}],
        )
        text = res.choices[0].message.content
        # 유료 호출이라 사용량을 남긴다. 스텁 클라이언트에는 usage가 없다.
        usage = getattr(res, "usage", None)
        if usage is not None:
            print(f"[formgen] {model} 토큰 in={usage.prompt_tokens} "
                  f"out={usage.completion_tokens} total={usage.total_tokens}", file=sys.stderr)
    except FormGenError:
        raise
    except Exception as exc:
        raise FormGenError(f"AI 응답을 받지 못했습니다. 다시 시도해주세요. (사유: {exc})") from exc
    try:
        return json.loads(text)
    except (TypeError, json.JSONDecodeError) as exc:
        raise FormGenError("AI 응답을 해석할 수 없습니다. 다시 시도해주세요.") from exc


def _masked(obj) -> str:
    """외부로 나가는 payload는 전부 마스킹을 통과시킨다 (§9.2·§13.2)."""
    return ag.mask(json.dumps(obj, ensure_ascii=False))


# ── 문답 ─────────────────────────────────────────────────────────────────────

def rubric_gaps(spec: dict) -> list[str]:
    """사양이 생성 가능한 상태인지 확인하고, 부족한 것을 돌려준다 (§4.3 가드 1).

    LLM이 spec_complete=true를 줘도 이 목록이 비지 않으면 완료로 인정하지 않는다.
    """
    gaps = []
    if not (spec.get("form_title") or "").strip():
        gaps.append("양식 이름")
    fields = spec.get("fields") or []
    if not fields:
        gaps.append("포함할 항목")
    for f in fields:
        name = (f.get("name") or "").strip() or "(이름 없는 항목)"
        if f.get("type") not in FIELD_TYPES:
            gaps.append(f"'{name}'의 자료 유형")
        if not isinstance(f.get("required"), bool):
            gaps.append(f"'{name}'의 필수/선택 구분")
        if f.get("type") in ("date", "amount") and not (f.get("format") or "").strip():
            gaps.append(f"'{name}'의 입력 형식")
    return list(dict.fromkeys(gaps))


def intake_turn(messages: list[dict], spec: dict, model: str,
                attachments: list[dict] | None = None, client=None) -> dict:
    """문답 한 턴. `{reply, spec_json, spec_complete, coverage}`를 돌려준다.

    messages는 원문 그대로 받고(저장본), 전송 직전에만 마스킹한다.
    """
    turn = len([m for m in messages if m.get("role") == "user"])
    payload = {
        "conversation": [{"role": m.get("role"), "content": m.get("content")} for m in messages],
        "current_spec": spec or {},
        "attachments": [{"name": a.get("original_name"), "kind": a.get("kind"),
                         "content": (a.get("extracted") or "")[:ATTACH_CHARS]}
                        for a in (attachments or [])],
    }
    if turn >= TURN_LIMIT:
        # §4.3 가드 2 — 무한 문답 방지. 지금까지 정보로 마감하게 한다
        payload["instruction"] = ("턴 수 상한에 도달했습니다. 더 묻지 말고 지금까지 확인된 "
                                 "내용으로 사양을 마감하세요.")

    out = _ask(INTAKE_SYSTEM, _masked(payload), model, client)
    new_spec = out.get("spec_json") if isinstance(out.get("spec_json"), dict) else (spec or {})
    coverage = out.get("coverage") if isinstance(out.get("coverage"), dict) else {}
    coverage.setdefault("confirmed_topics", [])
    coverage.setdefault("remaining_topics", [])

    gaps = rubric_gaps(new_spec)
    complete = bool(out.get("spec_complete")) and not gaps
    if bool(out.get("spec_complete")) and gaps:
        # 부족한 항목을 남은 주제로 되돌려 다음 턴에 LLM이 다시 묻게 한다
        coverage["remaining_topics"] = list(dict.fromkeys(list(coverage["remaining_topics"]) + gaps))

    return {
        "reply": str(out.get("reply") or "").strip() or "조금 더 알려주세요.",
        "spec_json": new_spec,
        "spec_complete": complete,
        "coverage": coverage,
        "gaps": gaps,
        # 지금 처리한 턴까지의 개수. `turn+1`은 아직 오지 않은 턴을 세는 셈이라
        # intake_sessions.turn_count가 실제보다 1 많게 남았다(6턴 문답이 7로 기록).
        "turn": turn,
    }


# ── 저작 · 검증 ───────────────────────────────────────────────────────────────

def _check_value(value, where: str) -> list[str]:
    if isinstance(value, bool) or value is None:
        return [] if value is None else [f"{where}: 참/거짓 값은 쓸 수 없습니다"]
    if isinstance(value, (int, float)):
        return []
    if not isinstance(value, str):
        return [f"{where}: 문자열·숫자만 쓸 수 있습니다"]
    if not value.startswith("="):
        return []
    bad = []
    used = {m.group(1).upper() for m in FUNC_RE.finditer(value)}
    for func in used - {f.upper() for f in ALLOWED_FUNCS}:
        bad.append(f"{where}: 허용되지 않은 함수 {func}")
    if "[" in value or "]" in value:
        bad.append(f"{where}: 다른 파일을 참조하는 수식은 쓸 수 없습니다")
    return bad


def normalize_workbook(wb: dict) -> dict:
    """같은 뜻의 다른 이름을 표준 키로 옮긴다. **값을 만들어내지는 않는다.**

    실측에서 모델이 드롭다운 목록을 `source` 대신 `values`로, 정수 검사를 `whole` 대신
    `integer`로 내보냈다. 회차마다 흔들리는 부분이라 프롬프트만으로는 막히지 않고, 이름이
    다를 뿐 같은 제약이라 그대로 받아 옮긴다. 여기서 걸러지지 않는 것은 검증이 막는다.
    """
    if not isinstance(wb, dict):
        return wb
    for sheet in wb.get("sheets") or []:
        if not isinstance(sheet, dict):
            continue
        for dv in sheet.get("validations") or []:
            if not isinstance(dv, dict):
                continue
            if dv.get("type") in ("integer", "int"):
                dv["type"] = "whole"
            if not dv.get("source"):
                for alias in ("values", "options", "list", "items"):
                    if isinstance(dv.get(alias), list) and dv[alias]:
                        dv["source"] = dv[alias]
                        break
    return wb


def validate_workbook(wb: dict) -> list[str]:
    """명세 §6.3 검증. 위반 목록을 돌려준다(빈 목록이면 통과).

    LLM 산출물을 그대로 파일로 만들지 않기 위한 관문이다.
    """
    problems: list[str] = []
    sheets = wb.get("sheets") if isinstance(wb, dict) else None
    if not isinstance(sheets, list) or not sheets:
        return ["시트가 없습니다"]
    if len(sheets) > MAX_SHEETS:
        problems.append(f"시트가 {len(sheets)}개입니다(최대 {MAX_SHEETS}개)")

    seen_names, total_cells = set(), 0
    for i, sheet in enumerate(sheets):
        where = f"시트 {i + 1}"
        if not isinstance(sheet, dict):
            problems.append(f"{where}: 형식이 올바르지 않습니다")
            continue
        name = str(sheet.get("name") or "").strip()
        if not name:
            problems.append(f"{where}: 시트명이 비었습니다")
        if len(name) > 31:
            problems.append(f"{where}: 시트명이 31자를 넘습니다")
        if set(name) & BAD_SHEET_CHARS:
            problems.append(f"{where}: 시트명에 쓸 수 없는 문자가 있습니다")
        if name in seen_names:
            problems.append(f"{where}: 시트명 '{name}'이 중복됩니다")
        seen_names.add(name)

        cells = sheet.get("cells") or []
        if not isinstance(cells, list):
            problems.append(f"{where}: cells 형식이 올바르지 않습니다")
            cells = []
        total_cells += len(cells)
        styles = sheet.get("styles") if isinstance(sheet.get("styles"), dict) else {}
        for key, style in styles.items():
            if not isinstance(style, dict):
                problems.append(f"{where}: 스타일 '{key}' 형식이 올바르지 않습니다")
                continue
            for attr in set(style) - set(STYLE_KEYS):
                problems.append(f"{where}: 스타일 '{key}'의 허용되지 않은 속성 {attr}")

        for cell in cells:
            if not isinstance(cell, dict) or not REF_RE.match(str(cell.get("ref") or "")):
                problems.append(f"{where}: 셀 위치 표기가 올바르지 않습니다 ({cell!r})"[:120])
                continue
            problems += _check_value(cell.get("value"), f"{where} {cell['ref']}")

        for rng in sheet.get("merges") or []:
            if not RANGE_RE.match(str(rng)) or ":" not in str(rng):
                problems.append(f"{where}: 병합 범위 표기가 올바르지 않습니다 ({rng!r})")
        fp = sheet.get("freeze_panes")
        if fp and not REF_RE.match(str(fp)):
            problems.append(f"{where}: freeze_panes 표기가 올바르지 않습니다 ({fp!r})")

        for dv in sheet.get("validations") or []:
            if not isinstance(dv, dict) or not RANGE_RE.match(str(dv.get("range") or "")):
                problems.append(f"{where}: 유효성 검사 범위 표기가 올바르지 않습니다")
                continue
            if dv.get("type") not in VALIDATION_TYPES:
                problems.append(f"{where}: 지원하지 않는 유효성 검사 종류 {dv.get('type')!r}")
            if dv.get("type") == "list" and not (dv.get("source") or []):
                problems.append(f"{where}: 드롭다운 목록이 비었습니다")

    if total_cells > MAX_CELLS:
        problems.append(f"셀이 {total_cells:,}개입니다(최대 {MAX_CELLS:,}개)")
    return problems


def author_workbook(spec: dict, model: str, attachments: list[dict] | None = None,
                    client=None, on_retry=None) -> dict:
    """사양대로 workbook_json을 저작한다. 검증을 통과한 것만 돌려준다.

    검증에 걸리면 위반 목록을 그대로 알려 다시 저작하게 한다(최대 AUTHOR_RETRY회).
    """
    client = client or _client()
    payload = {
        "spec": spec,
        "attachments": [{"name": a.get("original_name"), "kind": a.get("kind"),
                         "content": (a.get("extracted") or "")[:ATTACH_CHARS]}
                        for a in (attachments or [])],
    }
    problems: list[str] = []
    for attempt in range(AUTHOR_RETRY + 1):
        if problems:
            payload["fix_these"] = problems
            if on_retry:
                on_retry(attempt, problems)
        wb = normalize_workbook(_ask(AUTHOR_SYSTEM, _masked(payload), model, client))
        problems = validate_workbook(wb)
        if not problems:
            return wb
    raise FormGenError("양식 파일을 만들 수 없었습니다. 검증에서 걸린 항목: "
                       + "; ".join(problems[:5]))


# ── xlsx 변환 ────────────────────────────────────────────────────────────────

def materialize(wb: dict, out: Path) -> list[str]:
    """workbook_json을 xlsx로 옮긴다. 해석·보정을 하지 않는 1:1 변환이다.

    안전 기본값만 둔다 — 정의 안 된 스타일은 무서식, 너비 미지정 열은 12.
    돌려주는 것은 **낮춰서 처리한 것들의 목록**이다(빈 목록이면 그대로 옮겨졌다는 뜻).
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    notes: list[str] = []
    book = openpyxl.Workbook()
    book.remove(book.active)
    for sheet in wb["sheets"]:
        ws = book.create_sheet(str(sheet.get("name") or "양식")[:31])
        styles = sheet.get("styles") if isinstance(sheet.get("styles"), dict) else {}

        for cell in sheet.get("cells") or []:
            target = ws[cell["ref"]]
            target.value = cell.get("value")
            style = styles.get(cell.get("style")) or {}
            if not style:
                continue                                  # 정의 없는 스타일 참조 → 무서식
            if style.get("bold") or style.get("italic") or style.get("size") or style.get("color"):
                target.font = Font(bold=bool(style.get("bold")), italic=bool(style.get("italic")),
                                   size=style.get("size"),
                                   color=str(style["color"]).lstrip("#") if style.get("color") else None)
            if style.get("bg"):
                fill = str(style["bg"]).lstrip("#")
                target.fill = PatternFill("solid", start_color=fill, end_color=fill)
            if style.get("border"):
                side = Side(style=str(style["border"]))
                target.border = Border(left=side, right=side, top=side, bottom=side)
            if style.get("align"):
                target.alignment = Alignment(horizontal=str(style["align"]), vertical="center",
                                             wrap_text=True)
            if style.get("number_format"):
                target.number_format = str(style["number_format"])

        for col in sheet.get("columns") or []:
            try:
                idx = int(col.get("index"))
            except (TypeError, ValueError):
                continue
            ws.column_dimensions[get_column_letter(idx)].width = col.get("width") or 12

        for rng in sheet.get("merges") or []:
            ws.merge_cells(str(rng))
        if sheet.get("freeze_panes"):
            ws.freeze_panes = str(sheet["freeze_panes"])

        for dv in sheet.get("validations") or []:
            kind = dv.get("type")
            if kind == "list":
                items = ",".join(str(s).replace(",", " ") for s in dv.get("source") or [])
                obj = DataValidation(type="list", formula1=f'"{items}"', allow_blank=True)
            elif kind == "date":
                # 날짜 검사는 비교식을 요구한다. 실무 데이터를 걸러내지 않는 하한을 쓴다
                obj = DataValidation(type="date", operator="greaterThan",
                                     formula1="DATE(1900,1,1)", allow_blank=True)
            else:
                lo, hi = dv.get("min"), dv.get("max")
                if lo is not None and hi is not None:
                    obj = DataValidation(type=kind, operator="between", formula1=str(lo),
                                         formula2=str(hi), allow_blank=True)
                elif lo is not None:
                    obj = DataValidation(type=kind, operator="greaterThanOrEqual",
                                         formula1=str(lo), allow_blank=True)
                else:
                    # 숫자 검사인데 범위가 없다. operator·수식 없는 숫자 검사는 엑셀이
                    # 손상으로 보므로 그대로 쓸 수 없고, 상·하한을 지어내면 사용자가
                    # 요구하지 않은 제약이 양식에 박힌다. 안내 문구만 남기고 밝힌다.
                    obj = DataValidation(allow_blank=True)
                    notes.append(f"{ws.title} {dv['range']}: 숫자 범위가 정해지지 않아 "
                                 "입력 안내만 넣고 형식 검사는 걸지 않았습니다")
            obj.prompt, obj.promptTitle = dv.get("prompt") or "", "입력 형식"
            obj.error, obj.errorTitle = dv.get("error") or "", "입력 오류"
            obj.showInputMessage = bool(dv.get("prompt"))
            obj.showErrorMessage = bool(dv.get("error"))
            ws.add_data_validation(obj)
            obj.add(str(dv["range"]))

    book.save(out)
    return notes


# ── 작성기준 파생 (F1-5·F1-6) ────────────────────────────────────────────────

def derive_field_rules(spec: dict) -> list[dict]:
    """spec.fields → field_rules 행. LLM을 다시 부르지 않는 기계적 파생이다."""
    rows = []
    for f in spec.get("fields") or []:
        name = (f.get("name") or "").strip()
        if not name:
            continue
        config: dict = {"required": bool(f.get("required"))}
        if (f.get("format") or "").strip():
            config["format"] = f["format"].strip()
        if (f.get("notes") or "").strip():
            config["notes"] = f["notes"].strip()
        rows.append({
            "field_name": name,
            "rule_type": f["type"] if f.get("type") in FIELD_TYPES else "text",
            "rule_config_json": config,
        })
    return rows


def rules_for_aggregate(spec: dict) -> dict:
    """취합 엔진(F2)이 그대로 먹는 작성기준으로 바꾼다.

    `aggregate.review_stage1`의 rules는 `{헤더명: {"required": bool, "key": bool}}` 형태다.
    **키 컬럼은 넣지 않는다** — 사양에 대응 개념이 없어서 첫 항목이나 dept_code를 키로
    단정하면 없는 근거로 값을 만드는 셈이다. F2의 자동 판정에 맡기고 화면에서 바꾼다.
    """
    out = {}
    for f in spec.get("fields") or []:
        name = (f.get("name") or "").strip()
        if name:
            out[name] = {"required": bool(f.get("required"))}
    return out


# ── CLI ──────────────────────────────────────────────────────────────────────

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="AI 양식 생성 (문답 → xlsx)")
    ap.add_argument("prompt", nargs="?", help="첫 요청. 생략하면 입력받는다")
    ap.add_argument("-o", "--out", type=Path, default=Path("form.xlsx"))
    ap.add_argument("--rules-out", type=Path, help="작성기준 JSON 경로(취합의 --rules에 그대로 쓴다)")
    ap.add_argument("--model", default=None, help="기본값은 .env의 OPENAI_MODEL")
    a = ap.parse_args(argv)

    ag.load_env(Path(__file__).parent / ".env")
    import os
    model = a.model or os.environ.get("OPENAI_MODEL", "gpt-5-mini")

    first = a.prompt or input("어떤 양식이 필요하신가요? > ").strip()
    if not first:
        print("요청 내용이 없습니다.")
        return 1

    messages: list[dict] = [{"role": "user", "content": first}]
    spec: dict = {}
    try:
        while True:
            turn = intake_turn(messages, spec, model)
            spec = turn["spec_json"]
            messages.append({"role": "assistant", "content": turn["reply"]})
            print(f"\nAI: {turn['reply']}")
            if turn["spec_complete"]:
                break
            if turn["gaps"]:
                print(f"  (아직 필요한 것: {', '.join(turn['gaps'])})")
            answer = input("\n나: ").strip()
            if not answer:
                print("답변이 없어 문답을 멈춥니다.")
                return 1
            messages.append({"role": "user", "content": answer})

        print(f"\n사양 확정: {spec.get('form_title')} · 항목 {len(spec.get('fields') or [])}개")
        print("양식 파일을 만드는 중…")
        wb = author_workbook(spec, model,
                             on_retry=lambda n, p: print(f"  · 검증 미통과 {len(p)}건 → 재저작 {n}회차"))
        notes = materialize(wb, a.out)
        print(f"결과: {a.out}")
        for n in notes:
            print(f"  · 낮춰서 처리: {n}")

        rules = rules_for_aggregate(spec)
        if a.rules_out:
            a.rules_out.write_text(json.dumps(rules, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"작성기준: {a.rules_out}  (취합 시 --rules {a.rules_out})")
        else:
            print("작성기준:", json.dumps(rules, ensure_ascii=False))
    except FormGenError as exc:
        print(f"\n실패: {exc}", file=sys.stderr)
        return 1
    except (KeyboardInterrupt, EOFError):
        print("\n중단했습니다.")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
