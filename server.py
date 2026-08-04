#!/usr/bin/env python3
"""F2 엑셀 취합 웹 서버.

aggregate.py(CLI 엔진)를 수정 없이 import해 브라우저 흐름에 얹는다:
  업로드 → 검토(1단계 + 파일 단위 게이팅 후 2단계) → 취합 → 다운로드

세션은 모듈 전역 dict. 서버를 재시작하면 사라지는 것이 정상이다(임시 작업 도구).
실행: uvicorn server:app
"""

from __future__ import annotations

import os
import re
import shutil
import tempfile
import threading
import time
import uuid
from pathlib import Path
from urllib.parse import quote

from fastapi import (Body, Cookie, Depends, FastAPI, File, Form, HTTPException, Response,
                     UploadFile)
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

import aggregate as ag
import auth
import formgen as fg
import hwpx_merge as hm
import store

BASE = Path(__file__).parent
# 업로드 상한은 `_limits()`가 정책에서 읽는다(기본값은 store.DEFAULT_LIMITS) — Admin에서
# 바꿀 수 있어야 하고, 화면 문구와 검사가 같은 값을 봐야 한다(종전에는 상수 하나와 화면
# 세 곳의 문구가 따로 적혀 있어 상한을 바꾸면 화면이 거짓말을 하게 돼 있었다).
RETENTION_DEFAULT_DAYS = 90
SESSION_TTL_H = 6           # 이 시간이 지난 세션의 임시 폴더는 지운다 (F3-5)
MAX_ATTACHMENTS = 5         # 문답 세션당 첨부 상한 (인지 J6)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HWPX_MIME = "application/hwp+zip"

# 업로드 허용 형식 → (거부 메시지에 붙일 안내, 결과 MIME)
FORMATS = {
    "xlsx": ("엑셀에서 'Excel 통합 문서(*.xlsx)'로 저장 후 재시도해주세요.", XLSX_MIME),
    "hwpx": ("한글에서 '한글 문서(*.hwpx)'로 저장 후 재시도해주세요. "
             "구버전 .hwp(바이너리)는 지원하지 않습니다.", HWPX_MIME),
}

ag.load_env(BASE / ".env")
app = FastAPI(title="엑셀 취합")


def _auth_disabled() -> bool:
    """명시적으로 끈 경우에만 인증을 건너뛴다.

    키가 없으면 조용히 열리는 것이 아니라 503으로 닫힌다 — 설정 누락이
    접근 제어 구멍이 되지 않도록 fail-closed로 둔다.
    """
    return os.environ.get("AUTH_DISABLED", "").lower() in ("1", "true", "yes")


if _auth_disabled():
    print("경고: AUTH_DISABLED=1 — 인증 없이 실행합니다. 로컬 개발·테스트 전용입니다.")


def require_user(access_token: str | None = Cookie(default=None, alias=auth.ACCESS_COOKIE)) -> dict:
    if _auth_disabled():
        return {"id": None, "employee_no": "(인증 비활성)", "role": "user", "status": "active"}
    return auth.current_user(access_token)


User = Depends(require_user)

# 세션: sid → {"dir": 임시폴더, "files": [UploadedFile], "bytes": 누적 바이트}
SESSIONS: dict[str, dict] = {}
# 잡: jid → 진행 상황. 세션과 마찬가지로 프로세스 메모리이며 재시작하면 사라진다
JOBS: dict[str, dict] = {}


def _apply_dept_names(session: dict) -> None:
    """사용자가 정한 부서명을 파일에 다시 씌운다.

    재검토·취합 때 원본을 다시 읽으면 uf.dept가 파일명으로 돌아가므로,
    파일을 만질 때마다 이걸 통과시킨다.
    """
    names = session.get("dept_names") or {}
    for fid, uf in enumerate(session["files"]):
        chosen = (names.get(str(fid)) or "").strip()
        if chosen:
            uf.dept = chosen


def _limits() -> dict:
    """업로드 상한. Supabase가 없으면(로컬 개발) 기본값으로 돈다."""
    if not store.enabled():
        return dict(store.DEFAULT_LIMITS)
    try:
        return store.get_limits()
    except Exception:
        # 정책을 못 읽는다고 업로드를 막을 이유는 없다. 기본값으로 진행한다
        return dict(store.DEFAULT_LIMITS)


def _retention_days() -> int:
    days = store.get_policy().get("retention_days")
    return int(days) if str(days).isdigit() else RETENTION_DEFAULT_DAYS


def _session(sid: str, user: dict) -> dict:
    session = SESSIONS.get(sid)
    # 남의 세션은 존재 자체를 알리지 않는다 (§6.2 사용자별 데이터 격리)
    if session is None or session["owner"] != user["id"]:
        raise HTTPException(404, "세션을 찾을 수 없습니다. 페이지를 새로고침 후 다시 시도해주세요.")
    return session


def _start_job(fn, owner) -> dict:
    """긴 작업을 백그라운드 스레드로 돌리고 즉시 job_id를 돌려준다.

    명세 규모(30개/10만 행)에서 읽기·합성이 수십 초라 동기 응답으로는
    진행 상황을 알릴 방법이 없다. fn은 report(phase, done, total)를 받아
    단계를 남기고, 클라이언트는 /api/job/{jid}를 폴링한다.
    """
    jid = uuid.uuid4().hex
    job = JOBS[jid] = {"status": "running", "phase": "준비 중", "done": 0, "total": 0,
                       "result": None, "error": None, "owner": owner}

    def report(phase: str, done: int = 0, total: int = 0) -> None:
        job.update(phase=phase, done=done, total=total)

    def run() -> None:
        try:
            job["result"] = fn(report)
            job.update(status="done", phase="완료")
        except HTTPException as exc:
            job.update(status="error", error=str(exc.detail))
        except Exception as exc:                       # 잡 안에서 죽어도 폴링으로 사유가 전달돼야 한다
            job.update(status="error", error=f"{type(exc).__name__}: {exc}")

    threading.Thread(target=run, daemon=True).start()
    return {"job_id": jid}


@app.post("/api/auth/signup")
def auth_signup(body: dict = Body(...)) -> dict:
    profile = auth.signup(body.get("employee_no", ""), body.get("password", ""),
                          body.get("reset_email", ""))
    return {"profile": profile}


@app.post("/api/auth/login")
def auth_login(response: Response, body: dict = Body(...)) -> dict:
    result = auth.login(body.get("employee_no", ""), body.get("password", ""))
    auth.set_session_cookies(response, result["token"])
    # 토큰은 httpOnly 쿠키로만 나간다. 응답 본문에는 담지 않는다
    return {"profile": result["profile"]}


@app.post("/api/auth/logout")
def auth_logout(response: Response) -> dict:
    auth.clear_session_cookies(response)
    return {"ok": True}


@app.post("/api/auth/reset-password-request")
def auth_reset(body: dict = Body(...)) -> dict:
    auth.request_password_reset(body.get("employee_no", ""))
    # 사번 존재 여부와 무관하게 같은 응답 (계정 열거 방지)
    return {"ok": True}


@app.get("/api/auth/me")
def auth_me(user: dict = User) -> dict:
    return {"profile": user, "auth_enabled": not _auth_disabled()}


def require_admin(user: dict = User) -> dict:
    """§3.2 Admin 전용. 권한이 없으면 403."""
    if not _auth_disabled() and user.get("role") != "admin":
        raise HTTPException(403, "관리자 권한이 필요합니다.")
    if not store.enabled():
        raise HTTPException(503, "Supabase가 설정되지 않아 관리자 기능을 쓸 수 없습니다.")
    return user


Admin = Depends(require_admin)


@app.get("/api/admin/dashboard")
def admin_dashboard(admin: dict = Admin) -> dict:
    return store.dashboard()


@app.get("/api/admin/accounts")
def admin_accounts(admin: dict = Admin) -> dict:
    return {"accounts": store.list_accounts()}


@app.patch("/api/admin/accounts/{user_id}")
def admin_patch_account(user_id: str, body: dict = Body(...), admin: dict = Admin) -> dict:
    if user_id == admin.get("id") and (body.get("role") == "user"
                                       or body.get("status") == "suspended"):
        # 마지막 관리자가 스스로를 잠가버리면 CLI 말고는 되돌릴 방법이 없다
        raise HTTPException(400, "본인 계정의 권한·상태는 이 화면에서 바꿀 수 없습니다.")
    row = store.patch_account(user_id, body)
    if row is None:
        raise HTTPException(404, "계정을 찾을 수 없거나 바꿀 수 있는 항목이 없습니다.")
    changed = ", ".join(f"{k}={v}" for k, v in body.items() if k in ("role", "status"))
    store.log_action(admin.get("id"), "계정 변경", f"{row['employee_no']} ({changed})")
    return {"account": row}


@app.delete("/api/admin/accounts/{user_id}")
def admin_delete_account(user_id: str, admin: dict = Admin) -> dict:
    if user_id == admin.get("id"):
        raise HTTPException(400, "본인 계정은 삭제할 수 없습니다.")
    target = next((a for a in store.list_accounts() if a["id"] == user_id), None)
    if target is None:
        raise HTTPException(404, "계정을 찾을 수 없습니다.")
    store.delete_account(user_id)
    # 계정이 사라지면 누가 지웠는지만 남는다(대상 사번은 문자열로 보존)
    store.log_action(admin.get("id"), "계정 삭제", target["employee_no"])
    return {"deleted": target["employee_no"]}


@app.put("/api/admin/policy/retention")
def admin_set_retention(body: dict = Body(...), admin: dict = Admin) -> dict:
    days = body.get("retention_days")
    if not isinstance(days, int) or not 1 <= days <= 3650:
        raise HTTPException(400, "보관 기간은 1~3650일 사이의 정수로 입력해주세요.")
    policy = store.set_policy("retention_days", days, admin.get("id"))
    store.log_action(admin.get("id"), "정책 변경", f"retention_days={days}")
    return {"policy": policy}


@app.put("/api/admin/policy/limits")
def admin_set_limits(body: dict = Body(...), admin: dict = Admin) -> dict:
    """업로드 상한 변경. 화면·검사가 같은 값을 보게 정책 한 곳에 둔다."""
    bounds = {"max_file_mb": (1, 500), "max_session_files": (1, 500),
              "max_session_mb": (1, 5000)}
    changed = {}
    for key, (lo, hi) in bounds.items():
        if key not in (body or {}):
            continue
        value = body[key]
        if not isinstance(value, int) or not lo <= value <= hi:
            raise HTTPException(400, f"{key}는 {lo}~{hi} 사이의 정수로 입력해주세요.")
        changed[key] = value
    if not changed:
        raise HTTPException(400, "바꿀 항목이 없습니다.")
    for key, value in changed.items():
        store.set_policy(key, value, admin.get("id"))
    store.log_action(admin.get("id"), "정책 변경",
                     ", ".join(f"{k}={v}" for k, v in changed.items()))
    return {"limits": store.get_limits()}


@app.get("/api/admin/retention/preview")
def admin_retention_preview(admin: dict = Admin) -> dict:
    """무엇이 지워질지 먼저 보여준다. 지우기 전에 세어 보이지 않으면 되돌릴 수 없다."""
    return store.expired_projects(_retention_days())


@app.post("/api/admin/retention/purge")
def admin_retention_purge(admin: dict = Admin) -> dict:
    """F4-3 보관 기간 정리 — 기간이 지난 작업과 그 파일을 실제로 지운다.

    자동 실행(타이머)은 넣지 않았다. 이 서버는 세션·잡을 프로세스 메모리에 두는 도구라
    재시작이 잦고, 타이머로 무인 삭제를 돌리면 언제 무엇이 지워졌는지 아무도 모른다.
    주기 실행이 필요하면 cron이 `python3 manage_users.py purge-expired`를 부르면 된다.
    """
    days = _retention_days()
    plan = store.expired_projects(days)
    ids = [p["id"] for p in plan["projects"]]
    store.purge_projects(ids)
    store.log_action(admin.get("id"), "보관 기간 정리",
                     f"기준 {days}일 · 작업 {len(ids)}건 · 파일 {plan['files']}건")
    return {"purged": len(ids), "files": plan["files"], "bytes": plan["bytes"],
            "kept_with_templates": plan["kept_with_templates"], "days": days}


@app.get("/api/admin/logs")
def admin_logs(q: str = "", admin: dict = Admin) -> dict:
    return {"logs": store.list_logs(q=q)}


@app.get("/api/limits")
def api_limits(user: dict = User) -> dict:
    """업로드 상한. 화면이 문구를 여기서 채운다(세 곳에 따로 적어 두면 어긋난다)."""
    return _limits()


@app.get("/api/history")
def history(q: str = "", type: str = "", user: dict = User) -> dict:
    """F4-2 이력. 본인 프로젝트만 돌려준다 (§6.2).

    type은 design.md 화면 9의 유형 필터(전체/엑셀/hwpx)에 대응한다.
    """
    if not store.enabled() or not user["id"]:
        return {"projects": []}
    kind = {"excel": "excel", "hwpx": "hwpx"}.get(type, "")
    return {"projects": store.list_projects(user["id"], q=q, kind=kind)}


@app.get("/api/history/{project_id}")
def history_detail(project_id: str, user: dict = User) -> dict:
    if not store.enabled() or not user["id"]:
        raise HTTPException(404, "이력을 찾을 수 없습니다.")
    detail = store.project_detail(user["id"], project_id)
    if detail is None:
        # 남의 프로젝트는 존재 자체를 알리지 않는다
        raise HTTPException(404, "이력을 찾을 수 없습니다.")
    return detail


@app.get("/api/history/{project_id}/download/{job_id}")
def history_download(project_id: str, job_id: str, user: dict = User):
    """이력에서 과거 취합 결과를 내려받는다. 세션이 사라진 뒤에도 동작한다."""
    detail = history_detail(project_id, user)
    job = next((j for j in detail.get("aggregation_jobs", []) if j["id"] == job_id), None)
    if job is None or not job.get("result_url"):
        raise HTTPException(404, "결과 파일이 없습니다.")
    data = store.get_object(store.RESULT_BUCKET, job["result_url"])
    # 확장자는 저장된 결과 파일을 따라간다 — 엑셀 취합은 xlsx, 한글 병합은 hwpx다
    ext = Path(job["result_url"]).suffix.lstrip(".").lower() or "xlsx"
    return Response(content=data, media_type=FORMATS.get(ext, (None, XLSX_MIME))[1],
                    headers={"Content-Disposition": f'attachment; filename="merged.{ext}"'})


@app.get("/api/job/{jid}")
def job_status(jid: str, user: dict = User) -> dict:
    job = JOBS.get(jid)
    if job is None or job["owner"] != user["id"]:
        raise HTTPException(404, "작업을 찾을 수 없습니다. 페이지를 새로고침 후 다시 시도해주세요.")
    return {k: v for k, v in job.items() if k != "owner"}


def _file_view(fid: int, uf, siblings: list[str] | None = None) -> dict:
    return {
        "fid": fid,
        "name": uf.name,
        "dept": uf.dept,
        # 파일명에서 뽑은 부서명 후보. 사용자가 고르거나 직접 입력한다(§9)
        "dept_candidates": ag.suggest_dept_names(uf.path.stem, siblings),
        "readable": uf.readable,
        # §2.1 읽기 실패 사유는 read_file이 남긴 첫 이슈의 사유 그대로
        "reason": (uf.issues[0].reason if not uf.readable and uf.issues else None),
        # headers는 UI의 키 컬럼 선택지 재료
        "sheets": [{"name": s.name, "header_row": s.header_row, "rows": len(s.rows),
                    "cols": len(s.headers), "images": len(s.images), "headers": s.headers}
                   for s in uf.sheets],
    }


def _sweep_sessions() -> None:
    """오래된 세션의 임시 폴더를 지운다 (F3-5 임시 파일 자동 정리).

    세션 생성 때마다 훑는다 — 주기 작업을 따로 돌리지 않아도 새 작업이 시작될 때
    청소되고, 아무도 안 쓰면 지울 것도 없다. 엑셀·한글 세션 모두에 적용된다.
    """
    cutoff = time.time() - SESSION_TTL_H * 3600
    for sid, session in list(SESSIONS.items()):
        if session.get("created", 0) < cutoff:
            shutil.rmtree(session["dir"], ignore_errors=True)
            SESSIONS.pop(sid, None)


@app.post("/api/session")
def create_session(body: dict = Body(default={}), user: dict = User) -> dict:
    """작업 1건 = 프로젝트 1개. 사용자가 고르는 UI는 아직 없어 자동 생성한다.

    kind는 엑셀 취합(excel, 기본)과 한글 병합(hwpx)을 가른다.
    """
    _sweep_sessions()
    kind = "hwpx" if (body or {}).get("kind") == "hwpx" else "excel"
    sid = uuid.uuid4().hex
    session = SESSIONS[sid] = {"dir": Path(tempfile.mkdtemp(prefix="agg-")), "files": [],
                               "bytes": 0, "owner": user["id"], "project_id": None,
                               "file_ids": {}, "dept_names": {}, "kind": kind,
                               "created": time.time()}
    if store.enabled() and user["id"]:
        # 이름에 시각을 넣지 않는다 — 서버 시간대로 굳어버려 DB의 created_at(UTC)을
        # 보는 사람 시간대로 변환한 값과 어긋난다. 시각은 created_at만 쓴다.
        session["project_id"] = store.create_project(
            user["id"], "한글 병합" if kind == "hwpx" else "엑셀 취합")
    return {"sid": sid, "project_id": session["project_id"], "kind": kind}


async def _save_uploads(session: dict, files: list[UploadFile], ext: str) -> list[Path]:
    """확장자·용량·개수 상한(§6)을 확인하고 세션 폴더에 저장한다.

    거부 사유는 즉시 400으로 돌려줘야 하므로 잡으로 넘기지 않고 여기서 검사한다.
    엑셀 취합(F2)과 한글 병합(F3-6)이 같은 상한을 쓴다.
    """
    hint = FORMATS[ext][0]
    limits = _limits()
    pending: list[Path] = []
    for up in files:
        name = Path(up.filename or "").name
        if not name.lower().endswith("." + ext):
            raise HTTPException(400, f"'{name}'은(는) .{ext} 파일이 아닙니다. {hint}")
        data = await up.read()
        if len(data) > limits["max_file_mb"] * 1024 * 1024:
            raise HTTPException(400, f"'{name}'의 용량이 {limits['max_file_mb']}MB를 초과합니다.")
        if len(session["files"]) + len(pending) + 1 > limits["max_session_files"]:
            raise HTTPException(400, f"업로드 가능한 파일은 최대 {limits['max_session_files']}개입니다.")
        if session["bytes"] + len(data) > limits["max_session_mb"] * 1024 * 1024:
            raise HTTPException(400, f"업로드 총 용량이 {limits['max_session_mb']}MB를 초과합니다.")
        path = session["dir"] / name
        if path.exists():
            # 같은 이름이 또 오면 덮어쓰지 않는다 — 병합 순서에 둘 다 들어갈 수 있다
            path = session["dir"] / f"{path.stem}_{len(session['files']) + len(pending) + 1}{path.suffix}"
        path.write_bytes(data)
        session["bytes"] += len(data)
        pending.append(path)
    return pending


@app.post("/api/session/{sid}/files")
async def upload_files(sid: str, files: list[UploadFile] = File(...), user: dict = User) -> dict:
    """저장까지만 동기로 하고, 느린 구조 인식은 잡으로 넘긴다."""
    session = _session(sid, user)
    pending = await _save_uploads(session, files, "xlsx")

    def work(report):
        for i, path in enumerate(pending):
            report(f"{path.name} 구조를 읽는 중", i, len(pending))
            uf = ag.read_file(path)
            session["files"].append(uf)
            if session["project_id"]:
                # 파일 자체를 Storage에 남겨야 재시작 후에도 원본이 살아 있다
                row = store.put_upload(session["project_id"], session["owner"], path, path.name)
                session["file_ids"][len(session["files"]) - 1] = row["id"]
        report("구조 인식 완료", len(pending), len(pending))
        # 부서명 후보는 같이 올라온 파일명을 함께 봐야 정확해진다(공통어 = 제목)
        stems = [f.path.stem for f in session["files"]]
        _apply_dept_names(session)
        # fid는 세션 내 인덱스이므로 누적된 전체 목록을 돌려준다
        return {"files": [_file_view(i, uf, stems) for i, uf in enumerate(session["files"])]}

    return _start_job(work, user["id"])


@app.post("/api/session/{sid}/review")
def review(sid: str, body: dict = Body(default={}), user: dict = User) -> dict:
    session = _session(sid, user)
    no_ai = bool((body or {}).get("no_ai"))
    model = os.environ.get("OPENAI_MODEL", "gpt-5-mini")
    # 키 컬럼·선택 입력 컬럼을 사용자가 고르면 엔진의 rules 경로로 넘긴다
    key_col = (body or {}).get("key_col")
    rules: dict = {key_col: {"key": True}} if key_col else {}
    # 비고·특이사항처럼 비워둘 수 있는 칸. 지정이 없으면 엔진 기본값(전 컬럼 필수)
    for col in (body or {}).get("optional_cols") or []:
        rules.setdefault(col, {})["required"] = False
    # 부서명은 화면에서 고르거나 직접 입력한 값을 쓴다. 파일명은 기본값일 뿐이다
    if isinstance((body or {}).get("dept_names"), dict):
        session["dept_names"] = {str(k): v for k, v in body["dept_names"].items()}
    # 취합할 표를 골랐으면 그것만 읽는다. 실제 양식에는 작성 가이드·비워 둔 대장처럼
    # 취합 대상이 아닌 시트가 섞여 있고, 어느 것이 데이터인지는 담당자가 안다
    # 생략하면 지난 선택을 유지하고, null을 주면 전체로 되돌린다
    if "sheets" in (body or {}):
        picked = body["sheets"]
        session["sheets"] = sorted(str(s) for s in picked) if isinstance(picked, list) else None

    def work(report):
        total = len(session["files"])
        _apply_dept_names(session)
        # 재검토 시 판정이 누적되지 않도록 초기화한다. 취합이 한 번이라도 돌았으면
        # preprocess가 셀 값을 고쳐놓았으므로 그때만 원본에서 다시 읽는다
        # (파일당 재파싱이 검토 시간의 절반을 차지한다).
        # 표 선택이 바뀐 경우도 다시 읽어야 한다 — 뺐던 표를 되살릴 근거가 파일뿐이다.
        selected = session.get("sheets")
        if session.get("preprocessed") or selected != session.get("loaded_sheets"):
            fresh = []
            for i, uf in enumerate(session["files"]):
                report(f"{uf.name} 원본을 다시 읽는 중", i, total)
                fresh.append(ag.read_file(uf.path,
                                          include_sheets=set(selected) if selected else None))
            session["files"] = fresh
            session["preprocessed"] = False
            session["loaded_sheets"] = selected
            _apply_dept_names(session)          # 다시 읽으면 dept가 파일명으로 돌아간다
        else:
            for uf in session["files"]:
                # 읽지 못한 파일은 read_file이 남긴 실패 사유가 유일한 기록이다. 지우면
                # 1단계도 건너뛰므로 이슈 0건이 되어 화면에 '정상'으로 뜬다(실제로 그랬다).
                if not uf.readable:
                    continue
                uf.issues.clear()
                uf.fixes.clear()
                uf.ai_unverified = False
                for sheet in uf.sheets:
                    sheet.col_types = []

        # 1단계는 전부 로컬 계산이라 빠르다(명세 규모에서 0.2초)
        for fid, uf in enumerate(session["files"]):
            report(f"{uf.name} 규칙 검증 중", fid, total)
            if uf.readable:
                ag.review_stage1(uf, rules)

        # 2단계는 파일당 십수 초 걸리는 API 호출이라 한꺼번에 병렬로 돌린다.
        # 게이팅(§4.2)은 review_stage2_many가 파일 단위로 적용한다.
        if not no_ai:
            gated = [uf for uf in session["files"] if uf.readable and not uf.issues]
            if gated:
                done = 0

                def ai_progress(uf):
                    nonlocal done
                    done += 1
                    report(f"AI 재검증 중 ({done}/{len(gated)})", done, len(gated))

                report("AI 재검증 중", 0, len(gated))
                ag.review_stage2_many(session["files"], model, on_done=ai_progress)

        out, columns = [], []
        for fid, uf in enumerate(session["files"]):
            for sheet in uf.sheets:
                columns += [h for h in sheet.headers if h]
            file_id = session["file_ids"].get(fid)
            if file_id:
                store.save_review(file_id, uf.grade,
                                  [{"sheet": i.sheet, "cell": i.cell, "column": i.attr,
                                    "kind": i.kind, "stage": i.stage, "grade": i.grade,
                                    "reason": i.reason} for i in uf.issues])
            out.append({
                "fid": fid,
                "name": uf.name,
                "status": uf.status,                          # §6.1 정상/이상 이진 표기
                "ai_unverified": uf.ai_unverified,            # §6.4 정상(AI 미검증)
                "default_checked": uf.grade != ag.ERROR,      # §7 내부 등급 오류만 기본 해제
                "selectable": bool(uf.readable and uf.sheets),
                "issues": [{"sheet": i.sheet, "cell": i.cell, "column": i.attr,
                            "message": f"{i.tag} {i.reason}"} for i in uf.issues],
            })
        report("검토 완료", total, total)
        return {"files": out, "columns": list(dict.fromkeys(columns))}

    return _start_job(work, user["id"])


@app.post("/api/session/{sid}/aggregate")
def aggregate(sid: str, body: dict = Body(default={}), user: dict = User) -> dict:
    session = _session(sid, user)
    if isinstance(body.get("dept_names"), dict):
        session["dept_names"] = {str(k): v for k, v in body["dept_names"].items()}
    _apply_dept_names(session)
    files = session["files"]
    included = set(body.get("included") or [])
    # included는 사용자의 최종 선택이므로 등급으로 재필터링하지 않는다 (§7 강제 포함)
    picked = [i for i, uf in enumerate(files) if i in included and uf.readable and uf.sheets]
    selected = [files[i] for i in picked]

    if not selected:
        raise HTTPException(400, "취합 가능한 파일이 없습니다.")

    mode = body.get("mode", "B")
    job_id = None
    if session["project_id"]:
        job_id = store.create_job(session["project_id"], mode,
                                  [session["file_ids"][i] for i in picked if i in session["file_ids"]],
                                  body.get("summary_cols") or [])
        session["job_id"] = job_id

    def work(report_fn):
        steps = 4

        def step(label: str, done: int) -> None:
            report_fn(label, done, steps)
            if job_id:
                store.update_job(job_id, progress=int(done / steps * 100))

        try:
            step("오류 리포트를 쓰는 중", 0)
            report_path = session["dir"] / "error_report.xlsx"
            ag.write_report(files, report_path)
            session["report"] = report_path

            step("자동교정하는 중", 1)
            for uf in selected:
                ag.preprocess(uf)
            session["preprocessed"] = True   # 셀 값이 바뀌었으니 재검토 시 원본을 다시 읽어야 한다

            step("시트를 합치는 중", 2)
            wb, notes = ag.synthesize(selected, mode,
                                      body.get("group_map") or {}, body.get("summary_cols") or [])

            step("결과 파일을 저장하는 중", 3)
            result = session["dir"] / "merged.xlsx"
            wb.save(result)
            session["result"] = result
        except Exception as exc:
            # 실패도 기록에 남겨야 이력에서 무슨 일이 있었는지 알 수 있다
            if job_id:
                store.update_job(job_id, status="failed", error_message=f"{type(exc).__name__}: {exc}")
            raise
        step("취합 완료", steps)

        # Fix에는 컬럼명이 없으므로 같은 셀의 이슈에서 작성기준을 끌어온다
        attr_of = {(i.file, i.sheet, i.cell): i.attr for uf in selected for i in uf.issues}
        payload = {
            "metrics": {
                "aggregated": len(selected),
                "sheets": len(wb.sheetnames),
                "autofixed": sum(len(uf.fixes) for uf in selected),
                "excluded": len(files) - len(selected),
            },
            "result_sheets": wb.sheetnames,
            "excluded_files": [{"name": uf.name, "issue_count": len(uf.issues)}
                               for i, uf in enumerate(files) if i not in set(picked)],
            "autofixes": [{"sheet": f.sheet, "cell": f.cell,
                           "column": attr_of.get((f.file, f.sheet, f.cell), ""),
                           "before": f.original, "after": f.corrected}
                          for uf in selected for f in uf.fixes],
            # §10.1 강제 포함한 오류 등급 파일의 오류 이슈 — 다운로드는 막지 않는다
            "forced_notice": [i.line() for uf in selected if uf.grade == ag.ERROR
                              for i in uf.issues if i.grade == ag.ERROR],
            "notes": list(dict.fromkeys(notes)),
        }
        if job_id:
            result_url = store.put_result(session["project_id"], job_id, result, "merged")
            store.put_result(session["project_id"], job_id, report_path, "report")
            store.update_job(job_id, status="done", progress=100, result_url=result_url,
                             stats_json=payload["metrics"])
            store.log_action(session["owner"], "취합 완료",
                            f"모드 {mode} · {len(selected)}건")
        return payload

    return _start_job(work, user["id"])


# ── F1 AI 양식 생성 ───────────────────────────────────────────────────────────
# 문답 상태는 F2·F3 세션과 달리 DB(intake_sessions)에 둔다. 업로드 파일은 Storage에
# 남아 다시 시작할 수 있지만, 여러 턴에 걸쳐 사용자가 답한 내용은 되살릴 근거가 없다.

def _require_consent(user: dict) -> None:
    """AI 전송 동의 없이는 F1을 쓸 수 없다 (명세 §9.1). 입력이 외부 API로 나간다."""
    if _auth_disabled() or not user.get("id"):
        return
    if not store.ai_consented(user["id"]):
        raise HTTPException(403, "AI_CONSENT_REQUIRED")


def _form_store_ready(user: dict) -> bool:
    return bool(store.enabled() and user.get("id"))


def _my_session(sid: str, user: dict) -> dict:
    """세션을 소유자 확인과 함께 가져온다. 남의 것은 존재를 알리지 않는다 (§6.2).

    Supabase를 못 쓰는 상태에서는 조회 자체가 성립하지 않는다. 여기서 막지 않으면
    REST 오류가 500으로 새어나간다.
    """
    if not _form_store_ready(user):
        raise HTTPException(404, "문답을 찾을 수 없습니다.")
    session = store.get_intake_session(sid)
    if session is None:
        raise HTTPException(404, "문답을 찾을 수 없습니다.")
    if user.get("id") and store.owner_of_project(session["project_id"]) != user["id"]:
        raise HTTPException(404, "문답을 찾을 수 없습니다.")
    return session


def _primary_attachment(session: dict) -> dict | None:
    """등록 대상으로 표시된 xlsx 첨부. 이 값이 있으면 세션은 인지 모드다 (F6-1)."""
    for a in session.get("attachments_json") or []:
        if a.get("primary") and a.get("structural"):
            return a
    return None


def _attachment_view(session: dict) -> list[dict]:
    """화면에 내려줄 첨부 목록. 구조 전체는 무거워서 표 이름·헤더 수만 보낸다."""
    out = []
    for a in session.get("attachments_json") or []:
        tables = [{"name": s.get("name"), "headers": s.get("headers") or []}
                  for s in (a.get("structural") or {}).get("sheets") or []]
        out.append({"id": a.get("id"), "name": a.get("original_name"), "kind": a.get("kind"),
                    "primary": bool(a.get("primary")), "note": a.get("note"), "tables": tables})
    return out


def _hwpx_reference(path: Path) -> str:
    """hwpx 첨부는 **참고 자료로만** 쓴다(인지 §6.4 — 등록 대상이 될 수 없다).

    본문 텍스트만 뽑는다. 병합 엔진의 읽기를 그대로 쓰므로 hwpx 판정 규칙이 갈라지지 않는다.
    """
    doc = hm._read(path)
    text = re.sub(r"<[^>]+>", " ", " ".join(doc["sections"]))
    return re.sub(r"\s+", " ", text).strip()[:fg.ATTACH_CHARS]


def _structure_preview(spec: dict, structural: dict) -> list[dict]:
    """구조 추출을 F1 생성 결과와 같은 모양(`sheets[].cells[]`)으로 바꾼다.

    화면의 미리보기 표를 두 벌 만들지 않기 위해서다 — 등록 양식은 셀 단위 저작물이 없으니
    헤더와 샘플 행을 그 자리에 놓는다. 파일을 고치는 것이 아니라 보여주기만 한다.
    """
    from openpyxl.utils import get_column_letter
    sheets = []
    for table in fg.selected_tables(spec, structural):
        headers = table.get("headers") or []
        cells = [{"ref": f"{get_column_letter(i + 1)}1", "value": h}
                 for i, h in enumerate(headers)]
        for r, row in enumerate(table.get("sample_rows") or [], start=2):
            for i, h in enumerate(headers):
                value = row.get(h)
                if value is not None:
                    cells.append({"ref": f"{get_column_letter(i + 1)}{r}", "value": value})
        sheets.append({"name": table.get("name"), "cells": cells})
    return sheets


@app.post("/api/form/attachments")
async def form_attachments(session_id: str = Form(""), files: list[UploadFile] = File(...),
                           user: dict = User) -> dict:
    """F1-3 참고 첨부 · F6-2 등록 대상 양식 업로드.

    xlsx는 구조를 읽어 등록 후보로 삼고(F6), hwpx는 참고 자료로만 쓴다. png는 받지 않는다.
    세션이 없으면 여기서 만든다 — 첨부가 첫 턴의 의도 판정 재료라 문답보다 먼저 온다.
    """
    _require_consent(user)
    if not store.enabled() or not user.get("id"):
        raise HTTPException(503, "양식 첨부는 Supabase 설정이 필요합니다.")
    if session_id:
        session = _my_session(session_id, user)
        if session["status"] == "closed":
            raise HTTPException(400, "이미 끝난 문답입니다. 새로 시작해주세요.")
    else:
        project_id = store.create_project(user["id"], "AI 양식 생성")
        session = store.create_intake_session(project_id, None)

    attachments = list(session.get("attachments_json") or [])
    has_primary = _primary_attachment(session) is not None
    tmp = Path(tempfile.mkdtemp(prefix="attach-"))
    try:
        for up in files:
            name = Path(up.filename or "").name
            ext = Path(name).suffix.lower()
            if ext not in (".xlsx", ".hwpx"):
                raise HTTPException(400, f"'{name}'은(는) 첨부할 수 없습니다. 등록할 양식은 "
                                         "xlsx, 참고 자료는 hwpx만 붙일 수 있습니다.")
            if len(attachments) + 1 > MAX_ATTACHMENTS:
                raise HTTPException(400, f"첨부는 최대 {MAX_ATTACHMENTS}개입니다.")
            data = await up.read()
            cap = _limits()["max_file_mb"]
            if len(data) > cap * 1024 * 1024:
                raise HTTPException(400, f"'{name}'의 용량이 {cap}MB를 초과합니다.")
            path = tmp / f"{uuid.uuid4().hex}{ext}"
            path.write_bytes(data)

            entry = {"id": uuid.uuid4().hex, "original_name": name,
                     "kind": "xlsx" if ext == ".xlsx" else "hwpx"}
            if ext == ".xlsx":
                try:
                    entry["structural"] = fg.read_structure(path)
                    if has_primary:
                        # 등록 대상은 하나뿐이다(인지 J6). 나머지는 참고로만 쓴다고 밝힌다
                        entry["note"] = "등록 대상 양식이 이미 있어 참고 자료로만 씁니다."
                    else:
                        entry["primary"], has_primary = True, True
                except fg.FormGenError as exc:
                    entry["note"] = f"양식으로 등록할 수 없어 참고 자료로만 씁니다: {exc}"
            else:
                try:
                    entry["extracted"] = _hwpx_reference(path)
                except hm.MergeError as exc:
                    raise HTTPException(400, str(exc))
                entry["note"] = "hwpx는 참고 자료로만 쓰이며 등록 대상이 될 수 없습니다."
            entry["storage_path"] = store.put_attachment(session["project_id"], path, name)
            attachments.append(entry)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    store.update_intake_session(session["id"], attachments_json=attachments)
    session["attachments_json"] = attachments
    return {"session_id": session["id"], "project_id": session["project_id"],
            "mode": "recognize" if has_primary else "generate",
            "attachments": _attachment_view(session)}


@app.post("/api/form/consent")
def form_consent(user: dict = User) -> dict:
    """AI 전송 동의는 **1회만** 기록한다 (명세 §9.1).

    종전에는 누를 때마다 다시 써서 동의 시각이 최신 클릭으로 밀리고 감사 로그가 쌓였다
    (실측에서 로그가 2건 남았다). 동의 시각은 언제 처음 동의했는지가 기록의 목적이다.
    """
    if not _auth_disabled() and user.get("id") and not store.ai_consented(user["id"]):
        store.set_ai_consent(user["id"])
        store.log_action(user["id"], "AI 전송 동의", "")
    return {"ok": True}


@app.get("/api/form/consent")
def form_consent_state(user: dict = User) -> dict:
    if _auth_disabled() or not user.get("id"):
        return {"consented": True}
    return {"consented": store.ai_consented(user["id"])}


# 리터럴 경로는 /api/form/{sid} 보다 먼저 선언해야 한다. 뒤에 두면 FastAPI가
# /api/form/templates 를 sid="templates" 로 잡아 404를 돌려준다(실제로 그랬다).
@app.get("/api/form/templates")
def form_templates(user: dict = User) -> dict:
    """생성해 둔 양식 목록. 취합 화면에서 작성기준을 고르는 재료다 (F1-6)."""
    if not store.enabled() or not user.get("id"):
        return {"templates": []}
    return {"templates": store.list_form_templates(user["id"])}


@app.get("/api/form/templates/{template_id}/rules")
def form_template_rules(template_id: str, user: dict = User) -> dict:
    """양식의 작성기준을 취합이 그대로 쓰는 형태로 돌려준다."""
    if not store.enabled() or not user.get("id"):
        raise HTTPException(404, "작성기준을 찾을 수 없습니다.")
    rules = store.rules_of_template(user["id"], template_id)
    if rules is None:
        raise HTTPException(404, "작성기준을 찾을 수 없습니다.")
    return {"rules": rules}


@app.post("/api/form/session")
def form_session(body: dict = Body(...), user: dict = User) -> dict:
    """F1-1 첫 프롬프트. 프로젝트와 문답 세션을 만들고 첫 턴을 돌린다."""
    _require_consent(user)
    message = (body or {}).get("message", "").strip()
    if not message:
        raise HTTPException(400, "어떤 양식이 필요한지 알려주세요.")
    if not store.enabled() or not user.get("id"):
        # 생성한 양식과 작성기준이 남지 않으면 배포·취합으로 이어질 수 없다
        raise HTTPException(503, "양식 생성은 Supabase 설정이 필요합니다.")

    project_id = store.create_project(user["id"], "AI 양식 생성")
    session = store.create_intake_session(project_id, {"role": "user", "content": message})
    return {"session_id": session["id"], "project_id": project_id} | _turn(session, user)


@app.post("/api/form/{sid}/messages")
def form_message(sid: str, body: dict = Body(...), user: dict = User) -> dict:
    """F1-2 문답 한 턴. 사양이 바뀌면 완료 판정도 다시 한다."""
    _require_consent(user)
    session = _my_session(sid, user)
    message = (body or {}).get("message", "").strip()
    if not message:
        raise HTTPException(400, "답변을 입력해주세요.")
    if session["status"] == "closed":
        raise HTTPException(400, "이미 생성이 끝난 문답입니다. 새로 시작해주세요.")
    session["messages_json"] = list(session["messages_json"]) + [{"role": "user", "content": message}]
    return _turn(session, user)


def _turn(session: dict, user: dict) -> dict:
    """LLM 한 턴을 돌리고 세션에 반영한다. 저장은 원문, 마스킹은 전송 때만(§9.2).

    등록 대상 xlsx가 붙어 있으면 인지 모드로 돈다(F6-1). 모드를 따로 저장하지 않는 이유는
    첨부의 `primary` 표시가 곧 그 상태라서다 — 별도 컬럼을 두면 둘이 어긋날 수 있다.
    """
    model = os.environ.get("OPENAI_MODEL", "gpt-5-mini")
    attachments = list(session.get("attachments_json") or [])
    primary = _primary_attachment(session)
    others = [a for a in attachments if a is not primary]
    try:
        if primary is not None:
            out = fg.recognize_turn(session["messages_json"], session.get("spec_json") or {},
                                    primary["structural"], model, attachments=others)
            if out["intent"] == "generate":
                # 첨부를 참고로만 쓰겠다는 뜻이다(F1-3). 등록 대상 표시를 내리면 다음 턴부터
                # 생성 모드로 돈다. 이번 턴 사양은 그대로 두되 완료로는 보지 않는다 —
                # 생성 완결성은 생성 루브릭으로 다시 판정해야 한다.
                for a in attachments:
                    a.pop("primary", None)
                store.update_intake_session(session["id"], attachments_json=attachments)
                out["spec_complete"] = False
                out["gaps"] = fg.rubric_gaps(out["spec_json"])
                primary = None
        else:
            out = fg.intake_turn(session["messages_json"], session.get("spec_json") or {},
                                 model, attachments=attachments)
    except fg.FormGenError as exc:
        # AI가 기능 자체라 건너뛸 수 없다. 사유를 그대로 전달한다(§9.4)
        raise HTTPException(503, str(exc))

    messages = list(session["messages_json"]) + [{"role": "assistant", "content": out["reply"]}]
    store.update_intake_session(
        session["id"], messages_json=messages, spec_json=out["spec_json"],
        turn_count=out["turn"], status="spec_complete" if out["spec_complete"] else "active")
    return {"reply": out["reply"], "spec_json": out["spec_json"],
            "spec_complete": out["spec_complete"], "coverage": out["coverage"],
            "gaps": out["gaps"], "messages": messages,
            "mode": "recognize" if primary is not None else "generate",
            "intent": out.get("intent"),
            "attachments": _attachment_view(session | {"attachments_json": attachments})}


@app.get("/api/form/{sid}")
def form_session_state(sid: str, user: dict = User) -> dict:
    """새로고침·재진입 복원용."""
    session = _my_session(sid, user)
    primary = _primary_attachment(session)
    return {"session_id": session["id"], "project_id": session["project_id"],
            "status": session["status"], "messages": session["messages_json"],
            "spec_json": session["spec_json"], "turn_count": session["turn_count"],
            "spec_complete": session["status"] == "spec_complete",
            "mode": "recognize" if primary is not None else "generate",
            "attachments": _attachment_view(session),
            "gaps": (fg.recognize_gaps(session["spec_json"] or {}, primary["structural"],
                                       session["messages_json"])
                     if primary is not None else fg.rubric_gaps(session["spec_json"] or {}))}


@app.post("/api/form/{sid}/register")
def form_register(sid: str, user: dict = User) -> dict:
    """F6-5 등록 확정. 원본 파일을 **다시 쓰지 않고** 그대로 양식으로 삼는다(E1).

    §5.3대로 동기 처리다 — LLM이 파일을 저작하지 않아 잡·폴링이 필요하지 않다.
    """
    _require_consent(user)
    session = _my_session(sid, user)
    primary = _primary_attachment(session)
    if primary is None:
        raise HTTPException(400, "등록할 양식 파일이 없습니다. xlsx 양식을 첨부해주세요.")
    if session["status"] == "closed":
        raise HTTPException(400, "이미 끝난 문답입니다. 새로 시작해주세요.")

    spec = session.get("spec_json") or {}
    gaps = fg.recognize_gaps(spec, primary["structural"], session["messages_json"])
    if gaps:
        # 화면의 등록 버튼은 서버 판정으로만 켜지지만, 직접 호출도 막는다
        raise HTTPException(400, "항목 정보가 아직 확정되지 않았습니다: " + ", ".join(gaps[:5]))

    # 확장자 위변조 방지 — 등록 직전에 실제 바이트를 다시 본다 (인지 §7·§9.3)
    data = store.get_object(store.FORM_BUCKET, primary["storage_path"])
    if not data.startswith(b"PK\x03\x04"):
        raise HTTPException(400, "양식 파일이 xlsx가 아닙니다(내용이 zip 형식이 아닙니다).")

    template_id = store.register_form_template(session["project_id"], session["id"], spec,
                                               primary["storage_path"])
    store.save_field_rules(template_id, fg.derive_field_rules(spec))
    store.update_intake_session(session["id"], status="closed")
    store.log_action(user.get("id"), "양식 등록 완료",
                     f"{spec.get('form_title')} · 항목 {len(spec.get('fields') or [])}개")
    tables = fg.selected_tables(spec, primary["structural"])
    return {
        "template_id": template_id,
        "form_title": spec.get("form_title"),
        "source": "recognized_external",
        "original_name": primary.get("original_name"),
        "sheets": _structure_preview(spec, primary["structural"]),
        "rules": fg.rules_for_aggregate(spec),
        "notes": [f"'{primary.get('original_name')}'을 고친 곳 없이 그대로 등록했습니다. "
                  "레이아웃·서식은 원본 그대로이며, 바꿀 수 있는 것은 작성기준입니다.",
                  f"회신받을 표: {', '.join(t.get('name') or '' for t in tables)}"]
        + ([f"이 파일에서 표 {len((primary['structural'].get('sheets') or []))}개를 찾아 "
            f"{len(tables)}개를 골랐습니다."]
           if len(primary["structural"].get("sheets") or []) > len(tables) else []),
    }


def _author_job(session: dict, spec: dict, template_id: str, version: int,
                user: dict, action: str) -> dict:
    """저작 → 검증 → xlsx → Storage → 기록. 생성(F1-4)과 수정(F1-7)이 같은 경로를 탄다."""
    model = os.environ.get("OPENAI_MODEL", "gpt-5-mini")
    tmp = Path(tempfile.mkdtemp(prefix="form-"))

    def work(report):
        report("양식 내용을 저작하는 중", 0, 3)
        try:
            wb = fg.author_workbook(spec, model,
                                    on_retry=lambda n, p: report(f"검증 미통과 → 재저작 {n}회차", 1, 3))
            report("파일로 만드는 중", 2, 3)
            out = tmp / "form.xlsx"
            notes = fg.materialize(wb, out)
            file_url = store.put_form(session["project_id"], template_id, out, version)
            if version > 1:
                store.revise_form_template(template_id, version, spec, wb, file_url)
            else:
                store.finish_form_template(template_id, wb, file_url)
            store.save_field_rules(template_id, fg.derive_field_rules(spec))
            store.update_intake_session(session["id"], status="closed")
            store.log_action(user.get("id"), action,
                            f"{spec.get('form_title')} · 항목 {len(spec.get('fields') or [])}개"
                            + (f" · v{version}" if version > 1 else ""))
        except fg.FormGenError as exc:
            if version == 1:
                store.fail_form_template(template_id, str(exc))
            # 수정 실패는 이전 버전을 살려둔다 — 쓸 수 있던 양식을 실패 상태로 덮으면
            # 사용자가 가지고 있던 것까지 잃는다
            raise HTTPException(503, str(exc))
        except Exception as exc:
            if version == 1:
                store.fail_form_template(template_id, f"{type(exc).__name__}: {exc}")
            raise
        report("생성 완료", 3, 3)
        return {
            "template_id": template_id,
            "form_title": spec.get("form_title"),
            "version": version,
            "sheets": [{"name": s.get("name"),
                        "cells": [{"ref": c.get("ref"), "value": c.get("value")}
                                  for c in (s.get("cells") or [])]}
                       for s in wb["sheets"]],
            "rules": fg.rules_for_aggregate(spec),
            "notes": notes,
        }

    return _start_job(work, user["id"])


@app.post("/api/form/{sid}/generate")
def form_generate(sid: str, user: dict = User) -> dict:
    """F1-4 양식 파일 생성. 저작·검증·변환을 잡으로 넘긴다."""
    _require_consent(user)
    session = _my_session(sid, user)
    spec = session.get("spec_json") or {}
    gaps = fg.rubric_gaps(spec)
    if gaps:
        # 화면의 생성 버튼은 서버 판정으로만 켜지지만, 직접 호출도 막는다
        raise HTTPException(400, "사양이 아직 확정되지 않았습니다: " + ", ".join(gaps[:5]))

    template_id = store.create_form_template(session["project_id"], session["id"], spec)
    return _author_job(session, spec, template_id, 1, user, "양식 생성 완료")


@app.post("/api/form/{sid}/revise")
def form_revise(sid: str, body: dict = Body(...), user: dict = User) -> dict:
    """F1-7 대화형 수정. "이 항목 빼줘"로 사양을 고쳐 **다시 저작하고** 버전을 올린다.

    새 프롬프트를 만들지 않는다 — 수정 요청은 결국 문답 한 턴이고(`intake_turn`은 이미
    "사용자가 정정하면 그 지시가 이긴다"를 지시받고 있다), 고쳐진 사양으로 저작·검증·변환을
    다시 태우면 **사양과 파일이 어긋날 수 없다.** 셀을 직접 고치는 쪽은 검증 관문을 하나 더
    만들어야 하고, 그 관문을 통과하지 못한 수정본이 사양과 다른 파일로 남는다.

    맞바꾼 것: 파일을 새로 저작하므로 버전 간 레이아웃이 조금 달라질 수 있다. 이전 버전
    파일은 버전별 경로로 Storage에 남는다.
    """
    _require_consent(user)
    session = _my_session(sid, user)
    message = (body or {}).get("message", "").strip()
    if not message:
        raise HTTPException(400, "무엇을 고칠지 알려주세요.")
    row = store.template_of_session(session["id"])
    if row is None:
        raise HTTPException(400, "수정할 양식이 없습니다. 먼저 양식을 생성해주세요.")
    if row.get("source") == "recognized_external":
        # 인지 명세 E3 — 등록 원본은 재저작 대상이 아니다
        raise HTTPException(400, "등록된 원본 양식은 수정할 수 없습니다. "
                                 "바꿀 수 있는 것은 작성기준입니다.")

    model = os.environ.get("OPENAI_MODEL", "gpt-5-mini")
    messages = list(session["messages_json"]) + [{"role": "user", "content": message}]
    try:
        out = fg.intake_turn(messages, session.get("spec_json") or {}, model,
                             attachments=session.get("attachments_json"))
    except fg.FormGenError as exc:
        raise HTTPException(503, str(exc))
    messages = messages + [{"role": "assistant", "content": out["reply"]}]
    store.update_intake_session(session["id"], messages_json=messages,
                                spec_json=out["spec_json"], turn_count=out["turn"])
    if not out["spec_complete"]:
        # 고칠 내용이 확정되지 않았으면 다시 저작하지 않고 되묻는다. 유료 호출을 아끼는
        # 것보다, 사양이 빈 채로 만들어진 파일이 나가지 않는 것이 중요하다
        return {"question": out["reply"], "gaps": out["gaps"], "messages": messages,
                "spec_json": out["spec_json"]}
    return _author_job(session, out["spec_json"], row["id"], (row.get("version") or 1) + 1,
                       user, "양식 수정 완료") | {"question": None, "reply": out["reply"],
                                              "messages": messages}


@app.get("/api/form/template/{template_id}/download")
def form_download(template_id: str, user: dict = User):
    if not _form_store_ready(user):
        raise HTTPException(404, "양식 파일이 없습니다.")
    row = store.get_form_template(template_id)
    if row is None or not row.get("file_url"):
        raise HTTPException(404, "양식 파일이 없습니다.")
    if user.get("id") and store.owner_of_project(row["project_id"]) != user["id"]:
        raise HTTPException(404, "양식 파일이 없습니다.")
    data = store.get_object(store.RESULT_BUCKET, row["file_url"])
    title = (row.get("spec_json") or {}).get("form_title") or "form"
    return Response(content=data, media_type=XLSX_MIME, headers={
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(title)}.xlsx"})


# ── F3 한글(hwpx) 병합 ────────────────────────────────────────────────────────
# 엑셀 취합과 세션·잡·상한·인증을 공유하고, 다른 것은 검사할 확장자와 엔진뿐이다.

@app.post("/api/hwpx/session/{sid}/files")
async def hwpx_upload(sid: str, files: list[UploadFile] = File(...), user: dict = User) -> dict:
    """F3-1 다중 hwpx 업로드. 병합 가능한 파일인지 여기서 판정해 목록에 표시한다."""
    session = _session(sid, user)
    pending = await _save_uploads(session, files, "hwpx")

    def work(report):
        for i, path in enumerate(pending):
            report(f"{path.name} 확인 중", i, len(pending))
            entry = {"name": path.name, "path": path, "size": path.stat().st_size,
                     "sections": 0, "readable": True, "reason": None}
            try:
                # 열어서 hwpx인지·본문이 있는지 확인한다. 여기서 걸러야 병합 시작 후
                # 실패하지 않는다. 구역 수는 병합 결과를 미리 알려주는 재료다
                entry["sections"] = len(hm._read(path)["sections"])
            except hm.MergeError as exc:
                entry.update(readable=False, reason=str(exc).split(": ", 1)[-1])
            session["files"].append(entry)
            if session["project_id"] and entry["readable"]:
                row = store.put_upload(session["project_id"], session["owner"], path,
                                       path.name, kind="hwpx")
                session["file_ids"][len(session["files"]) - 1] = row["id"]
        report("확인 완료", len(pending), len(pending))
        return {"files": [{"fid": i, **{k: v for k, v in f.items() if k != "path"}}
                          for i, f in enumerate(session["files"])]}

    return _start_job(work, user["id"])


@app.post("/api/hwpx/session/{sid}/merge")
def hwpx_merge_start(sid: str, body: dict = Body(default={}), user: dict = User) -> dict:
    """F3-2·3·4 지정한 순서로 병합. 순서는 화면에서 드래그로 정한 fid 목록이다."""
    session = _session(sid, user)
    files = session["files"]
    order = [i for i in (body or {}).get("order") or []
             if isinstance(i, int) and 0 <= i < len(files) and files[i]["readable"]]
    order = list(dict.fromkeys(order))          # 같은 파일을 두 번 넣지 않는다
    if len(order) < 2:
        raise HTTPException(400, "병합할 파일을 2개 이상 선택해주세요.")

    job_id = None
    if session["project_id"]:
        # 합성 모드(A/B/C/D)는 엑셀 취합 개념이라 한글 병합에는 없다 → NULL로 둔다
        job_id = store.create_job(session["project_id"], None,
                                  [session["file_ids"][i] for i in order
                                   if i in session["file_ids"]], [], kind="hwpx")
        session["job_id"] = job_id

    def work(report):
        out = session["dir"] / "merged.hwpx"
        try:
            rep = hm.merge([files[i]["path"] for i in order], out, progress=report)
        except hm.MergeError as exc:
            if job_id:
                store.update_job(job_id, status="failed", error_message=str(exc))
            raise HTTPException(400, str(exc))
        if rep["결과"] != "성공":
            # 검증 미통과 — 엔진이 파일을 쓰지 않았다. 깨진 산출물을 내보내지 않는다
            # 사유 후보를 빠뜨리면 사유 없는 실패 메시지가 나간다
            if rep["옮길수없는서식"]:
                reason = ("이 문서에는 병합이 옮길 수 없는 서식 정의가 있습니다(메모·개체 등): "
                          + "; ".join(rep["옮길수없는서식"][:5]))
            else:
                reason = "; ".join(rep["dangling"][:5]) or "; ".join(rep["itemCnt불일치"][:5])
            reason = reason or "수량 대조 불일치"
            if job_id:
                store.update_job(job_id, status="failed", error_message=reason)
            raise HTTPException(400, f"병합 결과 검증에 실패해 파일을 만들지 않았습니다: {reason}")

        session["result"] = out
        report("병합 완료", 1, 1)
        payload = {
            "metrics": {
                "merged": len(order),
                "sections": sum(rep["구역"].values()),
                "size_kb": round(out.stat().st_size / 1024),
            },
            "order": [files[i]["name"] for i in order],
            "sections_by_file": rep["구역"],
            "notes": rep["한계"],
            # header에 정의가 없어 그대로 둔 참조 — 몇 건인지 밝힌다
            "untouched_refs": rep["미대응"],
        }
        if job_id:
            result_url = store.put_result(session["project_id"], job_id, out, "merged")
            store.update_job(job_id, status="done", progress=100, result_url=result_url,
                             stats_json=payload["metrics"])
            store.log_action(session["owner"], "한글 병합 완료", f"{len(order)}건")
        return payload

    return _start_job(work, user["id"])


@app.get("/api/hwpx/session/{sid}/download")
def hwpx_download(sid: str, user: dict = User):
    session = _session(sid, user)
    path = session.get("result")
    if not path or not Path(path).exists():
        raise HTTPException(404, "아직 생성되지 않은 파일입니다.")
    return FileResponse(path, filename="merged.hwpx", media_type=HWPX_MIME)


@app.get("/api/session/{sid}/download/{kind}")
def download(sid: str, kind: str, user: dict = User):
    session = _session(sid, user)
    path = session.get("result" if kind == "result" else "report")
    if not path or not Path(path).exists():
        raise HTTPException(404, "아직 생성되지 않은 파일입니다.")
    return FileResponse(path, filename=Path(path).name,
                        media_type=XLSX_MIME)


# 정적 UI는 반드시 마지막에 마운트한다(먼저 걸면 /api 라우트가 가려진다)
app.mount("/", StaticFiles(directory=BASE / "ui", html=True))
