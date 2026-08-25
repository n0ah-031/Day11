#!/usr/bin/env python3
"""aggregate.py 자체 점검. 실행: python3 test_aggregate.py (프레임워크 없음)"""

import contextlib
import io
import json
import threading
import time
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils.units import EMU_to_pixels
from PIL import Image as PILImage

import aggregate as ag
import xlsx_format as xf


def make_book(path: Path, sheets: dict, title_row: bool = True) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        if title_row:
            # §2.2 예시 구조: 1행 병합 제목 → 2행 공백 → 3행 헤더 → 4행 데이터
            ws["A1"] = f"{name} 실적 보고"
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
            ws.append([None, None, None, None])
        for row in rows:
            ws.append(row)
    wb.save(path)


def test_structure_and_stage1(tmp: Path):
    """제목행 skip + 누락/정합성/충돌을 전부(첫 위반에서 멈추지 않고) 수집한다."""
    path = tmp / "기획부.xlsx"
    make_book(path, {"실적": [
        ["사번", "부서", "예산액", "비고"],
        ["A01", "기획부", 1000, "정상 집행"],
        ["A02", "기획부", None, "누락 케이스"],       # 누락 → 오류
        ["A03", "기획부", "abc", "숫자아님"],          # 정합성 → 오류
        ["A04", "기획부", "1,200", "콤마"],            # 정합성 → 경고(자동교정)
        ["A01", "기획부", 9999, "키충돌"],             # 충돌 → 오류
    ]})
    uf = ag.read_file(path)
    assert uf.readable, "xlsx를 읽어야 한다"
    assert len(uf.sheets) == 1, uf.issues
    sheet = uf.sheets[0]
    assert sheet.header_row == 3, f"제목행(1)·공백 아님을 넘어 3행이 헤더여야 함, got {sheet.header_row}"
    assert sheet.headers == ["사번", "부서", "예산액", "비고"], sheet.headers
    assert sheet.row_numbers[0] == 4, sheet.row_numbers

    ag.review_stage1(uf, rules={})
    kinds = sorted({i.kind for i in uf.issues})
    assert "누락" in kinds and "정합성" in kinds and "충돌" in kinds, f"{kinds} / {[i.line() for i in uf.issues]}"
    assert uf.grade == ag.ERROR and uf.status == "이상"
    # 예산액 컬럼은 값 타입 분포로 numeric 판정되어야 한다
    assert sheet.col_types[2] == "numeric", sheet.col_types
    # 자동교정 이력에 원본값이 보존되어야 한다 (§8 감사 추적)
    assert any(f.original == "1,200" for f in uf.fixes), uf.fixes
    print("  ✓ 구조 인식 + 1단계 규칙 검증(전량 스캔)")


def test_default_miss_collapse(tmp: Path):
    """기본 가정 폭주 감지(9차) — 미선언 컬럼의 누락이 임계를 넘으면 안내 1건으로 접는다.

    빈 칸이 정상인 점검표(월별 매트릭스, 실측 1,051건)가 개별 누락으로 도배되는 것을
    막되, 사용자가/양식이 선언한 기준의 누락은 몇 건이든 절대 접지 않는다.
    """
    # 월별 점검표 형태: 5개 월 컬럼이 대부분 비어 있다 (40행 × 5컬럼 = 검사 200칸)
    header = ["항목", "1월", "2월", "3월", "4월", "5월"]
    rows = [header] + [[f"활동-{i:02d}", "●" if i == 1 else None, None, None, None, None]
                       for i in range(1, 41)]
    path = tmp / "총무부.xlsx"
    make_book(path, {"점검표": rows})

    # ① 규칙 미지정 → 누락 199건이 아니라 '가정' 안내 1건으로 접힌다
    uf = ag.read_file(path)
    ag.review_stage1(uf, rules={})
    folded = [i for i in uf.issues if i.kind == "가정"]
    assert len(folded) == 1, [i.line() for i in uf.issues[:5]]
    assert not any(i.reason == "필수값 누락" for i in uf.issues), \
        "접었으면 개별 누락이 남으면 안 된다"
    assert folded[0].grade == ag.WARN and "필수 가정 재검토" in folded[0].reason
    assert "1월" in folded[0].reason, folded[0].reason   # 무엇을 지정할지 알려줘야 한다
    assert uf.status == "이상", "정상으로 만들지는 않는다 — 판정은 지정 후 재검토의 몫"

    # ② 같은 파일도 명시적으로 필수라 선언하면 접지 않는다(선언은 신뢰한다)
    uf2 = ag.read_file(path)
    ag.review_stage1(uf2, rules={h: {"required": True} for h in header})
    misses = [i for i in uf2.issues if i.reason == "필수값 누락"]
    assert len(misses) == 199 and not any(i.kind == "가정" for i in uf2.issues), len(misses)

    # ③ 선택 입력으로 선언하면 이슈 자체가 없다(기존 동작 유지)
    uf3 = ag.read_file(path)
    ag.review_stage1(uf3, rules={h: {"required": False} for h in header[1:]})
    assert not uf3.issues, [i.line() for i in uf3.issues]

    # ④ 소규모 누락(하한 미만)은 지금처럼 개별 보고한다
    small = tmp / "기획부.xlsx"
    make_book(small, {"실적": [["사번", "부서", "예산액", "비고"],
                              ["A01", "기획부", 1000, "정상"],
                              ["A02", "기획부", None, "누락1"],
                              ["A03", None, 300, "누락2"],
                              ["A04", "기획부", 400, "정상"]]})
    uf4 = ag.read_file(small)
    ag.review_stage1(uf4, rules={})
    assert sum(1 for i in uf4.issues if i.reason == "필수값 누락") == 2
    assert not any(i.kind == "가정" for i in uf4.issues)

    # ⑤ 키 가정 폭주 — 분류처럼 반복되는 첫 컬럼이 키로 잘못 가정되면 행 수만큼
    #    충돌이 쏟아진다. 자동 판정 키일 때만 안내 1건으로 접고, 중복 행
    #    '자동 제거' 제안도 함께 걷는다(틀린 가정으로 행을 지우면 안 된다)
    keyed = tmp / "안전부.xlsx"
    make_book(keyed, {"점검": [["분류", "항목", "수량"]] + [
        [f"분류-{i // 10}", f"작업-{i // 2:02d}", 100 + (i % 5)] for i in range(40)]})
    uf5 = ag.read_file(keyed)
    ag.review_stage1(uf5, rules={})
    folded5 = [i for i in uf5.issues if i.kind == "가정"]
    assert len(folded5) == 1 and "키 가정 재검토" in folded5[0].reason, \
        [i.line() for i in uf5.issues[:5]]
    assert not any(i.kind in ("충돌", "중복") for i in uf5.issues)
    assert not uf5.fixes, "접었으면 중복 자동 제거 제안도 남으면 안 된다"

    # ⑥ 사용자가 키를 지정했으면 충돌이 몇 건이든 접지 않는다(선언은 신뢰한다)
    uf6 = ag.read_file(keyed)
    ag.review_stage1(uf6, rules={"분류": {"key": True}})
    assert sum(1 for i in uf6.issues if i.kind == "충돌") >= 30
    assert not any(i.kind == "가정" for i in uf6.issues)
    print("  ✓ 기본 가정 폭주 감지(필수·키 접기, 선언 존중, 소규모 비접기)")


def test_matrix_checklist_header(tmp: Path):
    """월별 점검표(●표기·계층 헤더 2단·그룹 병합)의 헤더 오인식 회귀 고정 (9차).

    실측 결함: 값이 전부 텍스트(●)라 '다음 행에 숫자' 규칙이 뒤로 밀리다가, 예산 숫자가
    우연히 나오는 자리의 윗행(실적 행)이 헤더로 확정됐다. 그 뒤 데이터 구역의 그룹 병합
    (A7:A16)을 따라 헤더 블록이 10행으로 부풀어 법령 본문이 컬럼명이 됐다.
    """
    path = tmp / "총무부.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "점검표"
    ws["A1"] = "2026년 안전보건 모니터링"
    ws.merge_cells("A1:F1")
    # 3~4행 계층 헤더: A~C는 세로 병합, D3은 가로 병합(월 상위)
    for col, name in (("A", "활동목표"), ("B", "세부항목"), ("C", "구분")):
        ws[f"{col}3"] = name
        ws.merge_cells(f"{col}3:{col}4")
    ws["D3"] = "월별"
    ws.merge_cells("D3:E3")
    ws["D4"], ws["E4"], ws["F4"] = "1월", "2월", "예산"
    # 데이터: 활동목표를 세로 병합(그룹), 계획/실적 페어, 값은 ●만. 예산 숫자는 7행에야 나온다
    ws["A5"] = "① 준법경영"
    ws.merge_cells("A5:A8")
    ws["B5"] = "이행점검"
    ws.merge_cells("B5:B6")
    ws["C5"], ws["D5"], ws["E5"] = "계획", "●", "●"
    ws["C6"] = "실적"
    ws["B7"] = "법정검사"
    ws.merge_cells("B7:B8")
    ws["C7"], ws["D7"], ws["F7"] = "계획", "●", 250
    ws["C8"], ws["D8"] = "실적", "●"
    wb.save(path)

    uf = ag.read_file(path)
    sheet = uf.sheets[0]
    assert sheet.header_row == 3, f"3행이 헤더여야 함, got {sheet.header_row}"
    assert sheet.headers == ["활동목표", "세부항목", "구분", "월별 1월", "월별 2월", "예산"], \
        sheet.headers
    assert len(sheet.rows) == 4, sheet.rows
    # 그룹 병합은 데이터 행에 채워진다
    assert all(r[0] == "① 준법경영" for r in sheet.rows), [r[0] for r in sheet.rows]
    print("  ✓ 월별 점검표 헤더 인식(●표기·2단 헤더·그룹 병합 회귀)")


def test_plan_actual_pairs(tmp: Path):
    """계획-실적 페어 검사 (9차) — 조사 기간 내 계획이 있는 달은 실적이 있어야 한다.

    기간은 파일명에서 추정하고(사용자 결정), 위반은 경고로만 보고한다(미이행은 작성
    오류가 아니라 업무 사실일 수 있다). 주관부서 컬럼이 있으면 회신 부서 행만 본다.
    """
    # 기간 추정부터 — 분기·범위·단독 월은 잡고, 연도 숫자는 월로 오인하지 않는다
    assert ag._infer_period("2026년도 1분기 성과측정_총무부.xlsx") == (1, 3)
    assert ag._infer_period("상반기 점검") == (1, 6)
    assert ag._infer_period("1~3월 실적") == (1, 3)
    assert ag._infer_period("10월 보고") == (10, 10)
    assert ag._infer_period("2026년 결과보고") is None

    header = ["순번", "항목", "주관부서", "구분", "1월", "2월", "3월", "4월"]
    rows = [header,
            [1, "점검A", "총무부", "계획", "●", "●", None, "●"],
            [2, "점검A", "총무부", "실적", "●", None, None, None],   # 2월 미이행(4월은 기간 밖)
            [3, "점검B", "안전부", "계획", "●", None, None, None],
            [4, "점검B", "안전부", "실적", None, None, None, None]]  # 남의 부서 몫 — 빈 것이 정상
    rules = {"순번": {"key": True}} | {h: {"required": False} for h in header[1:]}
    path = tmp / "1분기 점검_총무부.xlsx"
    make_book(path, {"점검표": rows})

    # ① 파일명에서 1분기(1~3월)를 추정하고, 회신 부서(총무부) 행만 검사한다
    uf = ag.read_file(path)
    ag.review_stage1(uf, rules)
    pair = [i for i in uf.issues if i.kind == "계획실적"]
    assert uf.pair_period == (1, 3), uf.pair_period
    assert [i.line() for i in uf.issues if i.kind != "계획실적"] == [], uf.issues
    assert len(pair) == 1 and pair[0].grade == ag.WARN, [i.line() for i in pair]
    assert "2월 계획이 있으나 실적이 비어" in pair[0].reason, pair[0].reason
    assert pair[0].cell.endswith(str(uf.sheets[0].row_numbers[1])), pair[0].cell

    # ② 기간을 직접 지정하면 그 값이 이긴다 (1월만 → 위반 없음)
    uf2 = ag.read_file(path)
    ag.review_stage1(uf2, rules, period=(1, 1))
    assert not [i for i in uf2.issues if i.kind == "계획실적"]
    assert uf2.pair_period == (1, 1)

    # ③ 회신 부서와 맞는 행이 없으면 전체를 검사한다(조용한 전멸 방지) — 점검B 1월도 잡힘
    other = tmp / "1분기 점검_기획부.xlsx"
    make_book(other, {"점검표": rows})
    uf3 = ag.read_file(other)
    ag.review_stage1(uf3, rules)
    assert len([i for i in uf3.issues if i.kind == "계획실적"]) == 2

    # ④ 기간을 알 수 없으면 페어 검사만 조용히 건너뛴다(다른 검사는 그대로)
    unk = tmp / "정기점검_총무부.xlsx"
    make_book(unk, {"점검표": rows})
    uf4 = ag.read_file(unk)
    ag.review_stage1(uf4, rules)
    assert uf4.pair_period is None
    assert not [i for i in uf4.issues if i.kind == "계획실적"]
    print("  ✓ 계획-실적 페어 검사(기간 추정·부서 스코프·경고 등급)")


def test_clean_file_and_gating(tmp: Path):
    """위반 없는 파일은 정상, 위반 파일은 2단계 미진입(파일 단위 게이팅)."""
    clean = tmp / "총무부.xlsx"
    make_book(clean, {"실적": [
        ["사번", "부서", "예산액"],
        ["B01", "총무부", 500],
        ["B02", "총무부", 700],
    ]})
    uf = ag.read_file(clean)
    ag.review_stage1(uf, rules={})
    assert not uf.issues, [i.line() for i in uf.issues]
    assert uf.status == "정상"

    # AI 서비스 이용 불가(키 없음) → §13.3 폴백: 정상(AI 미검증)으로 남고 이상이 되지 않는다
    ag.review_stage2(uf, model="nonexistent-model")
    assert uf.status == "정상", [i.line() for i in uf.issues]
    assert uf.ai_unverified, "장애 시 AI 미검증 플래그가 서야 한다"
    print("  ✓ 정상 파일 판정 + AI 장애 폴백(정상(AI 미검증))")


def test_rules_and_masking(tmp: Path):
    """작성기준(§5.4)이 적용되고, 외부 전송 값은 마스킹된다(§13.2)."""
    path = tmp / "인사부.xlsx"
    make_book(path, {"명단": [
        ["사번", "부서", "이메일", "연봉"],
        ["C01", "없는부서", "a@b.com", 100],
        ["C02", "인사부", "bad-email", 100],
    ]})
    rules = {
        "부서": {"type": "code_list", "master_list": ["인사부", "총무부"]},
        "이메일": {"type": "pattern"},
        "연봉": {"type": "numeric", "min": 0, "max": 1000},
    }
    uf = ag.read_file(path)
    ag.review_stage1(uf, rules)
    reasons = " ".join(i.reason for i in uf.issues)
    assert "코드표" in reasons, reasons
    assert "허용 형식" in reasons, reasons

    assert ag.mask("문의: hong@corp.kr, 010-1234-5678, 900101-1234567") == \
        "문의: [이메일], [전화번호], [주민번호]", ag.mask("hong@corp.kr")
    # AI 전송 payload에도 마스킹이 적용되어야 한다
    uf2 = ag.read_file(path)
    ag.review_stage1(uf2, rules={})
    items = ag._ai_items(uf2)
    assert all("@" not in json.dumps(it, ensure_ascii=False) for it in items), items
    print("  ✓ 작성기준 적용 + 개인정보 마스킹(전송 payload 포함)")


def test_synthesis_modes(tmp: Path):
    """모드 A/B/C/D 결과 시트 구성 (§9)."""
    a = tmp / "기획부.xlsx"
    b = tmp / "총무부.xlsx"
    for path, dept in ((a, "기획부"), (b, "총무부")):
        make_book(path, {
            "예산": [["사번", "부서", "예산액"], [f"{dept[0]}1", dept, 100], [f"{dept[0]}2", dept, 200]],
            "인원": [["사번", "부서", "인원수"], [f"{dept[0]}1", dept, 3]],
        })
    files = []
    for path in (a, b):
        uf = ag.read_file(path)
        ag.review_stage1(uf, rules={})
        files.append(uf)

    wb, _, _ = ag.synthesize(files, "A", {}, [])
    assert len(wb.sheetnames) == 5, wb.sheetnames                 # 취합 개요 1 + n×m = 2×2
    assert "기획부_예산" in wb.sheetnames, wb.sheetnames

    wb, _, _ = ag.synthesize(files, "B", {}, [])
    assert sorted(wb.sheetnames) == ["예산", "인원", "취합 개요"], wb.sheetnames  # m = 2 + 취합 개요
    ws = wb["예산"]
    # 헤더는 1행이 아니라 원본 자리(3행: 제목1 + 공백2 다음)로 복원된다
    assert ws.cell(row=3, column=1).value == "부서", "부서 구분 컬럼이 추가돼야 한다"
    assert ws.max_row == 7, ws.max_row                # 제목 2 + 헤더 1 + 데이터 (2+2) 행

    wb, notes, _ = ag.synthesize(files, "C", {"예산": "재정"}, [])
    assert sorted(wb.sheetnames) == ["미분류", "재정", "취합 개요"], wb.sheetnames  # k = 2 + 취합 개요, 인원→미분류
    assert any("미분류" in n for n in notes), notes
    assert wb["재정"].cell(row=3, column=1).value == "구분", "원 시트명 컬럼이 추가돼야 한다"

    wb, _, _ = ag.synthesize(files, "D", {}, ["예산액"])
    assert wb.sheetnames[:2] == ["취합 개요", "종합요약"], wb.sheetnames  # 취합 개요가 맨 앞, 종합요약이 그 다음
    formula = wb["종합요약"].cell(row=2, column=2).value
    assert isinstance(formula, str) and formula.startswith("=SUMIF"), formula  # 하드코딩 금지
    print("  ✓ 합성 모드 A/B/C/D 시트 구성 + 모드 D 수식 기반 요약")


def test_image_anchor_relocation(tmp: Path):
    """합성 시 이미지가 원래 붙어 있던 '그 행 · 그 컬럼'을 따라가야 한다 (§9)."""
    def make_with_image(path: Path, dept: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "실적"
        ws["A1"] = f"{dept} 실적 보고"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
        for c, h in enumerate(["사번", "부서", "예산액", "증빙사진"], start=1):
            ws.cell(row=3, column=c, value=h)
        for i in range(2):
            ws.cell(row=4 + i, column=1, value=f"{dept[0]}{i + 1}")
            ws.cell(row=4 + i, column=2, value=dept)
            ws.cell(row=4 + i, column=3, value=100 * (i + 1))
        buf = io.BytesIO()
        PILImage.new("RGB", (300, 300), (10, 20, 30)).save(buf, format="PNG")
        buf.seek(0)
        picture = XLImage(buf)
        picture.width, picture.height = 96, 72      # 엑셀에서 셀 크기에 맞춰 줄여 놓은 상태
        ws.add_image(picture, "D4")                 # 데이터 첫 행의 증빙사진 칸
        wb.save(path)

    def placed(wb, tag):
        """anchor는 저장 시점에 확정되므로 저장 후 다시 읽어 실제 위치를 본다."""
        out = tmp / f"{tag}.xlsx"
        wb.save(out)
        found = {}
        for ws in openpyxl.load_workbook(out).worksheets:
            for im in getattr(ws, "_images", []):
                frm = im.anchor._from
                size = getattr(im.anchor, "ext", None)
                found.setdefault(ws.title, []).append((
                    frm.row + 1, frm.col + 1,
                    (EMU_to_pixels(size.cx), EMU_to_pixels(size.cy)) if size else None))
        return found

    files = []
    for dept in ("기획부", "총무부"):
        path = tmp / f"{dept}.xlsx"
        make_with_image(path, dept)
        uf = ag.read_file(path)
        ag.review_stage1(uf, rules={})
        files.append(uf)
    assert files[0].sheets[0].images, "원본 이미지를 읽어야 한다"

    # 모드 A: 원본 서식이 복원되어 헤더가 원래 자리(3행)로 돌아가므로
    # 데이터 첫 행도 원래 자리(4행)다. 컬럼은 그대로 D(4).
    wb, _, _ = ag.synthesize(files, "A", {}, [])
    got = placed(wb, "A")
    assert got["기획부_실적"] == [(4, 4, (96, 72))], got

    # 모드 B: '부서' 컬럼이 앞에 끼므로 증빙사진은 D→E(5). 행은 파일별 시작 행.
    wb, _, _ = ag.synthesize(files, "B", {}, [])
    assert placed(wb, "B")["실적"] == [(4, 5, (96, 72)), (6, 5, (96, 72))], placed(wb, "B")

    # 모드 C: '구분'+'부서' 2개가 끼므로 D→F(6).
    wb, _, _ = ag.synthesize(files, "C", {"실적": "상반기"}, [])
    assert placed(wb, "C")["상반기"] == [(4, 6, (96, 72)), (6, 6, (96, 72))], placed(wb, "C")
    print("  ✓ 이미지 anchor 재배치(행·열 추종 + 표시 크기 보존)")


def test_real_form_structure(tmp: Path):
    """실제 업무 양식 구조에서 오탐이 없어야 한다.

    실제 회신 파일(지사별 자체점검 리스트)에서 확인된 네 가지를 그대로 재현한다.
    이 구조에서 종전 엔진은 8개 파일 전부를 '오류'로 판정해 아무것도 취합되지
    않았다.
      ① 첫 시트가 '작성 주의사항' 안내문 — 표가 아님
      ② 제목 행(병합) + 안내 행 뒤 4행에 헤더
      ③ 지사·개소가 세로 병합 — 파일에는 첫 칸만 값이 있다
      ④ 표 끝에 '합 계' 행과 '*' 각주 행
    """
    path = tmp / "판교지사.xlsx"
    wb = openpyxl.Workbook()

    guide = wb.active
    guide.title = "작성 주의사항"
    guide["A1"] = "집중안전점검 사업장 자체점검 작성 안내"
    guide.merge_cells("A1:L1")
    for r, text in enumerate(["ㅇ 해당지사 : 전 지사", "ㅇ 주의사항", "1. 매달 취합합니다"], start=3):
        guide.cell(row=r, column=1, value=text)

    ws = wb.create_sheet("2025년 대상 리스트")
    ws["A1"] = "2025년 집중안전점검 지열지점 사업장 자체점검 리스트"
    ws.merge_cells("A1:H1")
    ws["A3"] = "  ※ 2024.11.01.~2025.04.30. 기간 기준"
    for c, h in enumerate(["지사", "개소", "위치", "관경(A)", "준공연도", "온도차(℃)", "보수여부"], start=1):
        ws.cell(row=4, column=c, value=h)
    ws["A5"], ws["B5"] = "판교지사", 2
    ws.merge_cells("A5:A7")          # 지사가 세 행에 걸쳐 병합
    ws.merge_cells("B5:B7")
    for i, (loc, dia, year, temp, fix) in enumerate([
        ("백현동 555-2", 200, 2015, 24.7, "보수완료"),
        ("삼평동 672-6", 150, 2015, 3.1, "보수완료"),
        ("판교동 12-3", 300, 2016, 5.5, "보수중"),
    ]):
        ws.cell(row=5 + i, column=3, value=loc)
        ws.cell(row=5 + i, column=4, value=dia)
        ws.cell(row=5 + i, column=5, value=year)
        ws.cell(row=5 + i, column=6, value=temp)
        ws.cell(row=5 + i, column=7, value=fix)
    ws["A9"], ws["B9"] = "합 계", 2                       # 합계 행
    ws["A10"] = " *  해당기간 내 점검 완료분만 기재"        # 각주 행
    # 양식은 아래쪽까지 미리 병합해 둔다 — 여기에 값이 새면 빈 행이 데이터로 살아난다
    ws.merge_cells("A12:A30")
    ws.merge_cells("B12:B30")
    wb.save(path)

    uf = ag.read_file(path)
    assert uf.readable, uf.issues
    # 안내 시트는 취합 대상에서 빠지고, 데이터 시트만 남는다
    assert [s.name for s in uf.sheets] == ["2025년 대상 리스트"], [s.name for s in uf.sheets]
    sheet = uf.sheets[0]
    assert sheet.header_row == 4, sheet.header_row
    assert sheet.headers[:3] == ["지사", "개소", "위치"], sheet.headers
    # 합계·각주 행은 제외, 미리 병합해 둔 빈 행은 되살아나지 않는다
    assert sheet.row_numbers == [5, 6, 7], sheet.row_numbers
    # 세로 병합된 지사·개소가 세 행 모두 채워진다
    assert [r[0] for r in sheet.rows] == ["판교지사"] * 3, [r[0] for r in sheet.rows]
    assert [r[1] for r in sheet.rows] == [2, 2, 2], [r[1] for r in sheet.rows]

    ag.review_stage1(uf, rules={})
    errors = [i for i in uf.issues if i.grade == ag.ERROR]
    assert not errors, [i.line() for i in errors]
    assert uf.grade == ag.WARN, [i.line() for i in uf.issues]
    # 등급이 오류가 아니므로 취합에 기본 포함된다 (§7)
    reasons = " ".join(i.reason for i in uf.issues)
    assert "표 구조가 아니어서" in reasons and "합계·각주" in reasons, reasons

    # 첫 컬럼(지사)이 반복되어도 키 충돌로 오탐하지 않는다 — 고유한 '위치'를 키로 잡는다
    assert ag._pick_key_column(sheet.headers, sheet.rows) == 2, sheet.headers

    # 합계 행이 취합에 섞이면 모드 B/D에서 이중 계상된다
    result, _, _ = ag.synthesize([uf], "B", {}, [])
    body = result["2025년 대상 리스트"]
    assert body.max_row == 7, body.max_row          # 제목 3(제목·공백·안내) + 헤더 1 + 데이터 3
    assert all(str(body.cell(row=r, column=2).value).replace(" ", "") not in ("합계", "계")
               for r in range(5, body.max_row + 1))
    print("  ✓ 실제 업무 양식 구조(안내 시트·세로 병합·합계·각주 행)")


def test_stacked_tables_and_tiered_header(tmp: Path):
    """한 시트에 표가 둘, 헤더가 여러 단인 양식에서 오탐이 없어야 한다.

    실제 회신 파일(분기별 실적 취합, 지사 12건)에서 확인된 것을 그대로 재현한다.
    이 구조에서 종전 엔진은 12개 파일 전부를 '오류'로 판정했고 이슈가 6,884건이었다.
      ① 표가 B열부터 시작 — A열이 비어 있다
      ② 헤더가 2단: 단일 컬럼은 세로 병합, 2단 컬럼은 상위 가로 병합 + 하위 행
      ③ 빈 행 한 줄만 두고 **둘째 표**가 이어진다 → 첫 표가 그 헤더·데이터를 삼켰다
      ④ 표 끝의 '합계' 행이 B열에 있어 제외되지 않았다(모드 B/D에서 이중 계상)
      ⑤ 순번만 1..N으로 미리 매겨 둔 빈 서식 행
      ⑥ 한 행에 사진 컬럼이 둘(문제 예시·개선 예시) — 이미지 개수 정책은 필드당이다
    """
    path = tmp / "1. 강남지사.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1. 아차사고 신고 월별 세부내용"
    ws["B2"] = "아차사고 신고 월별 세부내용"
    ws.merge_cells("B2:H2")

    # ── 첫 표: 재해유형 × 월(2단 헤더). 세로 병합 1단 + 가로 병합 상위 + 하위 행
    ws["B4"] = "재해유형"
    ws.merge_cells("B4:B5")                 # 단일 컬럼 → 세로 병합
    ws["C4"] = "1월"
    ws.merge_cells("C4:D4")                 # 상위 → 가로 병합
    ws["C5"], ws["D5"] = "신고", "조치"      # 하위
    for i, kind in enumerate(["끼임", "넘어짐", "떨어짐"]):
        ws.cell(row=6 + i, column=2, value=kind)
        ws.cell(row=6 + i, column=3, value=0)
        ws.cell(row=6 + i, column=4, value=0)
    ws["B9"], ws["C9"], ws["D9"] = "합계", 0, 0      # A열이 아니라 B열의 합계

    # ── 빈 행 한 줄 뒤 둘째 표(2단 헤더). 하위 단이 있는 컬럼만 이름이 합쳐져야 한다
    ws["B11"] = "구분"
    ws.merge_cells("B11:B12")
    ws["C11"] = "위험성평가"
    ws.merge_cells("C11:D11")
    ws["C12"], ws["D12"] = "가능성", "중대성"
    for i, (kind, a, b) in enumerate([("아차사고(내부)", 2, 3), ("아차사고(외부)", 1, 2)]):
        ws.cell(row=13 + i, column=2, value=kind)
        ws.cell(row=13 + i, column=3, value=a)
        ws.cell(row=13 + i, column=4, value=b)

    # ── 사진 컬럼이 둘인 시트 + 순번만 매겨 둔 빈 서식 행
    op = wb.create_sheet("3. 근로자 의견수렴, 개선")
    op["B2"] = "근로자 제안, 개선"
    op["B4"] = "No"
    op.merge_cells("B4:B5")
    op["C4"] = "관련사진"
    op.merge_cells("C4:D4")
    op["C5"], op["D5"] = "개선 전", "개선 후"
    op["B6"] = 1
    op.cell(row=6, column=5, value="완료")
    op["E4"] = "완료 여부"
    op.merge_cells("E4:E5")
    for i in range(2, 6):                   # 순번만 있는 빈 서식 행 4개
        op.cell(row=5 + i, column=2, value=i)
    for col in ("C6", "D6"):                # 같은 행에 사진 2장(서로 다른 컬럼)
        buf = io.BytesIO()
        PILImage.new("RGB", (300, 300), (10, 20, 30)).save(buf, format="PNG")
        buf.seek(0)
        op.add_image(XLImage(buf), col)
    wb.save(path)

    uf = ag.read_file(path)
    assert uf.readable, [i.line() for i in uf.issues]
    names = [s.name for s in uf.sheets]
    assert names == ["1. 아차사고 신고 월별 세부내용 (1)",
                     "1. 아차사고 신고 월별 세부내용 (2)",
                     "3. 근로자 의견수렴, 개선"], names

    first, second, third = uf.sheets
    # ② 계층 헤더가 한 줄로 합쳐진다. 하위 단이 없는 컬럼에는 상위 이름만 붙는다
    assert first.headers == ["", "재해유형", "1월 신고", "1월 조치"], first.headers
    # ③ 둘째 표를 삼키지 않고 따로 읽는다 ④ B열의 합계 행은 제외된다
    assert first.row_numbers == [6, 7, 8], first.row_numbers
    assert second.headers == ["", "구분", "위험성평가 가능성", "위험성평가 중대성"], second.headers
    assert second.row_numbers == [13, 14], second.row_numbers
    # ⑤ 순번만 있는 행은 레코드가 아니다
    assert third.row_numbers == [6], third.row_numbers
    assert third.headers[4] == "완료 여부", third.headers

    ag.review_stage1(uf, rules={})
    errors = [i for i in uf.issues if i.grade == ag.ERROR]
    assert not errors, [i.line() for i in errors]
    reasons = " ".join(i.reason for i in uf.issues)
    assert "표 2개를 찾아" in reasons and "순번만 있고" in reasons, reasons
    # ⑥ 사진 컬럼이 둘이어도 개수 위반이 아니다(명세 §5.3은 필드당)
    assert "이미지가" not in reasons or "장 있습니다" not in reasons, reasons

    # 표를 골라 읽을 수 있다 — 취합 대상이 아닌 표는 사유도 남기지 않는다
    only = ag.read_file(path, include_sheets={"1. 아차사고 신고 월별 세부내용 (2)"})
    assert [s.name for s in only.sheets] == ["1. 아차사고 신고 월별 세부내용 (2)"], only.sheets
    assert not any("3. 근로자" in i.sheet for i in only.issues), [i.line() for i in only.issues]
    print("  ✓ 한 시트 2표 · 계층 헤더 · B열 합계 · 순번만 있는 행 · 표 선택")


def test_stage2_parallel(tmp: Path):
    """2단계는 파일 간 병렬로 돌고, 한 파일의 실패가 다른 파일을 오염시키지 않는다.

    실제 API를 부르지 않는다 — _ai_call을 지연·실패가 흉내나는 스텁으로 바꿔
    호출 스케줄만 본다(네트워크·비용 없이 결정적으로 검증).
    """
    files = []
    for i in range(6):
        path = tmp / f"제{i}부서.xlsx"
        make_book(path, {"예산": [
            ["사번", "사업내용", "예산액"],
            [f"{i}-1", "교육 운영", 100],
            [f"{i}-2", "비품 구매", 200],
        ]})
        uf = ag.read_file(path)
        ag.review_stage1(uf, rules={})
        assert not uf.issues, [x.line() for x in uf.issues]   # 전부 게이팅 통과
        files.append(uf)

    DELAY = 0.15
    calls, peak, active = [], [0], [0]
    lock = threading.Lock()

    def fake_call(client, model, items):
        with lock:
            active[0] += 1
            peak[0] = max(peak[0], active[0])
            calls.append(items[0]["id"])
        time.sleep(DELAY)
        with lock:
            active[0] -= 1
        # '제3부서'는 서비스 장애를 흉내낸다
        if any("제3부서" in str(it.get("작성기준", "")) for it in items):
            raise RuntimeError("모의 장애")
        return {it["id"]: {"verdict": "부적합", "reason": "모의 판정"} for it in items}

    real_call, real_client = ag._ai_call, ag._ai_client
    ag._ai_call = fake_call
    ag._ai_client = lambda: object()          # 클라이언트 생성은 성공한 것으로 둔다
    try:
        # 파일명이 payload에 남지 않으므로, 장애 대상 식별용으로 작성기준에 심는다
        for uf in files:
            for sheet in uf.sheets:
                sheet.headers = [f"{h}({uf.dept})" for h in sheet.headers]
                sheet.col_types = []

        t0 = time.perf_counter()
        ag.review_stage2_many(files, model="stub", workers=6)
        elapsed = time.perf_counter() - t0
    finally:
        ag._ai_call, ag._ai_client = real_call, real_client

    assert len(calls) == 6, calls
    assert peak[0] > 1, f"동시에 호출되지 않았다(최대 동시 {peak[0]})"
    # 6개를 순차로 돌면 6*DELAY, 병렬이면 그보다 확실히 짧다
    assert elapsed < 6 * DELAY * 0.7, f"병렬 이득이 없다: {elapsed:.2f}s"

    failed = [uf for uf in files if uf.ai_unverified]
    assert [uf.dept for uf in failed] == ["제3부서"], [uf.dept for uf in failed]
    assert not failed[0].issues, "장애 파일은 2단계 판정을 남기지 않는다"
    for uf in files:
        if uf.ai_unverified:
            continue
        # 정상 파일은 부적합 판정이 경고 등급으로 붙는다 (§6.2)
        assert uf.issues and all(i.grade == ag.WARN and i.stage == "2단계" for i in uf.issues), \
            [i.line() for i in uf.issues]
    print("  ✓ 2단계 병렬 호출 + 파일 단위 실패 격리")


def test_cli_end_to_end(tmp: Path):
    """CLI 전체 경로: 이상 파일은 기본 제외되고 결과·리포트가 생성된다."""
    work = tmp / "e2e"
    work.mkdir()
    make_book(work / "기획부.xlsx", {"예산": [
        ["사번", "부서", "예산액"], ["A1", "기획부", 100], ["A2", "기획부", 200]]})
    make_book(work / "총무부.xlsx", {"예산": [
        ["사번", "부서", "예산액"], ["B1", "총무부", None]]})     # 누락 → 오류 → 기본 제외
    out, report = tmp / "merged.xlsx", tmp / "report.xlsx"
    stdout = io.StringIO()
    with contextlib.redirect_stdout(stdout):
        code = ag.main([str(work), "--mode", "B", "--no-ai",
                        "--out", str(out), "--report", str(report)])
    assert code == 0, code
    assert out.exists() and report.exists()
    ws = openpyxl.load_workbook(out)["예산"]
    # 헤더가 원본 자리(3행: 제목1 + 공백2 다음)로 복원되므로 데이터는 4행부터다
    depts = {ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)}
    assert depts == {"기획부"}, f"오류 파일은 기본 제외돼야 한다: {depts}"
    # 결과 파일은 '예산' 1개 시트뿐이다 — '취합 개요'까지 세면 서버(server.py)와
    # 다르게 시트 2개라고 말하게 된다
    assert "시트 1개" in stdout.getvalue(), stdout.getvalue()
    rep = openpyxl.load_workbook(report)
    # 합성 뒤에 쓰므로 결과 행수·기준 서식을 실은 취합 개요가 맨 앞에 붙는다
    assert rep.sheetnames == ["취합 개요", "오류 목록", "자동교정 이력"], rep.sheetnames
    assert rep["오류 목록"].max_row >= 2, "누락 오류가 리포트에 기록돼야 한다"

    # 강제 포함 시에는 오류 파일도 취합에 들어간다 (§7, §10.1)
    code = ag.main([str(work), "--mode", "B", "--no-ai", "--include-anomalous",
                    "--out", str(out), "--report", str(report)])
    assert code == 0
    ws = openpyxl.load_workbook(out)["예산"]
    depts = {ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)}
    assert depts == {"기획부", "총무부"}, depts
    print("  ✓ CLI end-to-end(기본 제외 / 강제 포함)")


def test_pivot_flatten(tmp: Path):
    """총괄표류(월×지표 피벗)를 한 줄로 눕혀 지사별 1행으로 모은다.

    실제 총괄표의 두 성질을 fixture에 재현한다 — ① 시트명이 파일마다 다르다,
    ② 표가 **B열부터** 시작해 A열 헤더가 비어 있다(라벨 컬럼을 0으로 단정하면 틀린다).
    """
    import openpyxl

    def pivot_file(path: Path, sheet_name: str, base: int) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws["B1"] = f"{sheet_name} 연간 집계"          # 제목 행
        ws["B3"], ws["C3"], ws["D3"] = "구분", "신고", "조치"   # 헤더는 B열부터
        for i, month in enumerate(("1월", "2월", "3월")):
            r = 4 + i
            ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = month, base + i, base + i + 10
        ws[f"B{4 + 3}"] = "합 계"                     # 합계 행은 레코드가 아니다
        ws[f"C{4 + 3}"] = base * 3
        wb.save(path)

    a, b = tmp / "강남.xlsx", tmp / "판교.xlsx"
    pivot_file(a, "강남지사 총괄표", 1)
    pivot_file(b, "판교지사 총괄표", 5)

    # 피벗 지정 없이 읽으면 시트명이 갈리고 월이 행으로 남는다(종전 동작)
    plain = ag.read_file(a, include_sheets={"강남지사 총괄표"})
    assert plain.sheets[0].name == "강남지사 총괄표", plain.sheets[0].name
    assert len(plain.sheets[0].rows) == 3, plain.sheets[0].rows

    files = []
    for path, name, dept in ((a, "강남지사 총괄표", "강남지사"), (b, "판교지사 총괄표", "판교지사")):
        uf = ag.read_file(path, include_sheets={name}, pivot_sheets={name})
        uf.dept = dept
        ag.review_stage1(uf, {})
        files.append(uf)

    flat = files[0].sheets[0]
    assert flat.name == ag.PIVOT_SHEET_NAME, flat.name          # 한 이름으로 모인다
    assert len(flat.rows) == 1, flat.rows                        # 파일 하나가 한 행
    # 라벨 컬럼은 A열(빈 헤더)이 아니라 이름이 있는 첫 컬럼(`구분`)이다
    assert flat.headers[:4] == ["1월 신고", "1월 조치", "2월 신고", "2월 조치"], flat.headers[:4]
    assert flat.rows[0][:4] == [1, 11, 2, 12], flat.rows[0][:4]
    assert len(flat.headers) == 6, flat.headers                  # 3개월 × 2지표, 합계 행 제외

    # 시트명이 달랐던 두 파일이 모드 B에서 한 시트 2행으로 합쳐진다
    wb, _notes, run = ag.synthesize(files, "B", {}, [])
    assert wb.sheetnames == [xf.OVERVIEW_SHEET, ag.PIVOT_SHEET_NAME], wb.sheetnames
    ws = wb[ag.PIVOT_SHEET_NAME]
    assert ws.max_row == 3 and ws.max_column == 7, (ws.max_row, ws.max_column)
    assert [c.value for c in ws[1]][:3] == ["부서", "1월 신고", "1월 조치"]
    assert [ws.cell(row=r, column=1).value for r in (2, 3)] == ["강남지사", "판교지사"]
    assert [ws.cell(row=r, column=2).value for r in (2, 3)] == [1, 5]
    # 피벗은 원본에 대응 시트가 없어 기본서식으로 떨어진다 — 사고가 아니라 결정이다
    assert run.source_label == "기본서식", run.source_label

    # 같은 라벨·컬럼이 겹치면 조용히 덮지 않고 번호를 붙인다
    dup = ag.SheetData("x", 1, ["구분", "신고"], [["1월", 7], ["1월", 9]], [2, 3])
    assert ag._flatten_pivot(dup).headers == ["1월 신고", "1월 신고 (2)"]
    assert ag._flatten_pivot(dup).rows[0] == [7, 9]

    # 눕히면 사진이 어느 칸에 붙어 있었는지 말할 수 없다 → 빼고 그 사실을 알린다
    img = ag.SheetData("x", 1, ["구분", "신고"], [["1월", 7]], [2],
                       images=[{"from": (2, 2), "to": (2, 2), "name": "p.png"}])
    assert ag._flatten_pivot(img).images == []
    print("  ✓ 피벗형 전개(B열 시작·시트명 통합·합계 제외·중복 번호·사진 제외)")


def _styled_form(path: Path) -> None:
    """실제 업무 양식의 구조를 재현한다 — 1행 병합 제목 / 3행 안내 / 4행 헤더 / 5행~ 데이터.

    실측(sampledata 8개)에서 확인한 모양이다: 제목은 A1:D1 병합, 헤더는 볼드·배경색·
    테두리, 열너비는 컬럼마다 다르고, 날짜 컬럼에 표시형식이 걸려 있다.
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "대상 리스트"
    ws["A1"] = "2025년 집중안전점검 리스트"
    ws["A1"].font = Font(bold=True, size=20)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.row_dimensions[1].height = 40
    ws["A3"] = "※ 2024.11.01 기준으로 작성"
    ws["A3"].font = Font(bold=True)
    for c, name in enumerate(["지사", "개소", "점검일시", "온도차"], start=1):
        cell = ws.cell(row=4, column=c, value=name)
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill("solid", fgColor="C0C0C0")
        cell.border = Border(bottom=thin, top=thin, left=thin, right=thin)
        cell.alignment = Alignment(horizontal="center")
    ws.append(["대구", 2, "2025-01-08", 3.3])
    for c in range(1, 5):
        ws.cell(row=5, column=c).border = Border(bottom=thin, top=thin, left=thin, right=thin)
    ws.cell(row=5, column=3).number_format = "mm-dd-yy"
    ws.cell(row=5, column=2).number_format = "0_);[Red]\\(0\\)"
    for letter, width in (("A", 16.2), ("C", 30.8), ("D", 33.0)):
        ws.column_dimensions[letter].width = width
    ws.print_title_rows = "$1:$4"
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)


def test_format_capture(tmp: Path):
    """원본 양식에서 서식만 읽어낸다. 못 여는 파일은 예외가 아니라 None이다."""
    path = tmp / "양식.xlsx"
    _styled_form(path)

    tpl = xf.capture(path, "대상 리스트", 4)
    assert tpl is not None
    assert tpl.header_row == 4
    # 제목 블록은 헤더 앞 3행(제목·빈행·안내) 전부
    assert len(tpl.title_rows) == 3
    assert tpl.title_rows[0].cells[0][1] == "2025년 집중안전점검 리스트"
    assert tpl.title_rows[0].height == 40
    assert tpl.title_merges == [(1, 1, 1, 4)]
    assert tpl.title_rows[2].cells[0][1] == "※ 2024.11.01 기준으로 작성"

    # 헤더·데이터 스타일은 첫 컬럼 셀에서 뜬다
    assert tpl.header_style.font.bold is True
    assert tpl.header_style.fill.fgColor.rgb.endswith("C0C0C0")
    assert tpl.data_style.border.bottom.style == "thin"

    # 열너비·숫자서식은 컬럼 위치가 아니라 헤더 이름으로 담긴다
    assert tpl.widths == {"지사": 16.2, "점검일시": 30.8, "온도차": 33.0}
    assert tpl.number_formats["점검일시"] == "mm-dd-yy"
    assert tpl.number_formats["개소"] == "0_);[Red]\\(0\\)"
    assert "온도차" not in tpl.number_formats          # General은 담지 않는다
    assert tpl.date_format == "mm-dd-yy"               # 날짜서식을 봤으면 폴백값도 그것

    assert tpl.print_title_rows == "$1:$4"
    assert tpl.fit_to_page is True
    assert tpl.source_label == "양식.xlsx"

    # 없는 시트 · 없는 파일 · 헤더 없는 시트는 예외가 아니라 None
    assert xf.capture(path, "없는시트", 4) is None
    assert xf.capture(tmp / "없는파일.xlsx", "대상 리스트", 4) is None
    assert xf.capture(path, "대상 리스트", 99) is None

    # 기본서식은 파일 없이도 만들어진다
    b = xf.basic()
    assert b.header_row == 1 and b.header_style.font.bold is True
    assert b.source_label == "기본서식"
    print("  ✓ 서식 캡처(제목 블록·스타일·헤더명 매칭 열너비/숫자서식·인쇄설정·폴백)")


def test_format_consensus(tmp: Path):
    """회신본마다 손댄 열너비를 걷어내고 원래 배포된 양식을 되찾는다.

    실측(sampledata 8개): 글꼴·색·숫자서식·제목은 8/8 만장일치이고 갈리는 것은
    열너비뿐이었다(한 지사만 A열을 38.8로 늘려 놓았다). 그래서 다수결은
    '손댄 흔적을 걷어내는' 장치다.
    """
    caps = []
    for i, width in enumerate([16.2, 16.2, 16.2, 38.8]):
        path = tmp / f"{i}.xlsx"
        _styled_form(path)
        wb = openpyxl.load_workbook(path)
        wb["대상 리스트"].column_dimensions["A"].width = width
        wb.save(path)
        caps.append(xf.capture(path, "대상 리스트", 4))

    merged = xf.consensus(caps)
    assert merged.widths["지사"] == 16.2            # 3:1로 다수값이 이긴다
    assert merged.widths["온도차"] == 33.0          # 갈리지 않은 값은 그대로
    assert merged.number_formats["점검일시"] == "mm-dd-yy"
    assert merged.header_style.font.bold is True
    assert merged.print_title_rows == "$1:$4"
    assert len(merged.title_rows) == 3
    assert "취합 파일 4개의 공통 서식" in merged.source_label
    assert "열너비 1개" in merged.source_label       # 소수 의견이 있었음을 밝힌다

    # 동수면 먼저 들어온 것(호출자가 파일명 오름차순으로 넣는다). 값의 크기가 아니라
    # 입력 순서를 따른다는 것은 양방향으로 뒤집어 봐야만 드러난다 — 한 방향만 보면
    # '더 작은 값이 이긴다'는 잘못된 구현도 우연히 통과할 수 있다
    two = [caps[0], caps[3]]
    assert xf.consensus(two).widths["지사"] == 16.2
    two_rev = [caps[3], caps[0]]
    assert xf.consensus(two_rev).widths["지사"] == 38.8

    # 캡처 실패가 섞여도 나머지로 만든다. 전부 실패면 None
    assert xf.consensus([None, caps[0], None]).widths["지사"] == 16.2
    assert xf.consensus([None, None]) is None
    assert xf.consensus([]) is None
    print("  ✓ 서식 다수결(열너비 소수의견 제거·동수는 첫 파일·캡처 실패 혼재)")


def test_format_apply(tmp: Path):
    """캡처한 서식이 출력 시트에 옮겨진다. 필터·인쇄영역은 결과 행수로 재계산된다."""
    src = tmp / "양식.xlsx"
    _styled_form(src)
    tpl = xf.capture(src, "대상 리스트", 4)

    wb = openpyxl.Workbook()
    ws = wb.active
    out_headers = ["부서", "지사", "개소", "점검일시", "온도차"]   # 앞에 부서 1칸
    first = xf.open_sheet(ws, tpl, out_headers)
    assert first == 5                     # 헤더가 원본과 같은 4행 → 데이터는 5행부터

    rows = [["대구", "대구", 2, "2025-01-08", 3.3],
            ["용인", "용인", 1, "2024-11-21", 10.9]]
    for r, row in enumerate(rows, start=first):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    xf.close_sheet(ws, tpl, out_headers, first, first + len(rows) - 1)

    # 제목 블록: 값·병합·행높이가 살아 있고, 병합은 출력 열수(5)까지 넓어진다
    assert ws["A1"].value == "2025년 집중안전점검 리스트"
    assert ws.row_dimensions[1].height == 40
    assert "A1:E1" in [str(m) for m in ws.merged_cells.ranges]
    assert ws["A3"].value == "※ 2024.11.01 기준으로 작성"

    # 헤더 행: 원본 자리(4행)에 원본 스타일
    assert [ws.cell(row=4, column=c).value for c in range(1, 6)] == out_headers
    assert ws.cell(row=4, column=1).font.bold is True
    assert ws.cell(row=4, column=1).fill.fgColor.rgb.endswith("C0C0C0")

    # 데이터 셀 서식 + 헤더 이름으로 따라온 숫자서식
    assert ws.cell(row=5, column=1).border.bottom.style == "thin"
    assert ws.cell(row=5, column=4).number_format == "mm-dd-yy"     # 점검일시
    assert ws.cell(row=5, column=3).number_format == "0_);[Red]\\(0\\)"   # 개소

    # 열너비: 앞에 낀 부서 컬럼만큼 밀려도 헤더 이름으로 제 컬럼을 찾는다. 부서 열은 값 길이에 맞춘다
    assert ws.column_dimensions["B"].width == 16.2     # 지사
    assert ws.column_dimensions["D"].width == 30.8     # 점검일시
    assert xf.MIN_WIDTH <= ws.column_dimensions["A"].width <= xf.MAX_WIDTH

    # 재계산 대상 — 원본은 A4:D5였지만 결과 행수로 다시 잡힌다
    assert ws.auto_filter.ref == "A4:E6"
    # openpyxl의 print_area는 읽을 때 시트명과 절대참조($)를 붙여 돌려준다(3.1.5 확정 동작) —
    # 지정한 범위 자체는 A1:E6이고 그 부분만 확인한다
    assert ws.print_area == "'Sheet'!$A$1:$E$6"
    assert ws.freeze_panes == "A5"
    assert ws.print_title_rows == "$1:$4"              # 반복행은 그대로

    # 기본서식으로도 돈다 — 제목 블록이 없고 헤더가 1행
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    b = xf.basic()
    f2 = xf.open_sheet(ws2, b, ["가", "나"])
    assert f2 == 2
    ws2.cell(row=2, column=1, value="x")
    ws2.cell(row=2, column=2, value="y")
    xf.close_sheet(ws2, b, ["가", "나"], 2, 2)
    assert ws2.cell(row=1, column=1).font.bold is True
    assert ws2.auto_filter.ref == "A1:B2"
    assert ws2.freeze_panes == "A2"

    # 상속할 표시형식이 없는 날짜 컬럼은 그냥 두면 '2025-01-08 00:00:00'으로 보인다.
    # 이 작업이 고치려는 증상이라 기본서식에서도 날짜 형식을 붙인다
    import datetime as _dt
    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    b3 = xf.basic()
    f3 = xf.open_sheet(ws3, b3, ["점검일", "메모"])
    ws3.cell(row=f3, column=1, value=_dt.datetime(2025, 1, 8))
    ws3.cell(row=f3, column=2, value="글자")
    xf.close_sheet(ws3, b3, ["점검일", "메모"], f3, f3)
    assert ws3.cell(row=f3, column=1).number_format == xf.DEFAULT_DATE_FORMAT
    assert ws3.cell(row=f3, column=2).number_format == "General"   # 날짜가 아닌 칸은 그대로

    # 기준 서식에서 본 날짜 형식이 있으면 그것을 쓴다(지어내지 않는다)
    wb4 = openpyxl.Workbook()
    ws4 = wb4.active
    tpl4 = xf.capture(src, "대상 리스트", 4)
    tpl4.number_formats.pop("점검일시")          # 그 컬럼만 상속 실패한 상황
    h4 = ["지사", "개소", "점검일시", "온도차"]
    f4 = xf.open_sheet(ws4, tpl4, h4)
    ws4.cell(row=f4, column=3, value=_dt.datetime(2025, 1, 8))
    xf.close_sheet(ws4, tpl4, h4, f4, f4)
    assert ws4.cell(row=f4, column=3).number_format == "mm-dd-yy"

    # 제목 블록을 쓰다 중간에 실패하면(잘못된 병합 범위 등) 이미 쓴 값을 되돌리고
    # 진짜 빈 시트에서 폴백해야 한다 — 안 그러면 남은 제목 텍스트 위에 데이터가 겹친다
    bad = xf.FormatTemplate(
        header_row=3,
        title_rows=[
            xf.TitleRow([(1, "2025년 리스트", xf.CellStyle())], 40),
            xf.TitleRow([], None),
            xf.TitleRow([(1, "※ 안내문", xf.CellStyle())], None),
        ],
        title_merges=[(5, 1, 1, 3)],   # min_row(5) > max_row(1) — merge_cells가 예외를 던진다
    )
    wb5 = openpyxl.Workbook()
    ws5 = wb5.active
    f5 = xf.open_sheet(ws5, bad, ["가", "나"])
    assert f5 == 2                                          # 헤더 1행 → 데이터 2행부터
    assert [ws5.cell(row=1, column=c).value for c in (1, 2)] == ["가", "나"]
    # 헤더 두 칸을 뺀 나머지 어디에도 제목·안내 텍스트가 남아 있지 않다
    assert ws5.max_row == 1 and ws5.max_column == 2
    assert not list(ws5.merged_cells.ranges)                # 남은 병합도 없다

    # 스타일링(열너비 등)이 실패해도 필터·인쇄영역 재계산은 반드시 돈다 —
    # 장식(서식)과 결함 수정(필터·인쇄영역)을 한 try에 묶으면 안 되는 이유
    tpl6 = xf.FormatTemplate(header_row=1)
    tpl6.widths["가"] = "abc"                 # 폭 대입에서 TypeError를 일으킨다
    wb6 = openpyxl.Workbook()
    ws6 = wb6.active
    h6 = ["가", "나"]
    f6 = xf.open_sheet(ws6, tpl6, h6)
    ws6.cell(row=f6, column=1, value=1)
    ws6.cell(row=f6, column=2, value=2)
    xf.close_sheet(ws6, tpl6, h6, f6, f6)
    assert ws6.auto_filter.ref == "A1:B2"
    assert ws6.freeze_panes == "A2"
    assert ws6.print_area == "'Sheet'!$A$1:$B$2"
    print("  ✓ 서식 적용(제목 블록·헤더 자리·필터/인쇄영역 재계산·기본서식·날짜 폴백·중간 실패 되돌리기·스타일 실패해도 재계산 진행)")


def test_close_sheet_uses_actual_header_row(tmp: Path):
    """open_sheet가 폴백해 실제 헤더를 1행에 쓰면 close_sheet도 그 실제 행을 써야 한다.

    tpl.header_row(원본의 4행)를 그대로 믿으면, 폴백으로 헤더가 1행에 쓰였는데도
    필터·틀고정이 4행 기준으로 계산돼 데이터 행을 헤더처럼 다룬다. 제목행에 걸린
    행높이(162pt)도 delete_rows가 지우지 않으므로 폴백 헤더에 그대로 남는다.
    """
    bad = xf.FormatTemplate(
        header_row=4,
        title_rows=[
            xf.TitleRow([(1, "2025년 리스트", xf.CellStyle())], 162),
            xf.TitleRow([], None),
            xf.TitleRow([(1, "※ 안내문", xf.CellStyle())], None),
        ],
        title_merges=[(5, 1, 1, 2)],   # min_row(5) > max_row(1) → merge_cells가 예외를 던진다
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    first = xf.open_sheet(ws, bad, ["가", "나"])
    assert first == 2                        # 폴백 → 헤더 1행, 데이터 2행부터
    for r, row in enumerate([["a", 1], ["b", 2], ["c", 3]], start=first):
        ws.cell(row=r, column=1, value=row[0])
        ws.cell(row=r, column=2, value=row[1])
    xf.close_sheet(ws, bad, ["가", "나"], first, first + 2)

    # tpl.header_row(4)가 아니라 실제 헤더 행(1)을 기준으로 재계산돼야 한다
    assert ws.auto_filter.ref == "A1:B4", ws.auto_filter.ref
    assert ws.freeze_panes == "A2", ws.freeze_panes
    # 제목행의 162pt 행높이가 폴백 헤더(1행)에 남아 있지 않다
    assert ws.row_dimensions[1].height is None, ws.row_dimensions[1].height
    print("  ✓ close_sheet가 폴백 후 실제 헤더 행을 쓴다(필터·틀고정·제목행 행높이 청소)")


def test_fit_width_ignores_formula_text(tmp: Path):
    """수식이 든 열은 화면에 보이는 값이 아니라 헤더 길이로 폭을 잰다.

    종합요약 시트의 지표 칸은 실제로는 '=SUMIF(...)' 같은 긴 수식 문자열을 담고
    있어서, len(str(value))로 재면 60자 상한에 바로 닿는다 — 이 브랜치 이전에는
    그 칸이 기본폭(8.43)이었던 회귀다.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["부서", "온도차"]
    tpl = xf.FormatTemplate(header_row=1)
    first = xf.open_sheet(ws, tpl, headers)
    ws.cell(row=first, column=1, value="대구")
    ws.cell(row=first, column=2,
            value="=SUMIF('대상 리스트'!A5:A5,$A2,'대상 리스트'!E5:E5)")
    xf.close_sheet(ws, tpl, headers, first, first)
    assert ws.column_dimensions["B"].width < 20, ws.column_dimensions["B"].width
    print("  ✓ 수식 열은 수식 텍스트가 아니라 헤더 길이로 폭을 잰다")


def test_synthesis_inherits_format(tmp: Path):
    """모드 B·C·D가 원본 양식의 서식을 이어받고, 열 오프셋만큼 밀어서 맞춘다."""
    files = []
    for dept in ("대구", "용인"):
        path = tmp / f"{dept}.xlsx"
        _styled_form(path)
        uf = ag.read_file(path)
        uf.dept = dept
        files.append(uf)

    wb, notes, run = ag.synthesize(files, "B", {}, [], all_files=files)
    ws = wb["대상 리스트"]

    # 제목 블록이 살아 있고 헤더는 원본 자리(4행), 앞에 부서 1칸
    assert ws["A1"].value == "2025년 집중안전점검 리스트"
    assert ws.cell(row=4, column=1).value == "부서"
    assert ws.cell(row=4, column=2).value == "지사"
    assert ws.cell(row=4, column=1).font.bold is True
    assert ws.cell(row=5, column=1).value == "대구"
    assert ws.column_dimensions["B"].width == 16.2        # 지사 = 원본 A열 너비
    assert ws.auto_filter.ref == "A4:E6"                  # 2행이 쌓였다
    assert ws.print_title_rows == "$1:$4"
    # 숫자서식도 최종 취합 결과까지 실제로 살아남는다 — 이게 없으면 날짜가
    # '2025-01-08 00:00:00'으로 보이는, 설계 §1이 지목한 그 증상이다
    assert ws.cell(row=5, column=4).number_format == "mm-dd-yy"   # 점검일시(D열)
    assert run.aggregated == 2 and run.result_rows == 2
    assert "공통 서식" in run.source_label
    assert any("기준 서식" in n for n in notes)           # 화면 안내로도 나간다

    # 모드 C는 구분·부서 2칸이 앞에 낀다
    wb_c, _, _ = ag.synthesize(files, "C", {"대상 리스트": "상반기"}, [], all_files=files)
    ws_c = wb_c["상반기"]
    assert [ws_c.cell(row=4, column=c).value for c in range(1, 4)] == ["구분", "부서", "지사"]
    assert ws_c.column_dimensions["C"].width == 16.2      # 지사가 2칸 밀렸다

    # 개요 시트가 결과 파일 맨 앞에 붙는다 — 결과만 받아도 누락이 보이게
    assert wb.sheetnames == ["취합 개요", "대상 리스트"]

    # 모드 D 종합요약은 헤더가 1행이 아니라 원본 헤더 행에 있는 시트를 참조해야 한다
    wb_d, _, _ = ag.synthesize(files, "D", {}, ["온도차"], all_files=files)
    assert wb_d.sheetnames[:2] == ["취합 개요", "종합요약"]
    formula = wb_d["종합요약"].cell(row=2, column=2).value
    # 부서(A열) 기준으로 온도차(E열)를 참조해야 한다 — 4행 헤더 아래(5행) 데이터
    assert formula.startswith("=SUMIF(") and "!A5:A5" in formula and "!E5:E5" in formula
    assert "#REF" not in formula

    # 규칙 1: 양식을 주면 다수결하지 않고 그 파일의 서식을 쓴다
    tmpl = tmp / "배포양식.xlsx"
    _styled_form(tmpl)
    wb0 = openpyxl.load_workbook(tmpl)
    wb0["대상 리스트"].column_dimensions["A"].width = 99.0
    wb0.save(tmpl)
    wb_t, _, run_t = ag.synthesize(files, "B", {}, [], all_files=files, template_path=tmpl)
    assert wb_t["대상 리스트"].column_dimensions["B"].width == 99.0   # 회신본 16.2가 아니다
    assert "등록 양식" in run_t.source_label

    # 폴백: 원본을 못 열어도 값은 전부 나온다 (서식은 부가물이다)
    for uf in files:
        uf.path = tmp / f"사라진_{uf.path.name}"
    wb_f, _, run_f = ag.synthesize(files, "B", {}, [], all_files=files)
    ws_f = wb_f["대상 리스트"]
    assert ws_f.cell(row=1, column=1).value == "부서"        # 기본서식은 헤더가 1행
    assert ws_f.cell(row=2, column=1).value == "대구"
    assert ws_f.cell(row=1, column=1).font.bold is True
    assert run_f.source_label == "기본서식"
    print("  ✓ 합성 서식 상속(모드 B·C 오프셋·모드 D 요약 참조 행·양식 우선·폴백)")


def test_mode_a_own_format(tmp: Path):
    """모드 A는 원본 보존형이라 각 시트가 자기 파일의 서식을 쓴다.

    헤더가 1행으로 당겨지지 않고 원본 행 번호에 놓이므로 이미지 앵커 보정이
    0이 된다 — 사진이 원본과 같은 데이터 행에 붙는다.
    """
    files = []
    # 대구 폭은 _styled_form 기본값(16.2)과도, 용인 폭(38.8)과도 달라야 한다 —
    # 16.2로 재설정하면 기본값과 우연히 같아져서, 대구 파일을 잘못된(엉뚱한) 원본에서
    # 캡처하는 결함이 있어도 이 단언이 통과해 버린다(실측: 용인 쪽 단언만 걸린다)
    for dept, width in (("대구", 24.6), ("용인", 38.8)):
        path = tmp / f"{dept}.xlsx"
        _styled_form(path)
        wb0 = openpyxl.load_workbook(path)
        wb0["대상 리스트"].column_dimensions["A"].width = width
        wb0.save(path)
        uf = ag.read_file(path)
        uf.dept = dept
        files.append(uf)

    # 원본 5행(데이터 첫 행)에 붙은 사진 하나
    # 주의: sheet.images의 실제 내부 스키마(marker_from/marker_to/display/resolution 등,
    # aggregate.py의 _read_images/_shift_anchor 참고)에 맞춰 구성한다. from·to는 1-base
    # (행, 열), marker_from은 0-base (열, 행, 열오프셋, 행오프셋)이다.
    buf = io.BytesIO()
    PILImage.new("RGB", (300, 300), "red").save(buf, format="PNG")
    data = buf.getvalue()
    files[0].sheets[0].images = [{
        "from": (5, 1), "to": (5, 1), "single_cell_anchor": True,
        "marker_from": (0, 4, 0, 0), "marker_to": None, "edit_as": "oneCell",
        "display": (300, 300), "ext": "png", "bytes": len(data), "data": data,
        "sha256": "", "resolution": (300, 300),
    }]

    wb, notes, run = ag.synthesize(files, "A", {}, [], all_files=files)
    assert wb.sheetnames == ["취합 개요", "대구_대상 리스트", "용인_대상 리스트"]

    # 파일마다 자기 열너비 — 다수결로 뭉개지 않는다
    assert wb["대구_대상 리스트"].column_dimensions["A"].width == 24.6
    assert wb["용인_대상 리스트"].column_dimensions["A"].width == 38.8

    ws = wb["대구_대상 리스트"]
    assert ws["A1"].value == "2025년 집중안전점검 리스트"
    assert ws.cell(row=4, column=1).value == "지사"      # 부서 컬럼 없음, 원본 자리
    assert ws.cell(row=5, column=1).value == "대구"
    assert ws.auto_filter.ref == "A4:D5"

    # 헤더가 제자리라 앵커 보정이 0 → 사진이 원본과 같은 행(5행 = 앵커 index 4)
    assert len(ws._images) == 1
    assert ws._images[0].anchor._from.row == 4
    assert "파일별 원본 서식" in run.source_label or "양식" in run.source_label
    print("  ✓ 모드 A(파일별 자기 서식·헤더 원본 자리·앵커 보정 0)")


def test_header_mismatch_counts_toward_total(tmp: Path):
    """헤더불일치 시트로 갈린 행도 취합 개요의 결과 행수에 잡혀야 한다.

    갈린 행이 total_rows에서 빠지면, 부서 하나가 통째로 헤더불일치로 걸러질 때
    취합 개요는 '결과 0행'이라고 말하는데 워크북에는 그 행들이 실제로 들어 있는
    모순이 생긴다 — 개요가 있는 이유(파일이 스스로 무엇이 빠졌는지 말한다)를
    정면으로 어기는 셈이다.
    """
    a = tmp / "기획부.xlsx"
    b = tmp / "총무부.xlsx"
    make_book(a, {"예산": [["사번", "부서", "예산액"], ["A1", "기획부", 100], ["A2", "기획부", 200]]})
    make_book(b, {"예산": [["사번", "부서", "예산액", "비고"],
                          ["B1", "총무부", 300, "메모"]]})   # 헤더가 하나 더 많다 → 불일치
    files = []
    for path in (a, b):
        uf = ag.read_file(path)
        ag.review_stage1(uf, rules={})
        files.append(uf)

    wb, notes, run = ag.synthesize(files, "B", {}, [], all_files=files)
    assert any("헤더불일치" in n for n in notes), notes
    err_title = next(name for name in wb.sheetnames if "헤더불일치" in name)
    assert wb[err_title].cell(row=4, column=1).value == "총무부"   # 갈린 행이 실제로 있다
    # 기획부 2행(정상 시트) + 총무부 1행(헤더불일치 시트) = 3행이 진짜 결과다
    assert run.result_rows == 3, run.result_rows
    print("  ✓ 헤더불일치 시트로 빠진 행도 개요의 결과 행수에 반영된다")


def test_overview_summary(tmp: Path):
    """결과 파일이 '무엇이 빠졌는지'를 스스로 말한다.

    종전에는 8개 중 2개가 통째로 빠져도 그 사실이 화면·stdout에만 있어서,
    파일만 받은 사람은 6개가 전부인 줄 알았다.
    """
    good = ag.UploadedFile(path=Path("대구.xlsx"), dept="대구")
    good.sheets = [ag.SheetData("리스트", 4, ["지사"], [["대구"], ["대구"]], [5, 6])]

    bad = ag.UploadedFile(path=Path("중앙지사.xlsx"), dept="중앙지사")
    bad.sheets = [ag.SheetData("리스트", 4, ["지사"], [], [])]
    bad.issues = [ag.Issue("중앙지사.xlsx", "리스트", f"A{r}", "누락", "1단계",
                           ag.ERROR, "필수값 누락", "지사") for r in range(5, 13)]

    unread = ag.UploadedFile(path=Path("깨진파일.xlsx"), dept="깨진파일", readable=False)

    run = ag.build_summary([good, bad, unread], [good], "B", "등록 양식 '리스트'",
                           result_rows=2)
    assert run.submitted == 3 and run.aggregated == 1 and run.excluded == 2
    assert run.mode == "B" and run.result_rows == 2
    rows = {f.name: f for f in run.files}
    assert rows["대구.xlsx"].included is True and rows["대구.xlsx"].rows == 2
    assert rows["대구.xlsx"].status == "정상" and rows["대구.xlsx"].reason == ""
    assert rows["중앙지사.xlsx"].included is False
    assert rows["중앙지사.xlsx"].reason == "필수값 누락 8건"
    assert rows["깨진파일.xlsx"].reason == "파일을 읽지 못했습니다"

    wb = openpyxl.Workbook()
    wb.create_sheet("결과")
    wb.remove(wb["Sheet"])
    xf.overview_sheet(wb, 0, run)
    assert wb.sheetnames == ["취합 개요", "결과"]        # 맨 앞이어야 열자마자 보인다
    ws = wb["취합 개요"]
    assert "제출 3" in ws["A4"].value and "취합 1" in ws["A4"].value
    assert "제외 2" in ws["A4"].value and "결과 2행" in ws["A4"].value
    assert "등록 양식" in ws["A3"].value
    assert [ws.cell(row=6, column=c).value for c in range(1, 7)] == \
        ["부서", "파일명", "판정", "취합", "행수", "사유"]
    assert ws.cell(row=6, column=1).font.bold is True
    body = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=4).value
            for r in range(7, 10)}
    assert body["대구.xlsx"] == "○" and body["중앙지사.xlsx"] == "✕"

    # 기준 서식이 없으면 없다고 적는다 — 0으로 단정하지 않는다
    empty = ag.build_summary([], [], "B", "")
    wb2 = openpyxl.Workbook()
    xf.overview_sheet(wb2, 0, empty)
    assert "없음(기본서식 적용)" in wb2["취합 개요"]["A3"].value
    print("  ✓ 취합 개요(제출/취합/제외·제외 사유·기준 서식 표기)")


def test_reason_truncates_and_discloses(tmp: Path):
    """사유가 4종 이상이면 앞 3개만 보여주고 나머지 개수를 '외 N종'으로 밝힌다.

    조용히 3개만 보여주면 사유가 그게 전부인 것처럼 읽힌다 — 잘랐다는 사실을
    밝히는 것이 이 작업의 바인딩 제약이다. 지금까지 이 분기를 실행하는
    테스트가 없어, 조용히 지어낸 것처럼 보이는 사유를 만들어도 걸리지 않았다.
    """
    bad = ag.UploadedFile(path=Path("총무부.xlsx"), dept="총무부")
    bad.issues = [
        ag.Issue("총무부.xlsx", "리스트", "A1", "누락", "1단계", ag.ERROR, "필수값 누락(A1)", "부서"),
        ag.Issue("총무부.xlsx", "리스트", "B1", "정합성", "1단계", ag.ERROR, "형식 오류(B1)", "예산액"),
        ag.Issue("총무부.xlsx", "리스트", "C1", "중복", "1단계", ag.ERROR, "중복 행(C1)", "사번"),
        ag.Issue("총무부.xlsx", "리스트", "D1", "충돌", "1단계", ag.ERROR, "충돌 값(D1)", "지사"),
    ]
    run = ag.build_summary([bad], [], "B", "")
    reason = run.files[0].reason
    assert "외 1종" in reason, reason               # 4종 중 3개만 보여주고 나머지 1개는 밝힌다
    for label in ("필수값 누락", "형식 오류", "중복 행"):
        assert label in reason, reason
    assert "충돌 값" not in reason, reason          # 4번째는 실제로 잘려 나갔다
    print("  ✓ 사유 4종 이상은 앞 3개 + '외 N종'으로 잘림을 밝힌다")


def test_report_format_and_order(tmp: Path):
    """리포트에 기본서식과 개요가 붙고, synthesize 뒤로 옮겨도 내용이 같다."""
    path = tmp / "대구.xlsx"
    _styled_form(path)
    uf = ag.read_file(path)
    uf.dept = "대구"
    ag.review_stage1(uf, {})
    before = ([(i.file, i.sheet, i.cell, i.reason) for i in uf.issues],
              [(f.cell, f.original, f.corrected) for f in uf.fixes])

    # 종전 순서(리포트 먼저)와 새 순서(합성 먼저)의 내용이 같아야 한다
    ag.preprocess(uf)
    wb, _, run = ag.synthesize([uf], "B", {}, [], all_files=[uf])
    after = ([(i.file, i.sheet, i.cell, i.reason) for i in uf.issues],
             [(f.cell, f.original, f.corrected) for f in uf.fixes])
    assert before == after, "preprocess·synthesize가 issues/fixes를 건드리면 순서를 바꿀 수 없다"

    out = tmp / "report.xlsx"
    ag.write_report([uf], out, overview=run)
    rb = openpyxl.load_workbook(out)
    assert rb.sheetnames == ["취합 개요", "오류 목록", "자동교정 이력"]
    ws = rb["오류 목록"]
    assert ws.cell(row=1, column=1).font.bold is True
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.startswith("A1:G")
    assert ws.column_dimensions["A"].width <= xf.MAX_WIDTH

    # 자동교정이 0건이면 빈 시트로 두지 않고 그렇게 적는다
    fixes = rb["자동교정 이력"]
    if fixes.max_row == 2:
        assert fixes.cell(row=2, column=1).value == "해당 없음"

    # overview 없이도 돈다(취합 가능한 파일이 0건인 경로)
    out2 = tmp / "report2.xlsx"
    ag.write_report([uf], out2)
    assert openpyxl.load_workbook(out2).sheetnames == ["오류 목록", "자동교정 이력"]
    print("  ✓ 리포트 서식·개요·호출 순서 불변")


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="agg-test-"))
    try:
        for fn in (test_structure_and_stage1, test_default_miss_collapse,
                   test_matrix_checklist_header, test_plan_actual_pairs,
                   test_clean_file_and_gating,
                   test_rules_and_masking, test_synthesis_modes,
                   test_image_anchor_relocation, test_real_form_structure,
                   test_stacked_tables_and_tiered_header, test_pivot_flatten,
                   test_stage2_parallel,
                   test_format_capture,
                   test_format_consensus,
                   test_format_apply,
                   test_close_sheet_uses_actual_header_row,
                   test_fit_width_ignores_formula_text,
                   test_overview_summary,
                   test_reason_truncates_and_discloses,
                   test_synthesis_inherits_format,
                   test_header_mismatch_counts_toward_total,
                   test_mode_a_own_format,
                   test_report_format_and_order,
                   test_cli_end_to_end):
            sub = tmp / fn.__name__
            sub.mkdir()
            fn(sub)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\n전체 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
