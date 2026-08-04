#!/usr/bin/env python3
"""인증(F4-1) 점검. 실행: python3 test_auth.py

test_server.py와 달리 **실제 Supabase 프로젝트를 상대로** 돈다.
.env의 SUPABASE_* 값이 필요하며 네트워크를 쓴다.

테스트 계정은 매 실행마다 새로 만들고 끝나면 지운다. 사번에 `zz-test-`
접두사를 붙여 실계정과 구분한다.
"""

from __future__ import annotations

import io
import os
import sys
import time
import uuid

import httpx
import openpyxl

import aggregate as ag
from pathlib import Path

ag.load_env(Path(__file__).parent / ".env")
os.environ.pop("AUTH_DISABLED", None)          # 인증을 켠 상태로 돌려야 한다

from fastapi.testclient import TestClient       # noqa: E402

import auth                                     # noqa: E402
import server                                   # noqa: E402

EMP = f"zz-test-{uuid.uuid4().hex[:10]}"
PW = "test-" + uuid.uuid4().hex[:12]
RESET_EMAIL = "aggregation-test@example.com"
created_user_id: str | None = None
client_for_wait: TestClient | None = None
storage_paths: list[str] = []


def book_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "예산"
    ws["A1"] = "실적 보고"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    ws.append([None, None, None])
    for row in [["사번", "부서", "예산액"], ["A1", "기획부", 100]]:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def hwpx_bytes(tag: str) -> bytes:
    """hwpx fixture는 test_hwpx.py 것을 그대로 쓴다(구조를 두 곳에서 관리하지 않는다)."""
    import tempfile
    from pathlib import Path

    import test_hwpx as th
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / f"{tag}.hwpx"
        th._fixture(p, tag, b"\x89PNG-" + tag.encode())
        return p.read_bytes()


def wait(res, timeout: float = 120.0):
    """잡 기반 엔드포인트: job_id를 받아 완료까지 폴링하고 결과만 돌려준다."""
    assert res.status_code == 200, res.text
    jid = res.json()["job_id"]
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client_for_wait.get(f"/api/job/{jid}").json()
        if job["status"] == "done":
            return job["result"]
        if job["status"] == "error":
            raise AssertionError(f"잡 실패: {job['error']}")
        time.sleep(0.02)
    raise AssertionError("잡이 시간 안에 끝나지 않았다")


def test_protected_without_login():
    """로그인 없이는 취합 API에 접근할 수 없어야 한다."""
    anon = TestClient(server.app)
    for method, path in [("post", "/api/session"), ("get", "/api/auth/me"),
                         ("get", "/api/job/아무거나")]:
        res = getattr(anon, method)(path)
        assert res.status_code == 401, f"{path} → {res.status_code} (401이어야 함)"
    print("  ✓ 미로그인 차단(401)")


def test_signup_and_login():
    global created_user_id
    client = TestClient(server.app)

    # 입력 검증
    bad = client.post("/api/auth/signup",
                      json={"employee_no": "a", "password": PW, "reset_email": RESET_EMAIL})
    assert bad.status_code == 400 and "사번" in bad.json()["detail"], bad.text
    bad = client.post("/api/auth/signup",
                      json={"employee_no": EMP, "password": "short", "reset_email": RESET_EMAIL})
    assert bad.status_code == 400 and "비밀번호" in bad.json()["detail"], bad.text

    res = client.post("/api/auth/signup",
                      json={"employee_no": EMP, "password": PW, "reset_email": RESET_EMAIL})
    assert res.status_code == 200, res.text
    profile = res.json()["profile"]
    created_user_id = profile["id"]
    assert profile["employee_no"] == EMP and profile["role"] == "user", profile
    assert profile["status"] == "active", profile

    # 같은 사번 재가입은 막힌다
    dup = client.post("/api/auth/signup",
                      json={"employee_no": EMP, "password": PW, "reset_email": RESET_EMAIL})
    assert dup.status_code == 409, dup.text

    # 대소문자만 다른 사번도 같은 사번이다. 가상 이메일이 소문자라 여기서
    # 걸러내지 않으면 Auth 쪽에서 이메일 중복으로 늦게 터진다
    dup2 = client.post("/api/auth/signup",
                       json={"employee_no": EMP.upper(), "password": PW, "reset_email": RESET_EMAIL})
    assert dup2.status_code == 409, f"대소문자 변형이 통과했다: {dup2.status_code} {dup2.text}"
    print("  ✓ 가입(입력 검증 + 사번 중복 409, 대소문자 무관)")

    # 틀린 비밀번호 — 사번 존재 여부를 노출하지 않는 같은 메시지
    wrong = client.post("/api/auth/login", json={"employee_no": EMP, "password": "wrong-pw-123"})
    missing = client.post("/api/auth/login", json={"employee_no": "zz-none", "password": PW})
    assert wrong.status_code == missing.status_code == 401, (wrong.text, missing.text)
    assert wrong.json()["detail"] == missing.json()["detail"], "계정 열거가 가능하면 안 된다"

    # 대문자로 입력해도 같은 계정으로 로그인된다
    upper = client.post("/api/auth/login", json={"employee_no": EMP.upper(), "password": PW})
    assert upper.status_code == 200, upper.text
    assert upper.json()["profile"]["employee_no"] == EMP, upper.json()

    ok = client.post("/api/auth/login", json={"employee_no": EMP, "password": PW})
    assert ok.status_code == 200, ok.text
    assert ok.json()["profile"]["employee_no"] == EMP
    # 토큰은 본문에 실리지 않고 httpOnly 쿠키로만 나간다
    assert "access_token" not in ok.text, "토큰이 응답 본문에 노출됐다"
    cookie = ok.cookies.get(auth.ACCESS_COOKIE)
    assert cookie, ok.cookies
    set_cookie = " ".join(ok.headers.get_list("set-cookie")).lower()
    assert "httponly" in set_cookie and "samesite=lax" in set_cookie, set_cookie
    print("  ✓ 로그인(계정 열거 방지 + httpOnly·SameSite 쿠키)")
    global client_for_wait
    client_for_wait = client
    return client


def test_authenticated_flow(client: TestClient):
    """로그인한 클라이언트는 취합 전 구간을 쓸 수 있다."""
    me = client.get("/api/auth/me")
    assert me.status_code == 200 and me.json()["profile"]["employee_no"] == EMP, me.text
    assert me.json()["auth_enabled"] is True

    sid = client.post("/api/session").json()["sid"]
    res = client.post(f"/api/session/{sid}/files",
                      files=[("files", ("기획부.xlsx", book_bytes(), "application/octet-stream"))])
    assert res.status_code == 200, res.text
    jid = res.json()["job_id"]
    for _ in range(2000):
        job = client.get(f"/api/job/{jid}").json()
        if job["status"] != "running":
            break
    assert job["status"] == "done", job
    assert "owner" not in job, "소유자 정보가 응답에 새어나가면 안 된다"
    print("  ✓ 로그인 후 취합 API 사용 가능")
    return sid


def test_isolation(sid: str):
    """다른 사용자의 세션·잡은 보이지 않아야 한다 (§6.2)."""
    other_emp = f"zz-test-{uuid.uuid4().hex[:10]}"
    other = TestClient(server.app)
    made = other.post("/api/auth/signup",
                      json={"employee_no": other_emp, "password": PW, "reset_email": RESET_EMAIL})
    assert made.status_code == 200, made.text
    other_id = made.json()["profile"]["id"]
    try:
        assert other.post("/api/auth/login",
                          json={"employee_no": other_emp, "password": PW}).status_code == 200
        # 남의 sid는 존재 자체를 알리지 않는다
        assert other.post(f"/api/session/{sid}/review", json={"no_ai": True}).status_code == 404
        assert other.get(f"/api/session/{sid}/download/result").status_code == 404
        print("  ✓ 사용자 간 세션 격리(404)")
    finally:
        _delete_storage_for_owner(other_id)
        _delete_user(other_id)


def test_role_and_suspend(client: TestClient):
    """manage_users.py로 바꾼 권한·상태가 즉시 반영돼야 한다."""
    import manage_users

    assert manage_users.main(["promote", EMP.upper()]) == 0, "대소문자 무관하게 찾아야 한다"
    assert client.get("/api/auth/me").json()["profile"]["role"] == "admin"

    assert manage_users.main(["demote", EMP]) == 0
    assert client.get("/api/auth/me").json()["profile"]["role"] == "user"

    # 정지는 이미 발급된 토큰에도 즉시 걸려야 한다 (매 요청 status 확인)
    assert manage_users.main(["suspend", EMP]) == 0
    blocked = client.get("/api/auth/me")
    assert blocked.status_code == 403, f"{blocked.status_code} {blocked.text}"
    assert "정지" in blocked.json()["detail"]
    # 정지 상태에서는 새로 로그인도 막힌다
    assert client.post("/api/auth/login",
                       json={"employee_no": EMP, "password": PW}).status_code == 403

    assert manage_users.main(["activate", EMP]) == 0
    assert client.get("/api/auth/me").status_code == 200
    print("  ✓ 권한 변경 + 계정 정지(발급된 토큰에도 즉시 적용)")


def test_persistence(client: TestClient):
    """업로드 파일은 Storage, 작업 기록은 DB에 남아야 한다 (세션 휘발성 해소)."""
    import store

    made = client.post("/api/session").json()
    sid, project_id = made["sid"], made["project_id"]
    assert project_id, "취합 시작 시 프로젝트가 자동 생성돼야 한다"

    res = client.post(f"/api/session/{sid}/files",
                      files=[("files", ("기획부.xlsx", book_bytes(), "application/octet-stream"))])
    jid = res.json()["job_id"]
    for _ in range(2000):
        if client.get(f"/api/job/{jid}").json()["status"] != "running":
            break
    wait(client.post(f"/api/session/{sid}/review", json={"no_ai": True}))
    agg = wait(client.post(f"/api/session/{sid}/aggregate", json={"mode": "B", "included": [0]}))
    assert agg["metrics"]["aggregated"] == 1, agg

    # 이력에 남았는지
    detail = client.get(f"/api/history/{project_id}")
    assert detail.status_code == 200, detail.text
    d = detail.json()
    assert len(d["uploaded_files"]) == 1, d["uploaded_files"]
    upl = d["uploaded_files"][0]
    assert upl["original_name"] == "기획부.xlsx" and upl["status"] == "reviewed", upl
    assert upl["review_results"] and upl["review_results"][0]["badge"] == "정상", upl
    job = d["aggregation_jobs"][0]
    assert job["status"] == "done" and job["progress"] == 100, job
    assert job["mode"] == "B" and job["result_url"], job

    # 원본 파일이 Storage에 실제로 있는지 (재시작해도 남는 근거)
    files = store._rest("GET", "/uploaded_files",
                        params={"id": f"eq.{upl['id']}", "select": "storage_path"}).json()
    blob = store.get_object(store.UPLOAD_BUCKET, files[0]["storage_path"])
    assert blob[:2] == b"PK" and len(blob) > 1000, len(blob)

    # 세션이 사라진 뒤에도 이력에서 결과를 내려받을 수 있는지
    import server
    server.SESSIONS.clear()
    dl = client.get(f"/api/history/{project_id}/download/{job['id']}")
    assert dl.status_code == 200 and dl.content[:2] == b"PK", dl.status_code

    # 목록에도 보이고, 남의 이력은 404
    projects = client.get("/api/history").json()["projects"]
    mine = next((p for p in projects if p["id"] == project_id), None)
    assert mine, projects
    assert mine["file_names"] == ["기획부.xlsx"] and mine["file_count"] == 1, mine
    assert mine["kind"] == "excel" and mine["jobs"][0]["status"] == "done", mine
    print("  ✓ 영속화(Storage 원본 + DB 기록 + 세션 없이 이력 다운로드)")

    # 화면 9의 검색·유형 필터가 실제로 걸러내야 한다
    def ids(**params):
        qs = "&".join(f"{k}={v}" for k, v in params.items())
        return [p["id"] for p in client.get(f"/api/history?{qs}").json()["projects"]]

    assert project_id in ids(q="기획부"), "파일명으로 찾아야 한다"
    assert project_id in ids(q="기획"), "부분 일치로 찾아야 한다"
    assert project_id in ids(q="엑셀"), "작업명으로도 찾아야 한다"
    assert project_id not in ids(q="없는파일명zzz"), "안 맞으면 빠져야 한다"
    assert project_id in ids(type="excel"), "엑셀 유형에 잡혀야 한다"
    assert project_id not in ids(type="hwpx"), "hwpx 유형에는 안 잡혀야 한다"
    print("  ✓ 이력 검색·유형 필터")
    return project_id


def _object_of(bucket: str, path: str, tries: int = 6) -> bytes:
    """방금 올린 Storage 객체를 읽는다.

    오브젝트 스토리지는 쓰기 직후 읽기가 즉시 보장되지 않는다. 엑셀 시나리오는 업로드와
    확인 사이에 검토·취합이 끼어 시간이 벌어지지만, 한글 병합은 업로드 직후라 간격이
    거의 없어 한 번 간헐 실패했다. 값을 눙치지 않고 잠깐 기다렸다 다시 읽는다.
    """
    import store
    last = None
    for i in range(tries):
        try:
            return store.get_object(bucket, path)
        except Exception as exc:                      # 404를 포함한 일시적 실패
            last = exc
            time.sleep(0.3 * (i + 1))
    raise AssertionError(f"Storage 객체를 읽지 못했다: {bucket}/{path} — {last}")


def test_hwpx_persistence(client: TestClient):
    """한글 병합도 Storage·DB에 남고, 이력에서 유형으로 걸러져야 한다 (F3 + F4-2)."""
    import store

    made = client.post("/api/session", json={"kind": "hwpx"}).json()
    sid, project_id = made["sid"], made["project_id"]
    assert made["kind"] == "hwpx" and project_id, made

    out = wait(client.post(f"/api/hwpx/session/{sid}/files", files=[
        ("files", ("가.hwpx", hwpx_bytes("A"), "application/octet-stream")),
        ("files", ("나.hwpx", hwpx_bytes("B"), "application/octet-stream"))]))
    assert all(f["readable"] for f in out["files"]), out

    merged = wait(client.post(f"/api/hwpx/session/{sid}/merge", json={"order": [1, 0]}))
    assert merged["metrics"] == {"merged": 2, "sections": 2,
                                 "size_kb": merged["metrics"]["size_kb"]}, merged

    d = client.get(f"/api/history/{project_id}").json()
    assert len(d["uploaded_files"]) == 2, d["uploaded_files"]
    job = d["aggregation_jobs"][0]
    assert job["status"] == "done" and job["progress"] == 100, job
    # 합성 모드(A/B/C/D)는 엑셀 취합 개념이라 한글 병합에는 없다 → 임의 값이 아니라 NULL
    assert job["mode"] is None, job
    assert job["result_url"].endswith(".hwpx"), job["result_url"]
    kinds = store._rest("GET", "/aggregation_jobs",
                        params={"id": f"eq.{job['id']}", "select": "kind"}).json()
    assert kinds[0]["kind"] == "hwpx", kinds

    # 기록의 유형이 hwpx이고 원본도 hwpx 확장자로 Storage에 올라갔는지
    rows = store._rest("GET", "/uploaded_files",
                       params={"project_id": f"eq.{project_id}",
                               "select": "kind,storage_path"}).json()
    assert {r["kind"] for r in rows} == {"hwpx"}, rows
    assert all(r["storage_path"].endswith(".hwpx") for r in rows), rows
    assert _object_of(store.UPLOAD_BUCKET, rows[0]["storage_path"])[:2] == b"PK", rows[0]

    # 세션이 사라진 뒤에도 이력에서 내려받을 수 있고, hwpx로 내려와야 한다
    import server
    server.SESSIONS.clear()
    # 결과도 방금 올라간 객체라 첫 요청이 이를 수 있다(위 _object_of와 같은 이유).
    # 사용자 경로는 세션의 로컬 파일을 주므로 이 지연에 걸리지 않는다.
    for i in range(6):
        dl = client.get(f"/api/history/{project_id}/download/{job['id']}")
        if dl.status_code == 200:
            break
        time.sleep(0.3 * (i + 1))
    assert dl.status_code == 200 and dl.content[:2] == b"PK", (dl.status_code, dl.text[:200])
    assert dl.headers["content-type"] == server.HWPX_MIME, dl.headers
    assert "merged.hwpx" in dl.headers.get("content-disposition", ""), dl.headers

    def ids(**params):
        qs = "&".join(f"{k}={v}" for k, v in params.items())
        return [p["id"] for p in client.get(f"/api/history?{qs}").json()["projects"]]

    assert project_id in ids(type="hwpx"), "한글 유형에 잡혀야 한다"
    assert project_id not in ids(type="excel"), "엑셀 유형에는 안 잡혀야 한다"
    assert project_id in ids(q="한글"), "작업명으로 찾아야 한다"
    print("  ✓ 한글 병합 영속화 + 이력 유형 필터(엑셀/한글 구분)")


def _form_spec() -> dict:
    """모든 자료 유형을 한 번에 쓰는 사양.

    유형을 골고루 넣는 이유는 field_rules.rule_type 제약이 formgen.FIELD_TYPES보다
    좁으면 저장이 23514로 죽기 때문이다. 실제로 그랬다 — 베이스라인 제약이 4종만
    허용해 text·number를 쓰는 양식(즉 대부분)이 전부 실패했다.
    """
    import formgen as fg
    fields = []
    for t in fg.FIELD_TYPES:
        f = {"name": f"항목-{t}", "type": t, "required": t != "text", "notes": None}
        # date·amount는 입력 형식이 없으면 완결성 루브릭이 사양을 미완으로 본다
        f["format"] = "YYYY-MM-DD" if t == "date" else ("숫자만" if t == "amount" else None)
        fields.append(f)
    return {"form_title": "zz-test 전 유형 양식", "fields": fields, "locale": "ko"}


def _form_workbook(spec: dict) -> dict:
    """사양대로 머리행만 있는 최소 저작물. 검증 관문을 통과하는 형태여야 한다."""
    cells = [{"ref": "A1", "value": spec["form_title"], "style": "title"}]
    for i, f in enumerate(spec["fields"]):
        cells.append({"ref": f"{chr(ord('A') + i)}3",
                      "value": f["name"] + ("*" if f["required"] else ""), "style": "header"})
    return {"sheets": [{
        "name": "양식", "freeze_panes": "A4",
        "styles": {"title": {"bold": True, "size": 14}, "header": {"bold": True, "bg": "EEF2FF"}},
        "cells": cells,
    }]}


def test_form_persistence(client: TestClient):
    """F1 문답·양식·작성기준·동의가 DB에 남고, 남의 것은 보이지 않아야 한다.

    LLM만 스텁으로 바꾸고 **store는 실제 Supabase를 탄다** — 스텁 store로는 스키마
    제약을 지나가지 않아 test_formgen.py가 못 잡는 지점이 여기다.
    """
    import store
    import formgen as fg
    from test_formgen import Stub

    spec = _form_spec()
    real_client = fg._client
    # 인스턴스는 하나만 둔다 — 호출마다 새로 만들면 매번 첫 응답만 돌려준다
    stub = Stub([
        # 1턴: LLM이 완료를 선언해도 사양이 비어 있으면 완료로 인정하지 않는다
        {"reply": "어떤 항목이 필요하신가요?", "spec_json": {"form_title": spec["form_title"]},
         "spec_complete": True},
        # 2턴: 사양 확정
        {"reply": "사양을 확정했습니다.", "spec_json": spec, "spec_complete": True},
        # 저작
        _form_workbook(spec),
    ])
    fg._client = lambda: stub
    try:
        # 동의 전에는 F1 전체가 403이어야 한다 (명세 §9.1)
        assert client.get("/api/form/consent").json()["consented"] is False
        blocked = client.post("/api/form/session", json={"message": "양식 만들어줘"})
        assert blocked.status_code == 403 and blocked.json()["detail"] == "AI_CONSENT_REQUIRED", \
            (blocked.status_code, blocked.text[:200])

        assert client.post("/api/form/consent").status_code == 200
        assert client.get("/api/form/consent").json()["consented"] is True
        me = client.get("/api/auth/me").json()["profile"]["id"]
        # 동의는 1회만 기록한다 (§9.1). 다시 눌러도 시각이 밀리거나 로그가 쌓이면 안 된다
        first_at = store._rest("GET", "/profiles", params={
            "id": f"eq.{me}", "select": "ai_consent_at"}).json()[0]["ai_consent_at"]
        assert client.post("/api/form/consent").status_code == 200
        again = store._rest("GET", "/profiles", params={
            "id": f"eq.{me}", "select": "ai_consent_at"}).json()[0]["ai_consent_at"]
        assert again == first_at, (first_at, again)
        assert len([r for r in store.list_logs(limit=50)
                    if r["actor"] == EMP and r["action"] == "AI 전송 동의"]) == 1
        rows = store._rest("GET", "/profiles",
                           params={"id": f"eq.{me}", "select": "ai_consent_at"}).json()
        assert rows[0]["ai_consent_at"], "동의 시각이 남아야 한다"

        # 1턴 — 완결성 가드가 성급한 완료 선언을 꺾는지
        first = client.post("/api/form/session", json={"message": "전 유형 양식 만들어줘"})
        assert first.status_code == 200, first.text
        made = first.json()
        sid, project_id = made["session_id"], made["project_id"]
        assert made["spec_complete"] is False and "포함할 항목" in made["gaps"], made

        # 2턴 — 사양 확정. 문답 원문이 DB에 누적돼야 한다
        second = client.post(f"/api/form/{sid}/messages", json={"message": "전 유형 다 넣어주세요"})
        assert second.status_code == 200, second.text
        assert second.json()["spec_complete"] is True, second.json()
        saved = store.get_intake_session(sid)
        assert saved["status"] == "spec_complete" and saved["turn_count"] == 2, saved
        assert [m["role"] for m in saved["messages_json"]] == \
            ["user", "assistant", "user", "assistant"], saved["messages_json"]
        assert saved["spec_json"]["form_title"] == spec["form_title"], saved["spec_json"]

        # 생성 — 양식 파일·작성기준·상태가 남아야 한다
        out = wait(client.post(f"/api/form/{sid}/generate"))
        template_id = out["template_id"]
        row = store.get_form_template(template_id)
        assert row["status"] == "done" and row["error_message"] is None, row
        assert row["version"] == 1 and row["output_format"] == "xlsx", row
        assert row["intake_session_id"] == sid, row
        assert row["workbook_json"]["sheets"][0]["name"] == "양식", row["workbook_json"]
        assert store.get_intake_session(sid)["status"] == "closed", "생성 후 문답은 닫힌다"

        # 생성물이 Storage에 실제로 있는지
        assert _object_of(store.RESULT_BUCKET, row["file_url"])[:2] == b"PK", row["file_url"]
        # 계정 삭제 시 정리되는 자리에 있어야 한다 — 프로젝트 폴더 바로 아래, 평면.
        # 종전 forms/{project_id}/... 는 정리를 빠져나가 계정을 지워도 남았다.
        assert row["file_url"].startswith(f"{project_id}/"), row["file_url"]
        assert "/" not in row["file_url"][len(project_id) + 1:], row["file_url"]

        # 작성기준 — rule_type이 formgen의 어휘 전부를 담을 수 있어야 한다
        saved_rules = store._rest("GET", "/field_rules",
                                  params={"form_template_id": f"eq.{template_id}",
                                          "select": "field_name,rule_type,rule_config_json"}).json()
        assert {r["rule_type"] for r in saved_rules} == set(fg.FIELD_TYPES), saved_rules
        text_rule = next(r for r in saved_rules if r["rule_type"] == "text")
        assert text_rule["rule_config_json"]["required"] is False, text_rule

        # 목록·작성기준 API가 취합이 그대로 먹는 형태로 내려주는지 (F1-6)
        listed = client.get("/api/form/templates").json()["templates"]
        mine = next((t for t in listed if t["id"] == template_id), None)
        assert mine and mine["field_count"] == len(spec["fields"]), listed
        rules = client.get(f"/api/form/templates/{template_id}/rules").json()["rules"]
        assert rules["항목-text"] == {"required": False}, rules
        assert all("key" not in v for v in rules.values()), "키 컬럼은 지어내지 않는다"

        dl = client.get(f"/api/form/template/{template_id}/download")
        assert dl.status_code == 200 and dl.content[:2] == b"PK", dl.status_code
        v1 = dl.content

        # F1-7 대화형 수정 — 사양을 고쳐 다시 저작하고 버전을 올린다
        assert client.post(f"/api/form/{sid}/revise", json={"message": " "}).status_code == 400
        trimmed = {**spec, "fields": [f for f in spec["fields"] if f["type"] != "text"]}
        # 고칠 내용이 확정되지 않으면 파일을 건드리지 않고 되묻는다
        stub.replies.append({"reply": "어떤 항목을 뺄까요?", "spec_json": {}, "spec_complete": True})
        asked = client.post(f"/api/form/{sid}/revise", json={"message": "하나 빼줘"})
        assert asked.status_code == 200, asked.text
        assert asked.json().get("job_id") is None and asked.json()["question"], asked.json()
        assert store.get_form_template(template_id)["version"] == 1, "되물을 때 버전이 올라갔다"

        stub.replies.append({"reply": "항목-text를 뺐습니다", "spec_json": trimmed,
                             "spec_complete": True})
        stub.replies.append(_form_workbook(trimmed))
        revised = wait(client.post(f"/api/form/{sid}/revise",
                                   json={"message": "항목-text 항목은 빼줘"}))
        assert revised["version"] == 2, revised
        row2 = store.get_form_template(template_id)
        assert row2["version"] == 2 and row2["file_url"].endswith("_v2.xlsx"), row2
        assert len(row2["spec_json"]["fields"]) == len(spec["fields"]) - 1, row2["spec_json"]
        assert row2["workbook_json"] != row["workbook_json"], "저작물이 갱신되지 않았다"
        # 작성기준도 함께 갈아끼워야 한다 — 파일에 없는 항목이 기준에 남으면 취합이 오판한다
        after = store._rest("GET", "/field_rules",
                            params={"form_template_id": f"eq.{template_id}",
                                    "select": "field_name,rule_type"}).json()
        assert "text" not in {r["rule_type"] for r in after}, after
        assert len(after) == len(spec["fields"]) - 1, after
        # 이전 버전 파일은 남는다(버전별 경로)
        assert _object_of(store.RESULT_BUCKET, row["file_url"])[:2] == b"PK", row["file_url"]
        dl2 = client.get(f"/api/form/template/{template_id}/download")
        assert dl2.status_code == 200 and dl2.content != v1, "최신 버전이 내려오지 않았다"
        assert "양식 수정 완료" in {r["action"] for r in store.list_logs(limit=50)
                                if r["actor"] == EMP}

        # 감사 로그에 두 액션이 남았는지
        # list_logs는 actor_id 대신 사번(actor)을 내려준다(화면이 쓰는 형태)
        actions = {r["action"] for r in store.list_logs(limit=50) if r["actor"] == EMP}
        assert {"AI 전송 동의", "양식 생성 완료"} <= actions, actions

        # 남의 문답·양식은 존재 자체를 알리지 않는다 (§6.2)
        other_emp = f"zz-test-{uuid.uuid4().hex[:10]}"
        other = TestClient(server.app)
        other_id = other.post("/api/auth/signup", json={
            "employee_no": other_emp, "password": PW,
            "reset_email": RESET_EMAIL}).json()["profile"]["id"]
        try:
            assert other.post("/api/auth/login",
                              json={"employee_no": other_emp, "password": PW}).status_code == 200
            # 동의 게이트(403)는 위에서 확인했다. 여기서 보려는 것은 격리라 동의를 준다
            assert other.post("/api/form/consent").status_code == 200
            for method, path in [("get", f"/api/form/{sid}"),
                                 ("post", f"/api/form/{sid}/messages"),
                                 ("get", f"/api/form/templates/{template_id}/rules"),
                                 ("get", f"/api/form/template/{template_id}/download")]:
                res = (other.post(path, json={"message": "보여줘"}) if method == "post"
                       else other.get(path))
                assert res.status_code == 404, f"{path} → {res.status_code} {res.text[:200]}"
            assert other.get("/api/form/templates").json()["templates"] == []
        finally:
            _delete_storage_for_owner(other_id)
            _delete_user(other_id)
        print("  ✓ F1 영속화(문답·양식·작성기준·동의 + 대화형 수정 v2 + 양식 격리)")
        return project_id
    finally:
        fg._client = real_client


def test_form_recognize(client: TestClient):
    """F6 기존 양식 인지·등록 — 원본을 고치지 않고 등록하고, 작성기준이 남아야 한다.

    여기서만 잡히는 것이 있다 — 등록은 첨부 객체를 **그대로** 양식 파일로 삼으므로
    Storage 경로 규칙(계정 삭제 시 정리되는 자리)과 바이트 무변경을 실제 Supabase를
    상대로 확인해야 한다.
    """
    import shutil
    import tempfile
    from pathlib import Path

    import store
    import formgen as fg
    from test_formgen import Stub

    tmp = Path(tempfile.mkdtemp(prefix="f6-test-"))
    real_client = fg._client
    try:
        # 배포 전 빈 양식(헤더 아래가 전부 비어 있는 실제 형태)을 만들어 그것을 올린다
        spec = _form_spec()
        form_path = tmp / "zz-test 기존양식.xlsx"
        fg.materialize(_form_workbook(spec), form_path)
        original = form_path.read_bytes()

        # 헤더는 `*`가 붙은 그대로다. 항목명이 헤더와 한 글자라도 다르면 취합에서
        # 그 열을 못 찾으므로, 인지 사양의 name은 헤더 문자열이어야 한다
        headers = [f["name"] + ("*" if f["required"] else "") for f in spec["fields"]]
        recognized = {"form_title": spec["form_title"], "locale": "ko",
                      "selected_sheets": ["양식"],
                      "fields": [dict(f, name=h, confidence="high")
                                 for f, h in zip(spec["fields"], headers)]}

        # png는 아예 받지 않는다(인지 §6.4)
        bad = client.post("/api/form/attachments",
                          files=[("files", ("그림.png", b"\x89PNG\r\n", "image/png"))])
        assert bad.status_code == 400 and "xlsx" in bad.json()["detail"], bad.text

        # 등록 대상 xlsx + 참고용 hwpx를 함께 올린다
        up = client.post("/api/form/attachments", files=[
            ("files", (form_path.name, original, server.XLSX_MIME)),
            ("files", ("참고.hwpx", hwpx_bytes("ref"), server.HWPX_MIME)),
        ])
        assert up.status_code == 200, up.text
        made = up.json()
        sid, project_id = made["session_id"], made["project_id"]
        assert made["mode"] == "recognize", made
        xlsx_att = next(a for a in made["attachments"] if a["kind"] == "xlsx")
        hwpx_att = next(a for a in made["attachments"] if a["kind"] == "hwpx")
        assert xlsx_att["primary"] is True and xlsx_att["tables"][0]["headers"] == headers, xlsx_att
        assert hwpx_att["primary"] is False and "참고 자료로만" in hwpx_att["note"], hwpx_att

        # 문답 — 1턴은 항목이 빠져 있어 완료로 인정되지 않는다(헤더 1:1 루브릭)
        short = {k: v for k, v in recognized.items()}
        short["fields"] = recognized["fields"][:-1]
        stub = Stub([
            {"intent": "recognize", "reply": "이렇게 읽었습니다", "spec_json": short,
             "recognize_complete": True},
            {"intent": "recognize", "reply": "확정했습니다", "spec_json": recognized,
             "recognize_complete": True},
        ])
        fg._client = lambda: stub

        first = client.post(f"/api/form/{sid}/messages",
                           json={"message": "이 양식 그대로 받을 거야"})
        assert first.status_code == 200, first.text
        assert first.json()["mode"] == "recognize" and first.json()["intent"] == "recognize"
        assert first.json()["spec_complete"] is False, first.json()
        assert any(headers[-1] in g for g in first.json()["gaps"]), first.json()["gaps"]
        # 확정되지 않은 상태에서 등록을 직접 부르면 막힌다
        assert client.post(f"/api/form/{sid}/register").status_code == 400

        second = client.post(f"/api/form/{sid}/messages", json={"message": "비고만 선택이야"})
        assert second.status_code == 200 and second.json()["spec_complete"] is True, second.text
        # 새로고침 복원 — 모드·첨부·사양이 되살아나야 한다
        restored = client.get(f"/api/form/{sid}").json()
        assert restored["mode"] == "recognize" and restored["gaps"] == [], restored
        assert len(restored["attachments"]) == 2, restored["attachments"]

        # 등록 — 동기 응답이다(잡·폴링 없음)
        reg = client.post(f"/api/form/{sid}/register")
        assert reg.status_code == 200, reg.text
        out = reg.json()
        template_id = out["template_id"]
        assert out["source"] == "recognized_external" and out["sheets"][0]["name"] == "양식", out
        assert out["rules"][headers[0]] == {"required": spec["fields"][0]["required"]}, out["rules"]

        row = store.get_form_template(template_id)
        assert row["status"] == "done" and row["source"] == "recognized_external", row
        # 원본이 산출물이므로 저작물 명세는 없는 것이 정상이다(E1)
        assert row["workbook_json"] is None, row["workbook_json"]
        assert row["intake_session_id"] == sid, row
        # Storage 경로는 프로젝트 폴더 바로 아래 평면 — 계정 삭제 시 정리되는 자리다
        assert row["file_url"].startswith(f"{project_id}/"), row["file_url"]
        assert "/" not in row["file_url"][len(project_id) + 1:], row["file_url"]

        # **무변경 등록**(E1의 핵심 보증) — 내려받은 바이트가 올린 것과 같아야 한다
        dl = client.get(f"/api/form/template/{template_id}/download")
        assert dl.status_code == 200 and dl.content == original, \
            (dl.status_code, len(dl.content), len(original))

        # 작성기준이 남고, 취합이 그대로 먹는 형태로 내려온다 (F6-6)
        saved = store._rest("GET", "/field_rules",
                            params={"form_template_id": f"eq.{template_id}",
                                    "select": "field_name,rule_type"}).json()
        assert {r["field_name"] for r in saved} == set(headers), saved
        rules = client.get(f"/api/form/templates/{template_id}/rules").json()["rules"]
        assert set(rules) == set(headers), rules
        listed = client.get("/api/form/templates").json()["templates"]
        mine = next(t for t in listed if t["id"] == template_id)
        assert mine["source"] == "recognized_external" and mine["field_count"] == len(headers), mine

        assert store.get_intake_session(sid)["status"] == "closed", "등록 후 문답은 닫힌다"
        assert client.post(f"/api/form/{sid}/register").status_code == 400, "두 번 등록되면 안 된다"
        # 등록 원본은 재저작 대상이 아니다(인지 E3) — F1-7이 손대면 원본 보존이 깨진다
        locked = client.post(f"/api/form/{sid}/revise", json={"message": "항목 하나 빼줘"})
        assert locked.status_code == 400 and "등록된 원본" in locked.json()["detail"], locked.text
        actions = {r["action"] for r in store.list_logs(limit=50) if r["actor"] == EMP}
        assert "양식 등록 완료" in actions, actions
        print("  ✓ F6 인지·등록(무변경 등록·헤더 1:1·작성기준·첨부 구분)")
    finally:
        fg._client = real_client
        shutil.rmtree(tmp, ignore_errors=True)


def test_admin_console(client: TestClient):
    """Admin API는 관리자만 쓸 수 있고, 변경은 감사 로그에 남아야 한다 (F4-3)."""
    import manage_users
    import store

    # 일반 사용자에게는 403
    for path in ("/api/admin/dashboard", "/api/admin/accounts", "/api/admin/logs"):
        assert client.get(path).status_code == 403, f"{path}가 일반 사용자에게 열렸다"

    assert manage_users.main(["promote", EMP]) == 0
    try:
        dash = client.get("/api/admin/dashboard")
        assert dash.status_code == 200, dash.text
        d = dash.json()
        assert d["users_total"] >= 1 and d["users_active"] >= 1, d
        assert d["retention_days"] == 90, d
        assert d["api_cost"] is None, "집계 경로가 없으면 0이 아니라 None이어야 한다"
        assert isinstance(d["storage_bytes"], int), d

        accounts = client.get("/api/admin/accounts").json()["accounts"]
        me = next(a for a in accounts if a["employee_no"] == EMP)
        assert me["role"] == "admin", me

        # 본인 계정을 스스로 잠그는 것은 막는다(CLI 말고는 되돌릴 길이 없다)
        assert client.patch(f"/api/admin/accounts/{me['id']}",
                            json={"status": "suspended"}).status_code == 400
        assert client.delete(f"/api/admin/accounts/{me['id']}").status_code == 400

        # 다른 계정은 정지/권한 변경이 된다
        other_emp = f"zz-test-{uuid.uuid4().hex[:10]}"
        made = TestClient(server.app).post(
            "/api/auth/signup",
            json={"employee_no": other_emp, "password": PW, "reset_email": RESET_EMAIL})
        other_id = made.json()["profile"]["id"]
        patched = client.patch(f"/api/admin/accounts/{other_id}", json={"status": "suspended"})
        assert patched.status_code == 200 and patched.json()["account"]["status"] == "suspended"

        # 보관 정책
        bad = client.put("/api/admin/policy/retention", json={"retention_days": 0})
        assert bad.status_code == 400, bad.text
        ok = client.put("/api/admin/policy/retention", json={"retention_days": 120})
        assert ok.status_code == 200 and ok.json()["policy"]["retention_days"] == 120, ok.text

        # 감사 로그를 남긴 계정도 지워져야 한다. audit_logs.actor_id가 ON DELETE
        # 절 없이 profiles를 참조하던 동안에는 '뭔가를 한 사용자'가 삭제 불가였다
        store.log_action(other_id, "테스트 행위", "삭제 가능 여부 확인")
        logged = [x for x in client.get("/api/admin/logs").json()["logs"]
                  if x["action"] == "테스트 행위"]
        assert logged, "로그가 남지 않아 이 검사가 무의미하다"

        # 삭제 + Storage까지 정리되는지
        deleted = client.delete(f"/api/admin/accounts/{other_id}")
        assert deleted.status_code == 200 and deleted.json()["deleted"] == other_emp, deleted.text
        assert all(a["employee_no"] != other_emp
                   for a in client.get("/api/admin/accounts").json()["accounts"])

        # 계정이 사라져도 로그는 남고, 행위자만 비워진다(감사 기록의 목적)
        after = [x for x in client.get("/api/admin/logs").json()["logs"]
                 if x["action"] == "테스트 행위"]
        assert after and after[0]["actor"] == "(삭제된 계정)", after

        # 감사 로그에 방금 한 일이 남아야 한다
        logs = client.get("/api/admin/logs").json()["logs"]
        actions = [x["action"] for x in logs]
        assert "계정 변경" in actions and "계정 삭제" in actions and "정책 변경" in actions, actions
        assert all(x["actor"] for x in logs), logs[:3]
        # 검색 필터
        only = client.get("/api/admin/logs?q=정책").json()["logs"]
        assert only and all("정책" in x["action"] or "정책" in (x["target"] or "") for x in only), only
        print("  ✓ Admin 콘솔(권한 403 / 지표 / 계정 변경·삭제 / 정책 / 감사 로그)")
    finally:
        client.put("/api/admin/policy/retention", json={"retention_days": 90})
        manage_users.main(["demote", EMP])


def test_public_key_reads_nothing():
    """프론트엔드에 실려 나가는 공개 키로는 아무 데이터도 읽히지 않아야 한다.

    RLS는 켜져 있고 정책은 0개다 — 모든 접근이 백엔드(service key)를 경유하는
    현 구조에서 이게 가장 안전한 상태다. 정책을 실수로 추가하거나 RLS를 끄면
    여기서 걸린다.
    """
    base = os.environ["SUPABASE_URL"].rstrip("/")
    pub = os.environ["SUPABASE_PUBLISHABLE_KEY"]
    headers = {"apikey": pub, "Authorization": f"Bearer {pub}"}
    for table in ("profiles", "projects", "uploaded_files", "review_results", "aggregation_jobs"):
        res = httpx.get(f"{base}/rest/v1/{table}", headers=headers,
                        params={"select": "*"}, timeout=20)
        # RLS가 행을 감추면 PostgREST는 403이 아니라 빈 배열을 준다
        assert res.status_code in (200, 401, 403), f"{table} → {res.status_code}"
        if res.status_code == 200:
            assert res.json() == [], f"{table}이 공개 키에 노출됐다: {res.text[:200]}"
    print("  ✓ 공개 키로 데이터 접근 불가(RLS 전면 거부 유지)")


def test_logout(client: TestClient):
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/auth/me").status_code == 401, "로그아웃 후에도 접근되면 안 된다"
    print("  ✓ 로그아웃")


def _delete_user(user_id: str | None) -> None:
    if not user_id:
        return
    key = auth._secret_key()
    headers = {"apikey": key, "Authorization": f"Bearer {key}"}
    # profiles는 auth.users FK가 ON DELETE CASCADE라 함께 지워진다.
    # 실패를 조용히 넘기면 테스트 계정이 쌓이므로 상태 코드를 확인한다 —
    # 실제로 audit_logs FK 때문에 삭제가 막히던 것을 이걸 안 봐서 놓쳤다.
    res = httpx.delete(f"{os.environ['SUPABASE_URL']}/auth/v1/admin/users/{user_id}",
                       headers=headers, timeout=20)
    assert res.status_code < 400, f"테스트 계정 삭제 실패: {res.status_code} {res.text[:300]}"


def _delete_test_logs() -> None:
    """테스트가 남긴 감사 로그를 지운다.

    행위자 계정을 지우면 로그는 actor_id=NULL로 남는다(감사 기록의 목적).
    실제 운영 기록과 섞이지 않게 테스트 흔적은 걷어낸다.
    """
    key = auth._secret_key()
    headers = {"apikey": key, "Authorization": f"Bearer {key}",
               "Content-Type": "application/json"}
    base = os.environ["SUPABASE_URL"].rstrip("/")
    for action in ("테스트 행위", "계정 변경", "계정 삭제", "정책 변경", "취합 완료",
                   "한글 병합 완료", "AI 전송 동의", "양식 생성 완료",
                   "양식 등록 완료", "양식 수정 완료"):
        httpx.delete(f"{base}/rest/v1/audit_logs", headers=headers, timeout=20,
                     params={"actor_id": "is.null", "action": f"eq.{action}"})


def _delete_storage_for_owner(owner_id: str | None) -> None:
    """이 사용자의 모든 프로젝트 폴더를 정리한다.

    DB는 profiles → projects → uploaded_files가 CASCADE로 지워지지만
    Storage 객체는 대상이 아니라 남는다. 사용자를 지우기 전에 훑어야
    프로젝트 목록을 알 수 있다.
    """
    if not owner_id:
        return
    key = auth._secret_key()
    headers = {"apikey": key, "Authorization": f"Bearer {key}",
               "Content-Type": "application/json"}
    base = os.environ["SUPABASE_URL"].rstrip("/")
    rows = httpx.get(f"{base}/rest/v1/projects", headers=headers, timeout=20,
                     params={"owner_id": f"eq.{owner_id}", "select": "id"})
    for row in (rows.json() if rows.status_code == 200 else []):
        _delete_storage(row["id"])


def _delete_storage(project_id: str | None) -> None:
    if not project_id:
        return
    key = auth._secret_key()
    headers = {"apikey": key, "Authorization": f"Bearer {key}",
               "Content-Type": "application/json"}
    base = os.environ["SUPABASE_URL"].rstrip("/")
    for bucket in ("uploads", "results"):
        listed = httpx.post(f"{base}/storage/v1/object/list/{bucket}", headers=headers, timeout=20,
                            json={"prefix": f"{project_id}/", "limit": 200})
        if listed.status_code >= 400:
            continue
        names = [f"{project_id}/{o['name']}" for o in listed.json()]
        if names:
            httpx.request("DELETE", f"{base}/storage/v1/object/{bucket}", headers=headers,
                          json={"prefixes": names}, timeout=30)


def main() -> int:
    if not auth.configured():
        print("건너뜀: .env에 SUPABASE_URL·SUPABASE_PUBLISHABLE_KEY·SUPABASE_SECRET_KEY가 필요합니다.")
        return 0
    try:
        test_protected_without_login()
        client = test_signup_and_login()
        sid = test_authenticated_flow(client)
        test_isolation(sid)
        test_role_and_suspend(client)
        test_persistence(client)
        test_hwpx_persistence(client)
        test_form_persistence(client)
        test_form_recognize(client)
        test_admin_console(client)
        test_public_key_reads_nothing()
        test_logout(client)
    finally:
        # 사용자를 지우면 프로젝트도 CASCADE로 사라지므로 Storage를 먼저 훑는다
        _delete_storage_for_owner(created_user_id)
        _delete_user(created_user_id)
        _delete_test_logs()
    print("\n전체 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
