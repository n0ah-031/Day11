"""F1 양식 생성 자체 점검. 실행: python3 test_formgen.py

LLM 호출은 스텁으로 바꿔 **네트워크·비용 없이** 돈다. 확인하는 것은 우리 쪽 로직이다 —
완결성 가드가 LLM의 성급한 완료 선언을 꺾는지, 검증 관문이 위험한 산출물을 막는지,
변환이 저작물을 그대로 옮기는지, 작성기준이 취합 엔진에 맞는 형태로 나오는지.

프롬프트 품질과 실제 응답은 스텁으로 확인할 수 없다 — 그건 실제 키로 돌려봐야 하고,
결과는 HANDOFF §4에 기록한다.
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

import formgen as fg

GOOD_SPEC = {
    "form_title": "부서별 예산 집행 현황",
    "target_depts": ["전체 부서"],
    "fields": [
        {"name": "부서명", "type": "dept_code", "required": True, "format": None, "notes": None},
        {"name": "집행일자", "type": "date", "required": True, "format": "YYYY-MM-DD", "notes": None},
        {"name": "집행금액", "type": "amount", "required": True, "format": "숫자만", "notes": None},
        {"name": "비고", "type": "text", "required": False, "format": None, "notes": None},
    ],
    "layout_hints": None,
    "locale": "ko",
}

GOOD_WB = {
    "sheets": [{
        "name": "예산집행",
        "columns": [{"index": 1, "width": 16}, {"index": 2, "width": 14}],
        "merges": ["A1:D1"],
        "freeze_panes": "A4",
        "styles": {
            "title": {"bold": True, "size": 14, "align": "center"},
            "header": {"bold": True, "bg": "EEF2FF", "border": "thin", "align": "center"},
        },
        "cells": [
            {"ref": "A1", "value": "부서별 예산 집행 현황", "style": "title"},
            {"ref": "A3", "value": "부서명*", "style": "header"},
            {"ref": "B3", "value": "집행일자*", "style": "header"},
            {"ref": "C3", "value": "집행금액*", "style": "header"},
            {"ref": "D3", "value": "비고", "style": "header"},
            {"ref": "C104", "value": "=SUM(C4:C103)", "style": "header"},
        ],
        "validations": [
            {"range": "B4:B103", "type": "date", "prompt": "YYYY-MM-DD", "error": "날짜 형식 오류"},
            {"range": "A4:A103", "type": "list", "source": ["기획팀", "인사팀"]},
        ],
    }]
}


class Stub:
    """chat.completions.create 인터페이스만 흉내내는 스텁. 정해둔 응답을 순서대로 준다."""

    def __init__(self, replies: list[dict]):
        self.replies = list(replies)
        self.seen: list[str] = []          # 전송된 payload — 마스킹 확인용
        self.chat = self

    @property
    def completions(self):
        return self

    def create(self, model, response_format, messages):
        self.seen.append(messages[-1]["content"])
        body = self.replies.pop(0)

        class R:
            choices = [type("C", (), {"message": type("M", (), {"content": json.dumps(body, ensure_ascii=False)})()})()]
        return R()


def test_rubric_guard():
    """LLM이 완료라고 해도 사양에 빈 곳이 있으면 완료로 인정하지 않는다 (§4.3)."""
    assert fg.rubric_gaps(GOOD_SPEC) == [], fg.rubric_gaps(GOOD_SPEC)

    # 형식이 빠진 날짜 항목 / 필수여부가 없는 항목 / 유형이 틀린 항목
    holed = {"form_title": "x", "fields": [
        {"name": "일자", "type": "date", "required": True},                 # format 없음
        {"name": "금액", "type": "amount", "required": True, "format": ""},  # format 빈 문자열
        {"name": "메모", "type": "text"},                                    # required 없음
        {"name": "코드", "type": "이상한것", "required": False},              # 유형 미정의
    ]}
    gaps = fg.rubric_gaps(holed)
    assert "'일자'의 입력 형식" in gaps and "'금액'의 입력 형식" in gaps, gaps
    assert "'메모'의 필수/선택 구분" in gaps and "'코드'의 자료 유형" in gaps, gaps
    assert fg.rubric_gaps({}) == ["양식 이름", "포함할 항목"], fg.rubric_gaps({})

    # 성급한 완료 선언 → 꺾이고, 부족 항목이 남은 주제로 되돌아간다
    stub = Stub([{"reply": "다 됐습니다", "spec_json": holed, "spec_complete": True,
                  "coverage": {"confirmed_topics": ["양식명"], "remaining_topics": []}}])
    turn = fg.intake_turn([{"role": "user", "content": "예산 양식"}], {}, "m", client=stub)
    assert turn["spec_complete"] is False, turn
    assert "'일자'의 입력 형식" in turn["coverage"]["remaining_topics"], turn["coverage"]

    # 완결된 사양이면 통과한다
    stub = Stub([{"reply": "확정했습니다", "spec_json": GOOD_SPEC, "spec_complete": True,
                  "coverage": {"confirmed_topics": ["전부"], "remaining_topics": []}}])
    turn = fg.intake_turn([{"role": "user", "content": "예산 양식"}], {}, "m", client=stub)
    assert turn["spec_complete"] is True and turn["gaps"] == [], turn
    print("  ✓ 완결성 루브릭 가드(성급한 완료 선언을 꺾고 부족 항목을 되돌림)")


def test_turn_limit_and_masking():
    """턴 상한에서 마감 지시가 붙고, 전송 payload는 마스킹을 통과한다."""
    stub = Stub([{"reply": "정리했습니다", "spec_json": GOOD_SPEC, "spec_complete": True,
                  "coverage": {}}])
    long_talk = [{"role": "user", "content": f"답변 {i}"} for i in range(fg.TURN_LIMIT)]
    fg.intake_turn(long_talk, {}, "m", client=stub)
    assert "턴 수 상한" in stub.seen[0], stub.seen[0][:200]

    # 주민번호·전화·이메일은 전송 전에 치환된다 (§9.2)
    stub = Stub([{"reply": "확인했습니다", "spec_json": {}, "spec_complete": False, "coverage": {}}])
    fg.intake_turn([{"role": "user",
                     "content": "담당자 900101-1234567, 010-1234-5678, a@b.com 로 보내주세요"}],
                   {}, "m", client=stub)
    sent = stub.seen[0]
    assert "900101-1234567" not in sent and "010-1234-5678" not in sent, sent
    assert "a@b.com" not in sent, sent
    print("  ✓ 턴 상한 마감 지시 + 전송 payload 마스킹")


def test_validation_gate():
    """검증 관문이 위험한·깨진 저작물을 막는다 (§6.3)."""
    assert fg.validate_workbook(GOOD_WB) == [], fg.validate_workbook(GOOD_WB)

    def one(sheet_patch: dict) -> list[str]:
        wb = json.loads(json.dumps(GOOD_WB))
        wb["sheets"][0].update(sheet_patch)
        return fg.validate_workbook(wb)

    # 수식 주입 — 허용 함수만
    assert any("허용되지 않은 함수" in p for p in
               one({"cells": [{"ref": "A1", "value": "=WEBSERVICE(\"http://x\")"}]}))
    assert any("다른 파일" in p for p in
               one({"cells": [{"ref": "A1", "value": "=[다른.xlsx]Sheet1!A1"}]}))
    assert one({"cells": [{"ref": "A1", "value": "=SUMIFS(C4:C9,A4:A9,\"x\")"}]}) == []

    # 표기·상한
    assert any("셀 위치" in p for p in one({"cells": [{"ref": "A", "value": 1}]}))
    assert any("병합 범위" in p for p in one({"merges": ["A1"]}))
    assert any("freeze_panes" in p for p in one({"freeze_panes": "여기"}))
    assert any("31자" in p for p in one({"name": "가" * 32}))
    assert any("쓸 수 없는 문자" in p for p in one({"name": "예산/집행"}))
    assert any("허용되지 않은 속성" in p for p in
               one({"styles": {"title": {"bold": True, "제멋대로": 1}}}))
    assert any("유효성 검사 종류" in p for p in
               one({"validations": [{"range": "A1:A2", "type": "지원안됨"}]}))
    assert any("드롭다운 목록이 비었" in p for p in
               one({"validations": [{"range": "A1:A2", "type": "list", "source": []}]}))

    big = {"sheets": [{"name": "s", "cells": [{"ref": f"A{i + 1}", "value": i}
                                              for i in range(fg.MAX_CELLS + 1)]}]}
    assert any("최대" in p for p in fg.validate_workbook(big))
    dupes = {"sheets": [{"name": "같음", "cells": []}, {"name": "같음", "cells": []}]}
    assert any("중복" in p for p in fg.validate_workbook(dupes))
    assert fg.validate_workbook({"sheets": []}) == ["시트가 없습니다"]
    print("  ✓ 검증 관문(수식 주입·외부 참조·표기·상한·중복·스타일 속성)")


def test_normalize():
    """모델이 쓰는 다른 이름을 표준 키로 받는다(실측에서 나온 실제 변형).

    `values`/`integer`는 이름만 다른 같은 제약이다. 이걸 안 받으면 검증에서 막혀 재저작을
    소진하고 사용자에게는 아무 파일도 안 나간다.
    """
    wb = {"sheets": [{"name": "s", "cells": [], "validations": [
        {"range": "A4:A9", "type": "list", "values": ["정기", "수시"]},
        {"range": "B4:B9", "type": "integer"},
        {"range": "C4:C9", "type": "list", "options": ["가", "나"]},
    ]}]}
    assert fg.validate_workbook(json.loads(json.dumps(wb))), "정규화 전에는 걸려야 한다"
    dv = fg.normalize_workbook(wb)["sheets"][0]["validations"]
    assert dv[0]["source"] == ["정기", "수시"] and dv[1]["type"] == "whole"
    assert dv[2]["source"] == ["가", "나"], dv[2]
    assert fg.validate_workbook(wb) == [], fg.validate_workbook(wb)

    # 없는 목록을 만들어내지는 않는다 — source도 별칭도 없으면 그대로 걸린다
    empty = {"sheets": [{"name": "s", "cells": [],
                         "validations": [{"range": "A1:A2", "type": "list"}]}]}
    fg.normalize_workbook(empty)
    assert any("드롭다운 목록이 비었" in p for p in fg.validate_workbook(empty))
    print("  ✓ 저작물 키 정규화(values→source, integer→whole)와 빈 목록 거부")


def test_author_retry():
    """검증에 걸리면 위반 목록을 알려 다시 저작하고, 소진되면 사유를 들고 실패한다."""
    broken = json.loads(json.dumps(GOOD_WB))
    broken["sheets"][0]["cells"].append({"ref": "Z1", "value": "=WEBSERVICE(\"http://x\")"})

    stub = Stub([broken, GOOD_WB])
    retries = []
    wb = fg.author_workbook(GOOD_SPEC, "m", client=stub,
                            on_retry=lambda n, p: retries.append((n, p)))
    assert wb == GOOD_WB and len(retries) == 1, retries
    assert any("허용되지 않은 함수" in p for p in retries[0][1]), retries
    assert "fix_these" in stub.seen[1], stub.seen[1][:200]   # 위반 목록이 실제로 전달됐다

    stub = Stub([broken, broken, broken])
    try:
        fg.author_workbook(GOOD_SPEC, "m", client=stub)
        raise AssertionError("검증을 통과하지 못했는데 저작물을 돌려줬다")
    except fg.FormGenError as exc:
        assert "허용되지 않은 함수" in str(exc), exc
    assert len(stub.replies) == 0, "재저작 횟수가 상한과 다르다"
    print("  ✓ 검증 미통과 시 재저작 후 실패 사유 전달")


def test_materialize(tmp: Path):
    """workbook_json이 xlsx로 그대로 옮겨진다(해석·보정 없음)."""
    out = tmp / "form.xlsx"
    assert fg.materialize(GOOD_WB, out) == []      # 낮춰 처리한 것 없음
    ws = openpyxl.load_workbook(out).active
    assert ws.title == "예산집행"
    assert ws["A1"].value == "부서별 예산 집행 현황" and ws["A1"].font.bold
    assert ws["A1"].font.size == 14 and ws["A1"].alignment.horizontal == "center"
    assert ws["A3"].value == "부서명*" and ws["A3"].fill.start_color.rgb.endswith("EEF2FF")
    assert ws["A3"].border.left.style == "thin"
    assert ws["C104"].value == "=SUM(C4:C103)"      # 수식은 수식으로 남는다
    assert "A1:D1" in [str(r) for r in ws.merged_cells.ranges]
    assert ws.freeze_panes == "A4"
    assert ws.column_dimensions["A"].width == 16
    kinds = {dv.type: dv for dv in ws.data_validations.dataValidation}
    assert "list" in kinds and "date" in kinds, kinds
    assert "기획팀" in kinds["list"].formula1

    # 정의 없는 스타일을 참조해도 실패하지 않고 무서식으로 둔다(안전 기본값)
    wb2 = {"sheets": [{"name": "s", "cells": [{"ref": "A1", "value": "x", "style": "없는스타일"}]}]}
    fg.materialize(wb2, tmp / "b.xlsx")
    ws2 = openpyxl.load_workbook(tmp / "b.xlsx").active
    assert ws2["A1"].value == "x" and not ws2["A1"].font.bold
    print("  ✓ materializer 1:1 변환(서식·병합·고정·너비·유효성검사·수식)")


def test_rules_bridge(tmp: Path):
    """작성기준이 취합 엔진(F2)에 그대로 먹는 형태로 나온다 (F1-6)."""
    rows = fg.derive_field_rules(GOOD_SPEC)
    assert [r["field_name"] for r in rows] == ["부서명", "집행일자", "집행금액", "비고"]
    assert rows[1]["rule_type"] == "date"
    assert rows[1]["rule_config_json"] == {"required": True, "format": "YYYY-MM-DD"}
    assert rows[3]["rule_config_json"] == {"required": False}

    rules = fg.rules_for_aggregate(GOOD_SPEC)
    assert rules == {"부서명": {"required": True}, "집행일자": {"required": True},
                     "집행금액": {"required": True}, "비고": {"required": False}}, rules
    # 키 컬럼은 넣지 않는다 — 사양에 근거가 없다
    assert not any("key" in v for v in rules.values()), rules

    # 실제로 취합 엔진이 이 rules를 먹는지 — 비고를 비워둔 회신 파일이 '정상'이어야 한다
    import aggregate as ag
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "예산집행"
    ws.append(["부서명", "집행일자", "집행금액", "비고"])
    ws.append(["기획팀", "2026-08-01", 1000, None])
    ws.append(["인사팀", "2026-08-02", 2000, None])
    path = tmp / "회신.xlsx"
    wb.save(path)

    uf = ag.read_file(path)
    ag.review_stage1(uf, rules)
    assert uf.grade != ag.ERROR, [i.line() for i in uf.issues]
    assert uf.status.startswith("정상"), (uf.status, [i.line() for i in uf.issues])

    # 반대로 필수 항목이 비면 잡아낸다(기준이 실제로 적용되는 증거)
    ws["B3"] = None
    wb.save(path)
    uf2 = ag.read_file(path)
    ag.review_stage1(uf2, rules)
    assert uf2.grade == ag.ERROR, [i.line() for i in uf2.issues]
    print("  ✓ 작성기준 파생 + 취합 엔진 연결(선택 항목 오탐 없음, 필수 누락은 잡음)")


def test_read_structure(tmp: Path):
    """F6 구조 추출 — 빈 양식과 예시 데이터 양식 양쪽에서 헤더·형식·코드표가 나온다."""
    # ① 빈 양식(배포 전) — F1이 만든 양식을 그대로 읽는다. 헤더 아래가 전부 비어 있다.
    blank = tmp / "빈양식.xlsx"
    fg.materialize(GOOD_WB, blank)
    st = fg.read_structure(blank)
    sheet = st["sheets"][0]
    assert sheet["header_row"] == 3, sheet["header_row"]
    assert sheet["headers"] == ["부서명*", "집행일자*", "집행금액*", "비고"], sheet["headers"]
    assert sheet["sample_rows"] == [], sheet["sample_rows"]      # 빈 양식이면 빈 배열(E5)
    dv = {c["letter"]: c["existing_validation"] for c in sheet["columns"]}
    assert dv["A"]["source"] == ["기획팀", "인사팀"], dv["A"]     # 기존 드롭다운 = 코드표 근거
    assert dv["B"]["type"] == "date" and dv["D"] is None, dv

    # ② 예시 데이터가 든 양식 — 표시형식·샘플이 채워지고 성명 열 값은 담지 않는다(§9.2 J4)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "교육이수"
    ws["A1"] = "2026년 상반기 교육 이수 현황"          # 제목 행(단일 병합 제목 흉내)
    for col, head in zip("ABCD", ["부서명", "성명", "이수일자", "이수시간"]):
        ws[f"{col}2"] = head
    ws.append(["기획팀", "홍길동", "2026-03-02", 8])
    ws.append(["인사팀", "김철수", "2026-03-05", 4])
    ws["C3"].number_format = "yyyy-mm-dd"
    filled = tmp / "예시.xlsx"
    wb.save(filled)

    st2 = fg.read_structure(filled)
    s2 = st2["sheets"][0]
    assert s2["headers"] == ["부서명", "성명", "이수일자", "이수시간"], s2["headers"]
    assert s2["header_row"] == 2, s2["header_row"]
    fmt = {c["header"]: c["existing_number_format"] for c in s2["columns"]}
    assert fmt["이수일자"] == "yyyy-mm-dd", fmt        # 형식은 추측이 아니라 파일에 적힌 값
    # 형식을 지정하지 않은 열은 근거가 없다는 뜻이다. `General`을 그대로 올려보내면 모델이
    # 그걸 입력 형식으로 옮겨 적어 화면에 'General'이 뜬다(실측에서 그랬다)
    assert fmt["부서명"] is None, fmt
    assert len(s2["sample_rows"]) == 2, s2["sample_rows"]
    assert s2["sample_rows"][0]["부서명"] == "기획팀"
    assert s2["sample_rows"][0]["성명"] == "(성명 마스킹)", s2["sample_rows"][0]

    # ③ 등록 대상이 아닌 것은 사유를 들고 막는다
    (tmp / "문서.hwpx").write_bytes(b"PK\x03\x04")
    for bad, why in [(tmp / "문서.hwpx", "xlsx"), (tmp / "없는파일.xlsx", "열 수 없")]:
        try:
            fg.read_structure(bad)
            raise AssertionError(f"{bad.name}이 통과했습니다")
        except fg.FormGenError as exc:
            assert why in str(exc), exc

    # ④ 표가 없는 파일(안내문만) — 등록할 근거가 없으므로 실패
    guide = openpyxl.Workbook()
    guide.active["A1"] = "작성 안내: 각 지사는 기한 내 제출하십시오"
    guide.save(tmp / "안내.xlsx")
    try:
        fg.read_structure(tmp / "안내.xlsx")
        raise AssertionError("표가 없는 파일이 통과했습니다")
    except fg.FormGenError as exc:
        assert "표" in str(exc), exc
    print("  ✓ 구조 추출(빈 양식·예시 데이터·기존 코드표·성명 마스킹·등록 불가 사유)")


def test_recognize_turn(tmp: Path):
    """F6 인지 턴 — 의도 판정과 헤더 1:1 루브릭이 임의 판정을 막는다."""
    structural = {"sheets": [{"name": "예산집행", "header_row": 3,
                              "headers": ["부서명*", "집행일자*", "집행금액*", "비고"],
                              "columns": [], "sample_rows": []}]}
    # 헤더 그대로가 항목명이어야 한다 — 취합이 이 이름으로 열을 찾는다
    exact = json.loads(json.dumps(GOOD_SPEC))
    for f, head in zip(exact["fields"], structural["sheets"][0]["headers"]):
        f["name"], f["confidence"] = head, "high"

    # ① 이름을 다듬어 내려주면(별표 제거) 완료로 인정하지 않고 무엇이 어긋났는지 밝힌다
    trimmed = json.loads(json.dumps(exact))
    trimmed["fields"][0]["name"] = "부서명"
    stub = Stub([{"intent": "recognize", "reply": "이렇게 읽었습니다",
                  "spec_json": trimmed, "recognize_complete": True}])
    out = fg.recognize_turn([{"role": "user", "content": "이 양식 그대로 쓸게"}], {},
                            structural, "m", client=stub)
    assert out["spec_complete"] is False, out
    assert "양식의 '부서명*' 항목 정보" in out["gaps"], out["gaps"]
    assert "'부서명'은 양식에 없는 항목입니다" in out["gaps"], out["gaps"]

    # ② 1:1로 맞으면 완료
    stub = Stub([{"intent": "recognize", "reply": "확인했습니다",
                  "spec_json": exact, "recognize_complete": True}])
    out = fg.recognize_turn([{"role": "user", "content": "이 양식 그대로"}], {},
                            structural, "m", client=stub)
    assert out["spec_complete"] is True and out["gaps"] == [], out

    # ③ 의도가 generate면 등록으로 넘기지 않는다(첨부는 참고가 된다 — F1-3)
    stub = Stub([{"intent": "generate", "reply": "새로 만들어 드릴까요?",
                  "spec_json": exact, "recognize_complete": True}])
    out = fg.recognize_turn([{"role": "user", "content": "이거 비슷하게 새로"}], {},
                            structural, "m", client=stub)
    assert out["intent"] == "generate" and out["spec_complete"] is False, out

    # ④ 판정을 못 받으면 ambiguous로 떨어뜨려 사용자에게 되묻는다(R6)
    stub = Stub([{"reply": "음", "spec_json": exact, "recognize_complete": True}])
    out = fg.recognize_turn([{"role": "user", "content": "?"}], {}, structural, "m", client=stub)
    assert out["intent"] == "ambiguous" and out["spec_complete"] is False, out

    # ④-2 근거 없는 항목(confidence low)은 확정으로 보지 않는다 — 실측에서 모델이 근거 없이
    # 전 항목을 '선택'으로 내려보냈고, 그대로 등록하면 검토가 누락을 못 잡는다
    unsure = json.loads(json.dumps(exact))
    unsure["fields"][2]["confidence"] = "low"
    stub = Stub([{"intent": "recognize", "reply": "필수 항목을 알려주세요",
                  "spec_json": unsure, "recognize_complete": True}])
    out = fg.recognize_turn([{"role": "user", "content": "이 양식 그대로"}], {},
                            structural, "m", client=stub)
    assert out["spec_complete"] is False, out
    assert f"'{structural['sheets'][0]['headers'][2]}' 확인 필요" in out["gaps"], out["gaps"]

    # ⑤ 표가 여러 개면 어느 표로 받을지 정해야 한다(실제 양식은 가이드·대장이 섞여 있다)
    many = {"sheets": [structural["sheets"][0],
                       {"name": "<참고> 작성 가이드", "header_row": 2,
                        "headers": ["항목", "문제점", "개선방안"],
                        "columns": [], "sample_rows": []}]}
    stub = Stub([{"intent": "recognize", "reply": "이 표로 받겠습니다",
                  "spec_json": exact, "recognize_complete": True}])
    out = fg.recognize_turn([{"role": "user", "content": "이 양식 그대로"}], {},
                            many, "m", client=stub)
    assert out["spec_complete"] is False, out
    assert out["gaps"][0].startswith("회신받을 표"), out["gaps"]
    assert "<참고> 작성 가이드" in out["gaps"][0], out["gaps"][0]

    # 표를 고르면 고른 표의 헤더만 1:1 대상이다(가이드 표의 헤더를 요구하지 않는다)
    picked = json.loads(json.dumps(exact)) | {"selected_sheets": ["예산집행"]}
    stub = Stub([{"intent": "recognize", "reply": "확인", "spec_json": picked,
                  "recognize_complete": True}])
    out = fg.recognize_turn([{"role": "user", "content": "첫 표로"}], {}, many, "m", client=stub)
    assert out["spec_complete"] is True and out["gaps"] == [], out

    # 표가 정해진 뒤에는 나머지 표를 **이름만** 보낸다 — 실측에서 표 6개(헤더 142개) 파일이
    # 매 턴 11.7k 입력 토큰을 썼고, 고른 뒤로는 그 대부분이 쓰이지 않는다. 이름을 남기는
    # 것은 사용자가 표를 바꿀 수 있어야 하기 때문이다
    stub = Stub([{"intent": "recognize", "reply": "확인", "spec_json": picked,
                  "recognize_complete": True}])
    fg.recognize_turn([{"role": "user", "content": "그대로"}], picked, many, "m", client=stub)
    sent = stub.seen[0]
    assert "문제점" not in sent, "고르지 않은 표의 헤더까지 보내고 있다"
    assert "<참고> 작성 가이드" in sent, "표를 바꿀 수 있도록 이름은 남아야 한다"
    assert "예산집행" in sent, sent[:200]

    # ⑥ 턴 상한에서 마감 지시 + 전송 payload 마스킹
    stub = Stub([{"intent": "recognize", "reply": "마감", "spec_json": exact,
                  "recognize_complete": True}])
    talk = [{"role": "user", "content": "010-1234-5678 로 연락"}
            for _ in range(fg.RECOGNIZE_TURN_LIMIT)]
    fg.recognize_turn(talk, {}, structural, "m", client=stub)
    assert "턴 수 상한" in stub.seen[0], stub.seen[0][:200]
    assert "010-1234-5678" not in stub.seen[0], "마스킹되지 않았습니다"
    print("  ✓ 인지 턴(의도 판정·헤더 1:1 루브릭·ambiguous 폴백·턴 상한·마스킹)")


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="formgen-test-"))
    try:
        test_rubric_guard()
        test_turn_limit_and_masking()
        test_validation_gate()
        test_normalize()
        test_author_retry()
        test_materialize(tmp)
        test_rules_bridge(tmp)
        test_read_structure(tmp)
        test_recognize_turn(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\n전체 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
