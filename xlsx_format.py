#!/usr/bin/env python3
"""취합 결과에 원본 양식의 서식을 입힌다.

aggregate.py의 `_write_block`은 값만 쓴다. 이 모듈이 원본 xlsx에서 서식을 읽고
(capture) 여러 회신본의 공통값을 고르고(consensus) 출력 시트에 입힌다
(open_sheet → 값 쓰기 → close_sheet).

설계·근거: docs/superpowers/specs/2026-08-07-취합문서-서식상속-design.md

서식은 부가물이다. 공개 함수는 예외를 밖으로 내지 않는다 — 서식을 못 입혀도
값은 반드시 나와야 한다.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_WIDTH = 60.0                    # 파일명처럼 긴 값이 열을 화면 밖으로 밀지 않게
MIN_WIDTH = 8.43                    # 엑셀 기본 열너비
DEFAULT_DATE_FORMAT = "yyyy-mm-dd"
_WIDTH_PER_CHAR = 1.4               # 한글이 섞이면 글자당 1자 너비로는 좁다
_FIT_SCAN_ROWS = 200                # 너비 계산에 훑는 최대 행수(전량 훑을 이유가 없다)

_GREY = PatternFill("solid", fgColor="DDDDDD")
_THIN = Side(style="thin")
_BOX = Border(bottom=_THIN, top=_THIN, left=_THIN, right=_THIN)


@dataclass
class CellStyle:
    """셀 하나의 서식. openpyxl 스타일 객체는 공유되므로 넣고 뺄 때 복사한다."""
    font: object = None
    fill: object = None
    border: object = None
    alignment: object = None

    @classmethod
    def of(cls, cell) -> "CellStyle":
        return cls(copy(cell.font), copy(cell.fill), copy(cell.border), copy(cell.alignment))

    def put(self, cell) -> None:
        if self.font is not None:
            cell.font = copy(self.font)
        if self.fill is not None:
            cell.fill = copy(self.fill)
        if self.border is not None:
            cell.border = copy(self.border)
        if self.alignment is not None:
            cell.alignment = copy(self.alignment)

    def key(self) -> str:
        """다수결에서 세기 위한 값 비교용 키."""
        return f"{self.font}|{self.fill}|{self.border}|{self.alignment}"


@dataclass
class TitleRow:
    cells: list = field(default_factory=list)   # [(열번호 1-base, 값, CellStyle)]
    height: float | None = None


@dataclass
class FormatTemplate:
    header_row: int = 1
    title_rows: list = field(default_factory=list)      # 1 ~ header_row-1
    title_merges: list = field(default_factory=list)    # (min_row, min_col, max_row, max_col)
    header_style: CellStyle | None = None
    data_style: CellStyle | None = None
    widths: dict = field(default_factory=dict)          # 헤더명 → 너비
    number_formats: dict = field(default_factory=dict)  # 헤더명 → 표시형식
    orientation: str = "portrait"
    fit_to_page: bool = False
    print_title_rows: str | None = None
    date_format: str = DEFAULT_DATE_FORMAT
    source_label: str = ""


def _is_date_format(nf: str) -> bool:
    low = (nf or "").lower()
    return "yy" in low or "mm-dd" in low or "m/d" in low


def capture(path, sheet_name: str, header_row: int) -> FormatTemplate | None:
    """원본 xlsx 한 시트에서 서식만 읽는다. 못 읽으면 None(예외를 내지 않는다)."""
    try:
        if header_row < 1:
            return None
        wb = openpyxl.load_workbook(Path(path))
        if sheet_name not in wb.sheetnames:
            return None
        ws = wb[sheet_name]
        if header_row > ws.max_row:
            return None

        # 헤더 이름 ← 컬럼 번호. 이 대응이 열너비·숫자서식 매칭의 기준이다
        headers: dict[int, str] = {}
        for c in range(1, ws.max_column + 1):
            value = ws.cell(row=header_row, column=c).value
            if value is not None and str(value).strip():
                headers[c] = str(value)
        if not headers:
            return None

        first = min(headers)
        tpl = FormatTemplate(
            header_row=header_row,
            header_style=CellStyle.of(ws.cell(row=header_row, column=first)),
            data_style=CellStyle.of(ws.cell(row=header_row + 1, column=first)),
            orientation=ws.page_setup.orientation or "portrait",
            fit_to_page=bool(ws.sheet_properties.pageSetUpPr
                             and ws.sheet_properties.pageSetUpPr.fitToPage),
            print_title_rows=ws.print_title_rows,
            source_label=Path(path).name,
        )

        for c, name in headers.items():
            # dict 인덱싱([])은 openpyxl에서 미설정 컬럼에도 기본폭(13.0)을 만들어
            # 낸다 — 실제로 폭을 지정한 컬럼만 담기 위해 .get()으로 존재 여부를 본다
            col_dim = ws.column_dimensions.get(get_column_letter(c))
            width = col_dim.width if col_dim is not None else None
            if width:
                tpl.widths[name] = round(float(width), 1)
            fmt = ws.cell(row=header_row + 1, column=c).number_format
            if fmt and fmt != "General":
                tpl.number_formats[name] = fmt
                if _is_date_format(fmt):
                    # 기본서식 폴백에서 쓸 날짜 표시형식 — 지어내지 않고 원본에서 본 것을 쓴다
                    tpl.date_format = fmt

        last_col = max(headers)
        for r in range(1, header_row):
            cells = [(c, ws.cell(row=r, column=c).value, CellStyle.of(ws.cell(row=r, column=c)))
                     for c in range(1, last_col + 1)]
            tpl.title_rows.append(TitleRow(cells, ws.row_dimensions[r].height))
        for rng in ws.merged_cells.ranges:
            if rng.max_row < header_row:
                tpl.title_merges.append((rng.min_row, rng.min_col, rng.max_row, rng.max_col))
        tpl.title_merges.sort()
        return tpl
    except Exception:
        return None


def basic(date_format: str = DEFAULT_DATE_FORMAT) -> FormatTemplate:
    """상속할 원본이 없을 때의 기본서식(설계 §6). 지어낸 스타일이 아니라 최소한의 가독성."""
    return FormatTemplate(
        header_row=1,
        header_style=CellStyle(Font(bold=True), _GREY, _BOX,
                               Alignment(horizontal="center", vertical="center")),
        data_style=None,
        date_format=date_format,
        source_label="기본서식",
    )
